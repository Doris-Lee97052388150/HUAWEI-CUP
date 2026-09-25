from __future__ import annotations

import argparse
import copy
import csv
import math
import random
import shutil
import sys
import time
from collections import defaultdict
from dataclasses import dataclass, field
from pathlib import Path
from typing import Dict, List, Tuple, Optional, Iterable, Set, Any

import numpy as np
from scipy.io import loadmat
from openpyxl import load_workbook, Workbook
import matplotlib.pyplot as plt


# ============================================================
# 0. 可统一修改的路径与算法参数
# ============================================================

# 固定项目根目录路径
PROJECT_ROOT = Path(r"C:\Users\77033\Desktop\D题")

# 固定输出目录
DEFAULT_OUTPUT_DIR = PROJECT_ROOT / "问题二输出"


RANDOM_SEED = 20260924
EARTH_RADIUS_M = 6_371_008.8
G = 9.81
MAX_STOPS_PER_ROUTE = 4               # 控制多点架次长度，便于求解与解释
ALNS_REMOVE_MIN = 4
ALNS_REMOVE_MAX = 9
DEFAULT_ALNS_ITERS = 800

# 最终综合目标权重（均在初始解基准上归一化）
OBJ_WEIGHTS = {
    "lateness": 0.45,
    "makespan": 0.25,
    "energy": 0.20,
    "trips": 0.10,
}

# Repair 局部代理代价：不调用完整 Decoder 时使用
PROXY_ENERGY_W = 1.0
PROXY_DURATION_W = 0.40       # 小时尺度
PROXY_TRIP_FIXED = 0.35
PROXY_TIGHT_SLACK_W = 0.80


# ============================================================
# 1. 数据对象
# ============================================================

@dataclass(frozen=True)
class Node:
    node_id: str
    name: str
    lon: float
    lat: float
    ground_alt: float
    work_alt: float


@dataclass(frozen=True)
class Box:
    box_id: str
    service_id: str
    material_type: str
    mass: float
    volume: float
    first_batch: bool
    first_deadline: Optional[float]
    expected_time: float
    priority: float
    hard_deadline: Optional[float]


@dataclass(frozen=True)
class DroneType:
    type_id: str
    name: str
    empty_mass: float
    max_payload: float
    max_volume: float
    cruise_speed: float
    empty_range: float
    full_range: float
    battery_energy: float
    reserve_ratio: float
    prep_time: float
    load_time_per_box: float
    service_base_time: float
    service_time_per_box: float
    climb_speed: float
    descent_speed: float
    climb_efficiency: float
    descent_efficiency: float


@dataclass(frozen=True)
class DroneUnit:
    drone_id: str
    type_id: str


@dataclass(frozen=True)
class BatterySpec:
    type_id: str
    count: int
    full_charge_time: float


@dataclass(frozen=True)
class Arc:
    from_node: str
    to_node: str
    distance_m: float
    dem_max_m: float
    cruise_alt_m: float
    start_work_alt_m: float
    end_work_alt_m: float
    climb_m: float
    descent_m: float
    flight_time_by_type: Dict[str, float]


@dataclass
class Stop:
    service_id: str
    box_ids: List[str]


@dataclass
class Route:
    stops: List[Stop]
    internal_id: int = 0

    def clone(self) -> "Route":
        return Route([Stop(s.service_id, list(s.box_ids)) for s in self.stops], self.internal_id)

    def box_ids(self) -> List[str]:
        return [b for s in self.stops for b in s.box_ids]

    def service_ids(self) -> List[str]:
        return [s.service_id for s in self.stops]

    def signature(self) -> Tuple:
        return tuple((s.service_id, tuple(sorted(s.box_ids))) for s in self.stops)

    def normalize(self) -> None:
        """合并同服务区 Stop、删除空 Stop，保持首次出现顺序。"""
        order: List[str] = []
        acc: Dict[str, List[str]] = defaultdict(list)
        for s in self.stops:
            if s.service_id not in acc:
                order.append(s.service_id)
            acc[s.service_id].extend(s.box_ids)
        self.stops = [Stop(sid, sorted(set(acc[sid]))) for sid in order if acc[sid]]


@dataclass
class LegEval:
    from_node: str
    to_node: str
    payload_mass: float
    payload_volume: float
    distance_m: float
    climb_m: float
    descent_m: float
    flight_time: float
    energy_kwh: float


@dataclass
class StopEval:
    service_id: str
    box_ids: List[str]
    delivery_mass: float
    delivery_volume: float
    arrival_offset: float
    service_duration: float
    delivery_offset: float


@dataclass
class RouteEval:
    feasible: bool
    reason: str
    drone_type: str
    total_boxes: int = 0
    total_mass: float = 0.0
    total_volume: float = 0.0
    pre_operation_time: float = 0.0
    takeoff_offset: float = 0.0
    duration: float = 0.0
    total_energy: float = 0.0
    end_soc: float = 0.0
    latest_start: float = math.inf
    soft_due_start: float = math.inf
    stops: List[StopEval] = field(default_factory=list)
    legs: List[LegEval] = field(default_factory=list)
    box_delivery_offsets: Dict[str, float] = field(default_factory=dict)


@dataclass
class BatteryState:
    battery_id: str
    type_id: str
    available_time: float = 0.0


@dataclass
class DroneState:
    drone_id: str
    type_id: str
    available_time: float = 0.0


@dataclass
class ScheduledTrip:
    route_index: int
    drone_id: str
    drone_type: str
    battery_id: str
    start_time: float
    takeoff_time: float
    return_time: float
    energy_kwh: float
    end_soc: float
    charge_end_time: float
    service_sequence: List[str]
    box_delivery: Dict[str, float]
    route_eval: RouteEval
    trip_id: str = ""


@dataclass
class DecodeResult:
    feasible: bool
    reason: str
    trips: List[ScheduledTrip] = field(default_factory=list)
    box_delivery: Dict[str, float] = field(default_factory=dict)
    bottleneck_routes: List[int] = field(default_factory=list)


@dataclass
class Metrics:
    weighted_lateness: float
    makespan: float
    energy: float
    trips: int
    hard_violations: int


# ============================================================
# 2. 文件定位与 Excel 读取
# ============================================================

def resolve_input_file(root: Path, base_name: str, suffix: str) -> Path:
    """优先按正式文件名查找；找不到时允许附件名带 (3)/(4) 等后缀。"""
    exact = list(root.rglob(base_name + suffix))
    if exact:
        return exact[0]
    candidates = [p for p in root.rglob(f"{base_name}*{suffix}") if "output" not in {x.lower() for x in p.parts}]
    if not candidates:
        raise FileNotFoundError(f"未找到文件：{base_name}{suffix}，搜索根目录：{root}")
    candidates.sort(key=lambda p: (len(p.name), len(str(p))))
    return candidates[0]


def cell_float(v: Any, default: Optional[float] = None) -> Optional[float]:
    if v is None or v == "":
        return default
    return float(v)


def load_nodes(path: Path) -> Dict[str, Node]:
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    nodes: Dict[str, Node] = {}

    # O01 固定在第 3 行，服务区从第 7 行起；同时做字符串识别，避免模板轻微变化。
    for row in ws.iter_rows(values_only=True):
        nid = str(row[0]).strip() if row and row[0] else ""
        if nid == "O01" or (nid.startswith("S") and len(nid) == 4):
            name = str(row[1]) if row[1] is not None else nid
            lon, lat, alt = float(row[2]), float(row[3]), float(row[4])
            work_alt = alt if nid == "O01" else alt + 30.0
            nodes[nid] = Node(nid, name, lon, lat, alt, work_alt)
    wb.close()
    if "O01" not in nodes or len(nodes) != 16:
        raise ValueError(f"节点读取异常：共读取 {len(nodes)} 个节点，应为 16 个（O01+15 服务区）")
    return nodes


def load_boxes(path: Path) -> Dict[str, Box]:
    wb = load_workbook(path, data_only=True, read_only=True)
    if "逐箱货箱清单" not in wb.sheetnames:
        raise ValueError("物资需求文件中未找到 '逐箱货箱清单' Sheet")
    ws = wb["逐箱货箱清单"]
    boxes: Dict[str, Box] = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row[0]:
            continue
        box_id = str(row[0]).strip()
        sid = str(row[1]).strip()
        material = str(row[2]).strip()
        mass = float(row[3])
        volume = float(row[4])
        first = str(row[5]).strip() == "是"
        first_deadline = cell_float(row[6])
        expected = float(row[7])
        priority = float(row[8])

        # 题意口径：医疗物资的期望送达时间作为硬要求；首批箱满足首批截止时间。
        hard_candidates: List[float] = []
        if first and first_deadline is not None:
            hard_candidates.append(first_deadline)
        if material == "医疗物资":
            hard_candidates.append(expected)
        hard_deadline = min(hard_candidates) if hard_candidates else None

        boxes[box_id] = Box(
            box_id, sid, material, mass, volume, first, first_deadline,
            expected, priority, hard_deadline
        )
    wb.close()
    if len(boxes) != 80:
        raise ValueError(f"逐箱货箱读取异常：读取 {len(boxes)} 箱，应为 80 箱")
    return boxes


def load_drone_data(path: Path) -> Tuple[Dict[str, DroneType], List[DroneUnit], Dict[str, BatterySpec]]:
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]

    types: Dict[str, DroneType] = {}
    for r in range(3, 6):
        vals = [ws.cell(r, c).value for c in range(1, 19)]
        tid = str(vals[0]).strip()
        types[tid] = DroneType(
            type_id=tid,
            name=str(vals[1]),
            empty_mass=float(vals[2]),
            max_payload=float(vals[3]),
            max_volume=float(vals[4]),
            cruise_speed=float(vals[5]),
            empty_range=float(vals[6]),
            full_range=float(vals[7]),
            battery_energy=float(vals[8]),
            reserve_ratio=float(vals[9]) / 100.0,
            prep_time=float(vals[10]),
            load_time_per_box=float(vals[11]),
            service_base_time=float(vals[12]),
            service_time_per_box=float(vals[13]),
            climb_speed=float(vals[14]),
            descent_speed=float(vals[15]),
            climb_efficiency=float(vals[16]),
            descent_efficiency=float(vals[17]),
        )

    drones: List[DroneUnit] = []
    for r in range(9, 17):
        did = ws.cell(r, 1).value
        if did:
            drones.append(DroneUnit(str(did).strip(), str(ws.cell(r, 2).value).strip()))

    batteries: Dict[str, BatterySpec] = {}
    for r in range(20, 23):
        tid = str(ws.cell(r, 1).value).strip()
        batteries[tid] = BatterySpec(tid, int(ws.cell(r, 2).value), float(ws.cell(r, 3).value))
    wb.close()
    return types, drones, batteries


# ============================================================
# 3. DEM 与 Arc Library
# ============================================================

def haversine_m(lon1: float, lat1: float, lon2: float, lat2: float) -> float:
    p1, p2 = math.radians(lat1), math.radians(lat2)
    dphi = math.radians(lat2 - lat1)
    dlambda = math.radians(lon2 - lon1)
    a = math.sin(dphi / 2) ** 2 + math.cos(p1) * math.cos(p2) * math.sin(dlambda / 2) ** 2
    return 2 * EARTH_RADIUS_M * math.asin(min(1.0, math.sqrt(a)))


def supercover_cells(r0: int, c0: int, r1: int, c1: int) -> List[Tuple[int, int]]:
    """整数栅格上的 supercover 直线，包含直线穿过/接触到的像元。"""
    x0, y0, x1, y1 = c0, r0, c1, r1
    dx, dy = x1 - x0, y1 - y0
    nx, ny = abs(dx), abs(dy)
    sx = 0 if dx == 0 else (1 if dx > 0 else -1)
    sy = 0 if dy == 0 else (1 if dy > 0 else -1)
    x, y = x0, y0
    cells = {(y, x)}
    ix = iy = 0
    while ix < nx or iy < ny:
        decision = (1 + 2 * ix) * ny - (1 + 2 * iy) * nx
        if decision == 0:
            # 穿过网格角点时把两侧像元也纳入 supercover
            if ix < nx:
                cells.add((y, x + sx))
            if iy < ny:
                cells.add((y + sy, x))
            if ix < nx:
                x += sx; ix += 1
            if iy < ny:
                y += sy; iy += 1
        elif decision < 0:
            x += sx; ix += 1
        else:
            y += sy; iy += 1
        cells.add((y, x))
    return list(cells)


class ArcLibrary:
    def __init__(self, nodes: Dict[str, Node], drone_types: Dict[str, DroneType], dem_path: Path):
        self.nodes = nodes
        self.drone_types = drone_types
        mat = loadmat(dem_path)
        self.dem = np.asarray(mat["dem"], dtype=float)
        self.lat = np.asarray(mat["latitude"], dtype=float).reshape(-1)
        self.lon = np.asarray(mat["longitude"], dtype=float).reshape(-1)
        nodata = mat.get("nodata", np.array([[-32767.0]]))
        self.nodata = float(np.asarray(nodata).reshape(-1)[0])
        self.arcs: Dict[Tuple[str, str], Arc] = {}
        self._node_rc: Dict[str, Tuple[int, int]] = {}
        self._prepare_indices()
        self._build()

    def _prepare_indices(self) -> None:
        lon_min, lon_max = float(self.lon.min()), float(self.lon.max())
        lat_min, lat_max = float(self.lat.min()), float(self.lat.max())
        for nid, n in self.nodes.items():
            if not (lon_min <= n.lon <= lon_max and lat_min <= n.lat <= lat_max):
                raise ValueError(f"节点 {nid} 超出 DEM 范围")
            r = int(np.argmin(np.abs(self.lat - n.lat)))
            c = int(np.argmin(np.abs(self.lon - n.lon)))
            self._node_rc[nid] = (r, c)

    def _build(self) -> None:
        ids = sorted(self.nodes.keys(), key=lambda x: (x != "O01", x))
        pair_geo: Dict[Tuple[str, str], Tuple[float, float]] = {}
        for i, a in enumerate(ids):
            for b in ids[i + 1:]:
                na, nb = self.nodes[a], self.nodes[b]
                dist = haversine_m(na.lon, na.lat, nb.lon, nb.lat)
                ra, ca = self._node_rc[a]
                rb, cb = self._node_rc[b]
                cells = supercover_cells(ra, ca, rb, cb)
                vals = []
                for r, c in cells:
                    if 0 <= r < self.dem.shape[0] and 0 <= c < self.dem.shape[1]:
                        v = self.dem[r, c]
                        if np.isfinite(v) and abs(v - self.nodata) > 1e-6:
                            vals.append(float(v))
                if not vals:
                    raise ValueError(f"航段 {a}-{b} 没有有效 DEM 像元")
                pair_geo[(a, b)] = (dist, max(vals))

        for (a, b), (dist, dem_max) in pair_geo.items():
            for u, v in [(a, b), (b, a)]:
                nu, nv = self.nodes[u], self.nodes[v]
                cruise_alt = dem_max + 50.0
                climb = max(0.0, cruise_alt - nu.work_alt)
                descent = max(0.0, cruise_alt - nv.work_alt)
                times = {}
                for tid, dt in self.drone_types.items():
                    times[tid] = climb / dt.climb_speed + dist / dt.cruise_speed + descent / dt.descent_speed
                self.arcs[(u, v)] = Arc(
                    u, v, dist, dem_max, cruise_alt,
                    nu.work_alt, nv.work_alt, climb, descent, times
                )

        # 基本一致性检查
        for (a, b), arc in list(self.arcs.items()):
            rev = self.arcs[(b, a)]
            if abs(arc.distance_m - rev.distance_m) > 1e-6 or abs(arc.dem_max_m - rev.dem_max_m) > 1e-6:
                raise AssertionError(f"Arc 对称性检查失败：{a}<->{b}")

    def __getitem__(self, key: Tuple[str, str]) -> Arc:
        return self.arcs[key]

    def export_csv(self, path: Path) -> None:
        with path.open("w", newline="", encoding="utf-8-sig") as f:
            w = csv.writer(f)
            w.writerow(["from", "to", "distance_m", "dem_max_m", "cruise_alt_m", "climb_m", "descent_m", "time_A", "time_B", "time_C"])
            for (a, b), arc in sorted(self.arcs.items()):
                w.writerow([
                    a, b, arc.distance_m, arc.dem_max_m, arc.cruise_alt_m,
                    arc.climb_m, arc.descent_m,
                    arc.flight_time_by_type.get("A", ""),
                    arc.flight_time_by_type.get("B", ""),
                    arc.flight_time_by_type.get("C", ""),
                ])


# ============================================================
# 4. Route Evaluator
# ============================================================

class RouteEvaluator:
    def __init__(self, boxes: Dict[str, Box], drone_types: Dict[str, DroneType], arcs: ArcLibrary):
        self.boxes = boxes
        self.drone_types = drone_types
        self.arcs = arcs
        self.cache: Dict[Tuple[Tuple, str], RouteEval] = {}

    @staticmethod
    def equivalent_range(dt: DroneType, payload: float) -> float:
        ratio = min(max(payload / dt.max_payload, 0.0), 1.0)
        return dt.empty_range - (dt.empty_range - dt.full_range) * (ratio ** 1.5)

    def leg_energy(self, dt: DroneType, arc: Arc, payload: float) -> float:
        """
        计算约定：
        1) 水平能耗 = 电池可用能量 × 水平距离 / 当前载荷等效航程；
        2) 爬升附加能耗 = m g h / eta，换算为 kWh；
        3) 下降附加能耗效率为 0 时不额外计算下降能耗。
        """
        rng = self.equivalent_range(dt, payload)
        if rng <= 0:
            return math.inf
        e_hor = dt.battery_energy * arc.distance_m / rng
        total_mass = dt.empty_mass + payload
        e_up = total_mass * G * arc.climb_m / max(dt.climb_efficiency, 1e-9) / 3.6e6
        return e_hor + e_up

    def evaluate(self, route: Route, type_id: str) -> RouteEval:
        route.normalize()
        sig = (route.signature(), type_id)
        if sig in self.cache:
            return self.cache[sig]
        dt = self.drone_types[type_id]

        # ---------- 结构合法性 ----------
        if not route.stops:
            out = RouteEval(False, "EMPTY_ROUTE", type_id)
            self.cache[sig] = out; return out
        if len(route.stops) > MAX_STOPS_PER_ROUTE:
            out = RouteEval(False, "TOO_MANY_STOPS", type_id)
            self.cache[sig] = out; return out
        service_ids = [s.service_id for s in route.stops]
        if len(service_ids) != len(set(service_ids)):
            out = RouteEval(False, "DUPLICATE_SERVICE", type_id)
            self.cache[sig] = out; return out
        all_boxes = route.box_ids()
        if len(all_boxes) != len(set(all_boxes)):
            out = RouteEval(False, "DUPLICATE_BOX", type_id)
            self.cache[sig] = out; return out
        for s in route.stops:
            if not s.box_ids:
                out = RouteEval(False, "EMPTY_STOP", type_id)
                self.cache[sig] = out; return out
            for bid in s.box_ids:
                if bid not in self.boxes or self.boxes[bid].service_id != s.service_id:
                    out = RouteEval(False, "BOX_DESTINATION_MISMATCH", type_id)
                    self.cache[sig] = out; return out

        stop_mass = [sum(self.boxes[b].mass for b in s.box_ids) for s in route.stops]
        stop_vol = [sum(self.boxes[b].volume for b in s.box_ids) for s in route.stops]
        total_mass, total_vol = sum(stop_mass), sum(stop_vol)
        nbox = len(all_boxes)
        if total_mass > dt.max_payload + 1e-9:
            out = RouteEval(False, "MASS_CAPACITY_EXCEEDED", type_id, nbox, total_mass, total_vol)
            self.cache[sig] = out; return out
        if total_vol > dt.max_volume + 1e-12:
            out = RouteEval(False, "VOLUME_CAPACITY_EXCEEDED", type_id, nbox, total_mass, total_vol)
            self.cache[sig] = out; return out

        # 后缀载荷
        k = len(route.stops)
        suffix_mass = [0.0] * (k + 1)
        suffix_vol = [0.0] * (k + 1)
        for i in range(k - 1, -1, -1):
            suffix_mass[i] = suffix_mass[i + 1] + stop_mass[i]
            suffix_vol[i] = suffix_vol[i + 1] + stop_vol[i]

        pre = dt.prep_time + nbox * dt.load_time_per_box
        current_time = pre
        stop_evals: List[StopEval] = []
        leg_evals: List[LegEval] = []
        total_energy = 0.0
        box_offsets: Dict[str, float] = {}

        nodes = ["O01"] + service_ids + ["O01"]
        # 去程 + 服务
        for i, stop in enumerate(route.stops):
            a, b = nodes[i], nodes[i + 1]
            arc = self.arcs[(a, b)]
            payload = suffix_mass[i]
            pvol = suffix_vol[i]
            ft = arc.flight_time_by_type[type_id]
            en = self.leg_energy(dt, arc, payload)
            current_time += ft
            arrival = current_time
            service_time = dt.service_base_time + len(stop.box_ids) * dt.service_time_per_box
            current_time += service_time
            delivery = current_time
            for bid in stop.box_ids:
                box_offsets[bid] = delivery
            total_energy += en
            leg_evals.append(LegEval(a, b, payload, pvol, arc.distance_m, arc.climb_m, arc.descent_m, ft, en))
            stop_evals.append(StopEval(stop.service_id, list(stop.box_ids), stop_mass[i], stop_vol[i], arrival, service_time, delivery))

        # 返程：空载
        a, b = nodes[-2], "O01"
        arc = self.arcs[(a, b)]
        ft = arc.flight_time_by_type[type_id]
        en = self.leg_energy(dt, arc, 0.0)
        current_time += ft
        total_energy += en
        leg_evals.append(LegEval(a, b, 0.0, 0.0, arc.distance_m, arc.climb_m, arc.descent_m, ft, en))

        end_soc = 1.0 - total_energy / dt.battery_energy
        if end_soc + 1e-12 < dt.reserve_ratio:
            out = RouteEval(
                False, "ENERGY_RESERVE_VIOLATION", type_id, nbox, total_mass, total_vol,
                pre, pre, current_time, total_energy, end_soc,
                stops=stop_evals, legs=leg_evals, box_delivery_offsets=box_offsets
            )
            self.cache[sig] = out; return out

        hard_limits = []
        soft_limits = []
        for bid, off in box_offsets.items():
            bobj = self.boxes[bid]
            if bobj.hard_deadline is not None:
                hard_limits.append(bobj.hard_deadline - off)
            soft_limits.append(bobj.expected_time - off)
        latest = min(hard_limits) if hard_limits else math.inf
        soft_due = min(soft_limits) if soft_limits else math.inf
        if latest < -1e-9:
            out = RouteEval(
                False, "HARD_DEADLINE_IMPOSSIBLE", type_id, nbox, total_mass, total_vol,
                pre, pre, current_time, total_energy, end_soc,
                latest, soft_due, stop_evals, leg_evals, box_offsets
            )
            self.cache[sig] = out; return out

        out = RouteEval(
            True, "OK", type_id, nbox, total_mass, total_vol,
            pre, pre, current_time, total_energy, end_soc,
            latest, soft_due, stop_evals, leg_evals, box_offsets
        )
        self.cache[sig] = out
        return out

    def feasible_types(self, route: Route) -> Dict[str, RouteEval]:
        out = {}
        for tid in self.drone_types:
            ev = self.evaluate(route, tid)
            if ev.feasible:
                out[tid] = ev
        return out

    def proxy_cost(self, route: Route) -> float:
        fes = self.feasible_types(route)
        if not fes:
            return math.inf
        vals = []
        for ev in fes.values():
            tight = 0.0
            if math.isfinite(ev.latest_start):
                tight = max(0.0, 600.0 - ev.latest_start) / 600.0
            vals.append(
                PROXY_ENERGY_W * ev.total_energy
                + PROXY_DURATION_W * (ev.duration / 3600.0)
                + PROXY_TRIP_FIXED
                + PROXY_TIGHT_SLACK_W * tight
            )
        return min(vals)


# ============================================================
# 5. 路线构造与 Regret-2 Repair
# ============================================================

def ensure_route_integrity(routes: List[Route], boxes: Dict[str, Box], allow_missing: bool = False) -> None:
    seen: List[str] = []
    for r in routes:
        r.normalize()
        for s in r.stops:
            for b in s.box_ids:
                if boxes[b].service_id != s.service_id:
                    raise AssertionError(f"货箱 {b} 被挂到错误服务区 {s.service_id}")
                seen.append(b)
    if len(seen) != len(set(seen)):
        raise AssertionError("Solution 中存在重复货箱")
    if not allow_missing and set(seen) != set(boxes):
        miss = set(boxes) - set(seen)
        extra = set(seen) - set(boxes)
        raise AssertionError(f"货箱完整性失败：missing={sorted(miss)}, extra={sorted(extra)}")


def route_from_box(box: Box) -> Route:
    return Route([Stop(box.service_id, [box.box_id])])


def insert_box_into_route(route: Route, box: Box, position: Optional[int]) -> Route:
    nr = route.clone()
    # 同服务区优先直接挂载，不增加访问点
    for s in nr.stops:
        if s.service_id == box.service_id:
            s.box_ids.append(box.box_id)
            nr.normalize()
            return nr
    if position is None:
        position = len(nr.stops)
    nr.stops.insert(position, Stop(box.service_id, [box.box_id]))
    nr.normalize()
    return nr


def best_order_for_stops(route: Route, evaluator: RouteEvaluator, arc_lib: ArcLibrary) -> Route:
    """少量 Stop 下用全排列（<=4），选局部代理代价最小的顺序。"""
    import itertools
    if len(route.stops) <= 1:
        return route.clone()
    best = route.clone()
    best_c = evaluator.proxy_cost(best)
    # MAX_STOPS_PER_ROUTE <=4，因此最多 24 个排列，完全可承受
    for perm in itertools.permutations(route.stops):
        nr = Route([Stop(s.service_id, list(s.box_ids)) for s in perm])
        c = evaluator.proxy_cost(nr)
        if c < best_c - 1e-10:
            best, best_c = nr, c
    return best


def construct_initial_routes(boxes: Dict[str, Box], evaluator: RouteEvaluator) -> List[Route]:
    """
    初始解（以“先保证硬时限可调度”为第一原则）：
    1) 每个服务区的硬时限箱形成独立紧急种子，不往紧急种子中塞普通箱，
       避免原本 A/B 可执行的紧急任务因增载而被迫占用稀缺的 C 型机；
    2) 普通箱按服务区单独进行容量装箱；
    3) 只对“不含硬时限箱”的普通路线进行跨服务区贪心 merge。

    这样得到的初始解通常架次数略多，但资源 Decoder 更容易首先得到全局可行解；
    后续 ALNS 再负责把普通任务、甚至部分有裕量的任务重新合并。
    """
    by_service: Dict[str, List[Box]] = defaultdict(list)
    for b in boxes.values():
        by_service[b.service_id].append(b)

    routes: List[Route] = []
    assigned: Set[str] = set()

    # I. 硬时限种子：严格保持轻载，优先保证 A/B/C 机型选择灵活性
    services = sorted(by_service)
    for sid in services:
        hard = sorted(
            [b for b in by_service[sid] if b.hard_deadline is not None],
            key=lambda b: (b.hard_deadline, -b.priority, b.box_id)
        )
        if not hard:
            continue
        r = Route([Stop(sid, [b.box_id for b in hard])])
        if evaluator.feasible_types(r):
            routes.append(r)
            assigned.update(r.box_ids())
        else:
            # 极端情况下拆成单箱，确保不会因为组批导致物理不可行
            for b in hard:
                rb = route_from_box(b)
                if not evaluator.feasible_types(rb):
                    raise RuntimeError(f"单箱硬时限任务都不可行：{b.box_id}")
                routes.append(rb)
                assigned.add(b.box_id)

    # II. 普通箱：在同一服务区内做顺序装箱；不占用硬时限种子的容量
    for sid in services:
        soft = sorted(
            [b for b in by_service[sid] if b.box_id not in assigned],
            key=lambda b: (b.expected_time, -b.priority, -b.mass, b.box_id)
        )
        current: Optional[Route] = None
        for b in soft:
            if current is None:
                current = route_from_box(b)
                continue
            trial = insert_box_into_route(current, b, None)
            if evaluator.feasible_types(trial):
                current = trial
            else:
                routes.append(current)
                current = route_from_box(b)
        if current is not None:
            routes.append(current)

    ensure_route_integrity(routes, boxes)

    def contains_hard(r: Route) -> bool:
        return any(boxes[bid].hard_deadline is not None for bid in r.box_ids())

    # III. 仅合并普通路线，减少架次；紧急种子先不碰，确保初始 Decoder 稳定可行
    improved = True
    while improved:
        improved = False
        best_merge = None
        for i in range(len(routes)):
            if contains_hard(routes[i]):
                continue
            for j in range(i + 1, len(routes)):
                if contains_hard(routes[j]):
                    continue
                ra, rb = routes[i], routes[j]
                service_union = set(ra.service_ids()) | set(rb.service_ids())
                if len(service_union) > MAX_STOPS_PER_ROUTE:
                    continue
                merged_map: Dict[str, List[str]] = defaultdict(list)
                order = []
                for r in (ra, rb):
                    for s in r.stops:
                        if s.service_id not in merged_map:
                            order.append(s.service_id)
                        merged_map[s.service_id].extend(s.box_ids)
                merged = Route([Stop(sid, merged_map[sid]) for sid in order])
                merged = best_order_for_stops(merged, evaluator, evaluator.arcs)
                cm = evaluator.proxy_cost(merged)
                if not math.isfinite(cm):
                    continue
                old = evaluator.proxy_cost(ra) + evaluator.proxy_cost(rb)
                saving = old - cm
                if saving > 0.02 and (best_merge is None or saving > best_merge[0]):
                    best_merge = (saving, i, j, merged)
        if best_merge:
            _, i, j, merged = best_merge
            routes[i] = merged
            routes.pop(j)
            improved = True

    ensure_route_integrity(routes, boxes)
    return routes

def remove_boxes_from_routes(routes: List[Route], remove_ids: Set[str]) -> List[Route]:
    out = []
    for r in routes:
        nr = r.clone()
        for s in nr.stops:
            s.box_ids = [b for b in s.box_ids if b not in remove_ids]
        nr.normalize()
        if nr.stops:
            out.append(nr)
    return out


def enumerate_insert_options(routes: List[Route], box: Box, evaluator: RouteEvaluator) -> List[Tuple[float, int, Route]]:
    opts: List[Tuple[float, int, Route]] = []
    # 现有路线
    for i, r in enumerate(routes):
        old = evaluator.proxy_cost(r)
        if box.service_id in r.service_ids():
            nr = insert_box_into_route(r, box, None)
            c = evaluator.proxy_cost(nr)
            if math.isfinite(c):
                opts.append((c - old, i, nr))
        elif len(r.stops) < MAX_STOPS_PER_ROUTE:
            for pos in range(len(r.stops) + 1):
                nr = insert_box_into_route(r, box, pos)
                c = evaluator.proxy_cost(nr)
                if math.isfinite(c):
                    opts.append((c - old, i, nr))
    # 新路线，index == len(routes)
    nr = route_from_box(box)
    c = evaluator.proxy_cost(nr)
    if math.isfinite(c):
        opts.append((c, len(routes), nr))
    opts.sort(key=lambda x: x[0])
    return opts


def regret2_repair(routes: List[Route], unassigned_ids: List[str], boxes: Dict[str, Box], evaluator: RouteEvaluator) -> Optional[List[Route]]:
    routes = [r.clone() for r in routes]
    pool = set(unassigned_ids)
    while pool:
        best_choice = None
        for bid in list(pool):
            b = boxes[bid]
            opts = enumerate_insert_options(routes, b, evaluator)
            if not opts:
                return None
            best_cost = opts[0][0]
            second = opts[1][0] if len(opts) >= 2 else best_cost + 10.0
            regret = second - best_cost
            # 硬时限/早期望箱在 regret 相同情况下优先
            urgency = (b.hard_deadline if b.hard_deadline is not None else b.expected_time)
            key = (regret, -best_cost, -b.priority, -1.0 / max(urgency, 1.0))
            if best_choice is None or key > best_choice[0]:
                best_choice = (key, bid, opts[0])
        _, bid, (cost, idx, nr) = best_choice
        if idx == len(routes):
            routes.append(nr)
        else:
            routes[idx] = nr
        pool.remove(bid)
    return routes


# ============================================================
# 6. Resource Decoder
# ============================================================

def charge_time_seconds(soc: float, full_charge_time: float) -> float:
    soc = min(max(soc, 0.0), 1.0)
    if soc >= 1.0 - 1e-12:
        return 0.0
    if soc >= 0.9:
        return full_charge_time * 0.35 * (1.0 - soc) / 0.1
    return full_charge_time * (0.65 * (0.9 - soc) / 0.9 + 0.35)


class ResourceDecoder:
    def __init__(
        self,
        evaluator: RouteEvaluator,
        drone_units: List[DroneUnit],
        batteries: Dict[str, BatterySpec],
        boxes: Dict[str, Box],
    ):
        self.evaluator = evaluator
        self.drone_units = drone_units
        self.battery_specs = batteries
        self.boxes = boxes

    def decode(self, routes: List[Route]) -> DecodeResult:
        drones: Dict[str, DroneState] = {
            d.drone_id: DroneState(d.drone_id, d.type_id, 0.0) for d in self.drone_units
        }
        battery_states: Dict[str, BatteryState] = {}
        for tid, spec in self.battery_specs.items():
            for i in range(1, spec.count + 1):
                bid = f"BAT-{tid}{i:02d}"
                battery_states[bid] = BatteryState(bid, tid, 0.0)

        unscheduled = set(range(len(routes)))
        trips: List[ScheduledTrip] = []

        # 每条路线的三机型评价一次性取好
        route_evals: Dict[int, Dict[str, RouteEval]] = {}
        for i, r in enumerate(routes):
            fes = self.evaluator.feasible_types(r)
            if not fes:
                return DecodeResult(False, f"ROUTE_{i}_NO_FEASIBLE_TYPE", bottleneck_routes=[i])
            route_evals[i] = fes

        while unscheduled:
            route_choices = []
            for ri in unscheduled:
                type_options = []
                r = routes[ri]
                for tid, ev in route_evals[ri].items():
                    ds = [d for d in drones.values() if d.type_id == tid]
                    bs = [b for b in battery_states.values() if b.type_id == tid]
                    if not ds or not bs:
                        continue
                    d = min(ds, key=lambda x: (x.available_time, x.drone_id))
                    b = min(bs, key=lambda x: (x.available_time, x.battery_id))
                    start = max(d.available_time, b.available_time)
                    finish = start + ev.duration
                    deadline_ok = (not math.isfinite(ev.latest_start)) or start <= ev.latest_start + 1e-9
                    hard_slack = (ev.latest_start - start) if math.isfinite(ev.latest_start) else math.inf
                    soft_slack = ev.soft_due_start - start
                    # 类型评分：先按时，再早完成，再低能耗；轻任务不强制占 C
                    capacity_ratio = ev.total_mass / max(self.evaluator.drone_types[tid].max_payload, 1e-9)
                    type_score = finish + 80.0 * ev.total_energy + 60.0 * max(0.0, 0.25 - capacity_ratio)
                    type_options.append((deadline_ok, type_score, tid, d, b, start, finish, ev, hard_slack, soft_slack))

                feasible_now = [x for x in type_options if x[0]]
                if feasible_now:
                    # 路线紧迫度：硬时限优先，其次软期望
                    best_opt = min(feasible_now, key=lambda x: x[1])
                    has_hard = math.isfinite(best_opt[7].latest_start)
                    priority_key = (
                        0 if has_hard else 1,
                        best_opt[8] if has_hard else best_opt[9],
                        best_opt[6],
                    )
                    route_choices.append((priority_key, ri, best_opt))
                else:
                    # 当前资源状态下没有机型能满足硬最晚开始，但可能是排序导致；记录为危急候选
                    if type_options:
                        best_bad = min(type_options, key=lambda x: (x[8], x[1]))
                        route_choices.append((( -1, best_bad[8], best_bad[6]), ri, best_bad))

            if not route_choices:
                return DecodeResult(False, "NO_SCHEDULABLE_ROUTE", trips=trips, bottleneck_routes=sorted(unscheduled))

            route_choices.sort(key=lambda x: x[0])
            _, ri, opt = route_choices[0]
            deadline_ok, _, tid, d, b, start, finish, ev, hard_slack, soft_slack = opt
            if not deadline_ok:
                # 当前 EDF/LST 列表调度失败。返回关键路线，交给 ALNS 改结构。
                critical = [x[1] for x in route_choices[:min(5, len(route_choices))]]
                return DecodeResult(False, "GLOBAL_SCHEDULE_INFEASIBLE_HARD_DEADLINE", trips=trips, bottleneck_routes=critical)

            delivery = {bid: start + off for bid, off in ev.box_delivery_offsets.items()}
            spec = self.battery_specs[tid]
            chg = charge_time_seconds(ev.end_soc, spec.full_charge_time)
            charge_end = finish + chg
            trip = ScheduledTrip(
                route_index=ri,
                drone_id=d.drone_id,
                drone_type=tid,
                battery_id=b.battery_id,
                start_time=start,
                takeoff_time=start + ev.takeoff_offset,
                return_time=finish,
                energy_kwh=ev.total_energy,
                end_soc=ev.end_soc,
                charge_end_time=charge_end,
                service_sequence=routes[ri].service_ids(),
                box_delivery=delivery,
                route_eval=ev,
            )
            trips.append(trip)
            d.available_time = finish
            b.available_time = charge_end
            unscheduled.remove(ri)

        # 最终按开始时间编号，避免优化过程内部 ID 混乱
        trips.sort(key=lambda t: (t.start_time, t.drone_id, t.return_time))
        box_delivery: Dict[str, float] = {}
        for idx, t in enumerate(trips, start=1):
            t.trip_id = f"Q2T{idx:03d}"
            for bid, tm in t.box_delivery.items():
                if bid in box_delivery:
                    return DecodeResult(False, "DUPLICATE_BOX_AFTER_DECODING", trips=trips)
                box_delivery[bid] = tm
        return DecodeResult(True, "OK", trips, box_delivery)


# ============================================================
# 7. 指标、校验与 ALNS
# ============================================================

def compute_metrics(decoded: DecodeResult, boxes: Dict[str, Box]) -> Metrics:
    if not decoded.feasible:
        return Metrics(math.inf, math.inf, math.inf, 10**9, 10**9)
    late = 0.0
    hard_v = 0
    for bid, b in boxes.items():
        t = decoded.box_delivery[bid]
        late += b.priority * max(0.0, t - b.expected_time) / max(b.expected_time, 1.0)
        if b.hard_deadline is not None and t > b.hard_deadline + 1e-6:
            hard_v += 1
    makespan = max((t.return_time for t in decoded.trips), default=0.0)
    energy = sum(t.energy_kwh for t in decoded.trips)
    return Metrics(late, makespan, energy, len(decoded.trips), hard_v)


def normalized_score(m: Metrics, base: Metrics) -> float:
    if m.hard_violations > 0 or not all(math.isfinite(x) for x in [m.weighted_lateness, m.makespan, m.energy]):
        return math.inf
    late_scale = max(base.weighted_lateness, 1.0)
    makespan_scale = max(base.makespan, 1.0)
    energy_scale = max(base.energy, 1e-6)
    trip_scale = max(base.trips, 1)
    return (
        OBJ_WEIGHTS["lateness"] * m.weighted_lateness / late_scale
        + OBJ_WEIGHTS["makespan"] * m.makespan / makespan_scale
        + OBJ_WEIGHTS["energy"] * m.energy / energy_scale
        + OBJ_WEIGHTS["trips"] * m.trips / trip_scale
    )


def validate_final_solution(routes: List[Route], decoded: DecodeResult, boxes: Dict[str, Box], drone_types: Dict[str, DroneType]) -> None:
    ensure_route_integrity(routes, boxes)
    if not decoded.feasible:
        raise AssertionError(f"Decoder 不可行：{decoded.reason}")
    if set(decoded.box_delivery) != set(boxes):
        raise AssertionError("最终逐箱交付不完整")
    # 硬时限
    for bid, b in boxes.items():
        if b.hard_deadline is not None and decoded.box_delivery[bid] > b.hard_deadline + 1e-6:
            raise AssertionError(f"硬时限违反：{bid}, t={decoded.box_delivery[bid]:.2f}, ddl={b.hard_deadline}")
    # 无人机占用不重叠
    by_drone = defaultdict(list)
    by_bat = defaultdict(list)
    for t in decoded.trips:
        by_drone[t.drone_id].append((t.start_time, t.return_time, t.trip_id))
        by_bat[t.battery_id].append((t.start_time, t.charge_end_time, t.trip_id))
        if t.end_soc + 1e-9 < drone_types[t.drone_type].reserve_ratio:
            raise AssertionError(f"返航 SOC 不足：{t.trip_id}")
    for did, ints in by_drone.items():
        ints.sort()
        for a, b in zip(ints, ints[1:]):
            if a[1] > b[0] + 1e-7:
                raise AssertionError(f"无人机冲突 {did}: {a} vs {b}")
    for bid, ints in by_bat.items():
        ints.sort()
        for a, b in zip(ints, ints[1:]):
            if a[1] > b[0] + 1e-7:
                raise AssertionError(f"电池占用/充电冲突 {bid}: {a} vs {b}")


def destroy_random(routes: List[Route], q: int, rng: random.Random) -> Tuple[List[Route], List[str]]:
    all_ids = [b for r in routes for b in r.box_ids()]
    q = min(q, len(all_ids))
    rem = set(rng.sample(all_ids, q))
    return remove_boxes_from_routes(routes, rem), list(rem)


def destroy_related(routes: List[Route], q: int, boxes: Dict[str, Box], nodes: Dict[str, Node], rng: random.Random) -> Tuple[List[Route], List[str]]:
    all_ids = [b for r in routes for b in r.box_ids()]
    seed = rng.choice(all_ids)
    sb = boxes[seed]
    ns = nodes[sb.service_id]
    ranked = sorted(
        all_ids,
        key=lambda bid: (
            haversine_m(ns.lon, ns.lat, nodes[boxes[bid].service_id].lon, nodes[boxes[bid].service_id].lat)
            + 0.15 * abs(sb.expected_time - boxes[bid].expected_time)
        )
    )
    rem = set(ranked[:min(q, len(ranked))])
    return remove_boxes_from_routes(routes, rem), list(rem)


def destroy_worst(routes: List[Route], q: int, boxes: Dict[str, Box], decoded: DecodeResult, rng: random.Random) -> Tuple[List[Route], List[str]]:
    # 对迟到贡献大的箱优先；若均不迟到，则优先晚送达/低裕量箱
    scored = []
    for bid, b in boxes.items():
        t = decoded.box_delivery.get(bid, 0.0)
        late = b.priority * max(0.0, t - b.expected_time) / max(b.expected_time, 1.0)
        ratio = t / max(b.expected_time, 1.0)
        scored.append((late * 1000 + ratio, rng.random(), bid))
    scored.sort(reverse=True)
    rem = {x[2] for x in scored[:min(q, len(scored))]}
    return remove_boxes_from_routes(routes, rem), list(rem)


def alns_optimize(
    initial_routes: List[Route],
    boxes: Dict[str, Box],
    nodes: Dict[str, Node],
    evaluator: RouteEvaluator,
    decoder: ResourceDecoder,
    iterations: int,
    seed: int,
) -> Tuple[List[Route], DecodeResult, Metrics, List[Tuple[int, float]]]:
    rng = random.Random(seed)
    current_routes = [r.clone() for r in initial_routes]
    current_dec = decoder.decode(current_routes)
    if not current_dec.feasible:
        # 初始 Decoder 失败时，使用最保守退化方案：每个 Route 只保留单服务区/必要时单箱。
        conservative = []
        for r in current_routes:
            for s in r.stops:
                # 同服务区按 C 型容量逐箱顺序装箱，直到不能再加
                cr: Optional[Route] = None
                for bid in s.box_ids:
                    b = boxes[bid]
                    if cr is None:
                        cr = route_from_box(b)
                    else:
                        trial = insert_box_into_route(cr, b, None)
                        if evaluator.feasible_types(trial):
                            cr = trial
                        else:
                            conservative.append(cr)
                            cr = route_from_box(b)
                if cr is not None:
                    conservative.append(cr)
        current_routes = conservative
        current_dec = decoder.decode(current_routes)
        if not current_dec.feasible:
            raise RuntimeError(f"即使保守单点拆分方案仍无法调度：{current_dec.reason}")

    base_metrics = compute_metrics(current_dec, boxes)
    current_metrics = base_metrics
    current_score = normalized_score(current_metrics, base_metrics)
    best_routes = [r.clone() for r in current_routes]
    best_dec = current_dec
    best_metrics = current_metrics
    best_score = current_score
    history = [(0, best_score)]

    destroy_names = ["random", "related", "worst"]
    weights = {k: 1.0 for k in destroy_names}
    temp0 = 0.08

    for it in range(1, iterations + 1):
        # 自适应加权抽样 destroy 算子
        total_w = sum(weights.values())
        x = rng.random() * total_w
        acc = 0.0
        op = destroy_names[-1]
        for name in destroy_names:
            acc += weights[name]
            if x <= acc:
                op = name; break
        q = rng.randint(ALNS_REMOVE_MIN, ALNS_REMOVE_MAX)
        if op == "random":
            partial, rem = destroy_random(current_routes, q, rng)
        elif op == "related":
            partial, rem = destroy_related(current_routes, q, boxes, nodes, rng)
        else:
            partial, rem = destroy_worst(current_routes, q, boxes, current_dec, rng)

        cand_routes = regret2_repair(partial, rem, boxes, evaluator)
        if cand_routes is None:
            weights[op] *= 0.995
            continue
        try:
            ensure_route_integrity(cand_routes, boxes)
        except Exception:
            weights[op] *= 0.995
            continue

        cand_dec = decoder.decode(cand_routes)
        if not cand_dec.feasible:
            weights[op] *= 0.997
            continue
        cand_metrics = compute_metrics(cand_dec, boxes)
        cand_score = normalized_score(cand_metrics, base_metrics)
        if not math.isfinite(cand_score):
            continue

        temp = max(0.002, temp0 * (1.0 - it / max(iterations, 1)))
        delta = cand_score - current_score
        accept = delta <= 0 or rng.random() < math.exp(-delta / temp)
        if accept:
            current_routes, current_dec = cand_routes, cand_dec
            current_metrics, current_score = cand_metrics, cand_score
            weights[op] *= 1.002
        else:
            weights[op] *= 0.999

        if cand_score < best_score - 1e-10:
            best_routes = [r.clone() for r in cand_routes]
            best_dec = cand_dec
            best_metrics = cand_metrics
            best_score = cand_score
            weights[op] += 0.15

        if it % 10 == 0 or it == iterations:
            history.append((it, best_score))

    return best_routes, best_dec, best_metrics, history


# ============================================================
# 8. 输出：CSV、模板 Excel、详细 Excel、图表
# ============================================================

def write_csvs(decoded: DecodeResult, boxes: Dict[str, Box], output_dir: Path) -> Tuple[Path, Path]:
    trip_path = output_dir / "Q2_运输架次.csv"
    box_path = output_dir / "Q2_逐箱交付.csv"

    with trip_path.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["架次编号", "无人机编号", "机型编号", "电池编号", "开始时刻（s）", "访问服务区顺序", "返回O01时刻（s）", "架次能耗（kWh）"])
        for t in decoded.trips:
            w.writerow([
                t.trip_id, t.drone_id, t.drone_type, t.battery_id,
                round(t.start_time, 3), "→".join(t.service_sequence),
                round(t.return_time, 3), round(t.energy_kwh, 6)
            ])

    trip_by_box = {}
    for t in decoded.trips:
        for bid in t.box_delivery:
            trip_by_box[bid] = t.trip_id
    with box_path.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["货箱编号", "架次编号", "服务区编号", "交付完成时刻（s）"])
        for bid in sorted(boxes):
            w.writerow([bid, trip_by_box[bid], boxes[bid].service_id, round(decoded.box_delivery[bid], 3)])
    return trip_path, box_path


def fill_template(template_path: Path, decoded: DecodeResult, boxes: Dict[str, Box], output_path: Path) -> None:
    shutil.copy2(template_path, output_path)
    wb = load_workbook(output_path)
    ws_trip = wb["Q2_运输架次"]
    ws_box = wb["Q2_逐箱交付"]

    # 清理旧数据（保留第一行表头）
    for ws, max_col in [(ws_trip, 8), (ws_box, 4)]:
        for row in ws.iter_rows(min_row=2, max_col=max_col):
            for c in row:
                c.value = None

    for r, t in enumerate(decoded.trips, start=2):
        vals = [
            t.trip_id, t.drone_id, t.drone_type, t.battery_id,
            round(t.start_time, 3), "→".join(t.service_sequence),
            round(t.return_time, 3), round(t.energy_kwh, 6)
        ]
        for c, v in enumerate(vals, start=1):
            ws_trip.cell(r, c, v)

    trip_by_box = {}
    for t in decoded.trips:
        for bid in t.box_delivery:
            trip_by_box[bid] = t.trip_id
    for r, bid in enumerate(sorted(boxes), start=2):
        vals = [bid, trip_by_box[bid], boxes[bid].service_id, round(decoded.box_delivery[bid], 3)]
        for c, v in enumerate(vals, start=1):
            ws_box.cell(r, c, v)

    wb.save(output_path)


def write_detailed_xlsx(decoded: DecodeResult, boxes: Dict[str, Box], output_path: Path, metrics: Metrics) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "架次详情"
    ws.append(["架次编号", "无人机", "机型", "电池", "开始", "起飞", "返回", "路线", "能耗kWh", "返航SOC", "充电完成"])
    for t in decoded.trips:
        ws.append([t.trip_id, t.drone_id, t.drone_type, t.battery_id, t.start_time, t.takeoff_time, t.return_time,
                   "→".join(t.service_sequence), t.energy_kwh, t.end_soc, t.charge_end_time])

    ws2 = wb.create_sheet("逐箱详情")
    ws2.append(["货箱编号", "服务区", "物资类型", "交付时刻", "期望时刻", "硬截止", "优先系数", "迟到量", "迟到加权"])
    for bid in sorted(boxes):
        b = boxes[bid]; t = decoded.box_delivery[bid]
        late = max(0.0, t - b.expected_time)
        ws2.append([bid, b.service_id, b.material_type, t, b.expected_time, b.hard_deadline, b.priority, late,
                    b.priority * late / max(b.expected_time, 1.0)])

    ws3 = wb.create_sheet("航段详情")
    ws3.append(["架次", "from", "to", "载荷kg", "体积m3", "距离m", "爬升m", "下降m", "飞行时间s", "能耗kWh"])
    for t in decoded.trips:
        for le in t.route_eval.legs:
            ws3.append([t.trip_id, le.from_node, le.to_node, le.payload_mass, le.payload_volume,
                        le.distance_m, le.climb_m, le.descent_m, le.flight_time, le.energy_kwh])

    ws4 = wb.create_sheet("指标汇总")
    ws4.append(["指标", "数值"])
    ws4.append(["加权迟到", metrics.weighted_lateness])
    ws4.append(["全部任务完成时间(s)", metrics.makespan])
    ws4.append(["总运输能耗(kWh)", metrics.energy])
    ws4.append(["总架次数", metrics.trips])
    ws4.append(["硬时限违反数", metrics.hard_violations])
    wb.save(output_path)


def plot_routes(decoded: DecodeResult, nodes: Dict[str, Node], out: Path) -> None:
    fig, ax = plt.subplots(figsize=(10, 8))
    for t in decoded.trips:
        seq = ["O01"] + t.service_sequence + ["O01"]
        xs = [nodes[n].lon for n in seq]
        ys = [nodes[n].lat for n in seq]
        ax.plot(xs, ys, alpha=0.55, linewidth=1.2)
    for nid, n in nodes.items():
        ax.scatter(n.lon, n.lat, s=45 if nid == "O01" else 22)
        ax.text(n.lon, n.lat, nid, fontsize=8, ha="left", va="bottom")
    ax.set_title("Q2 Transport Routes")
    ax.set_xlabel("Longitude")
    ax.set_ylabel("Latitude")
    ax.grid(alpha=0.25)
    fig.tight_layout(); fig.savefig(out, dpi=180); plt.close(fig)


def plot_drone_gantt(decoded: DecodeResult, out: Path) -> None:
    ids = sorted(set(t.drone_id for t in decoded.trips))
    ymap = {d: i for i, d in enumerate(ids)}
    fig, ax = plt.subplots(figsize=(11, 5.5))
    for t in decoded.trips:
        ax.barh(ymap[t.drone_id], t.return_time - t.start_time, left=t.start_time, height=0.55)
        ax.text((t.start_time + t.return_time) / 2, ymap[t.drone_id], t.trip_id, ha="center", va="center", fontsize=7)
    ax.set_yticks(list(ymap.values()), list(ymap.keys()))
    ax.set_xlabel("Time (s)")
    ax.set_title("Q2 Drone Schedule Gantt")
    ax.grid(axis="x", alpha=0.25)
    fig.tight_layout(); fig.savefig(out, dpi=180); plt.close(fig)


def plot_delivery(decoded: DecodeResult, boxes: Dict[str, Box], out: Path) -> None:
    ordered = sorted(boxes.values(), key=lambda b: (b.expected_time, -b.priority, b.box_id))
    x = np.arange(len(ordered))
    actual = [decoded.box_delivery[b.box_id] for b in ordered]
    expected = [b.expected_time for b in ordered]
    hard = [b.hard_deadline if b.hard_deadline is not None else np.nan for b in ordered]
    fig, ax = plt.subplots(figsize=(12, 5.5))
    ax.plot(x, actual, marker=".", linewidth=1.0, label="actual")
    ax.plot(x, expected, linewidth=1.0, label="expected")
    ax.scatter(x, hard, marker="x", label="hard deadline")
    ax.set_xlabel("Boxes sorted by expected time")
    ax.set_ylabel("Time (s)")
    ax.set_title("Q2 Box Delivery Timeliness")
    ax.grid(alpha=0.25); ax.legend()
    fig.tight_layout(); fig.savefig(out, dpi=180); plt.close(fig)


def plot_battery(decoded: DecodeResult, out: Path) -> None:
    ids = sorted(set(t.battery_id for t in decoded.trips))
    ymap = {b: i for i, b in enumerate(ids)}
    fig, ax = plt.subplots(figsize=(11, max(5, 0.35 * len(ids) + 2)))
    for t in decoded.trips:
        y = ymap[t.battery_id]
        # 任务占用
        ax.barh(y, t.return_time - t.start_time, left=t.start_time, height=0.55)
        # 充电期用轮廓线表示
        ax.barh(y, t.charge_end_time - t.return_time, left=t.return_time, height=0.30, fill=False, edgecolor="black", linewidth=0.8)
    ax.set_yticks(list(ymap.values()), list(ymap.keys()))
    ax.set_xlabel("Time (s)")
    ax.set_title("Q2 Battery Use and Recharge")
    ax.grid(axis="x", alpha=0.25)
    fig.tight_layout(); fig.savefig(out, dpi=180); plt.close(fig)


def plot_alns(history: List[Tuple[int, float]], out: Path) -> None:
    if not history:
        return
    fig, ax = plt.subplots(figsize=(8, 4.5))
    ax.plot([x for x, _ in history], [y for _, y in history])
    ax.set_xlabel("ALNS iteration")
    ax.set_ylabel("Best normalized objective")
    ax.set_title("Q2 ALNS Convergence")
    ax.grid(alpha=0.25)
    fig.tight_layout(); fig.savefig(out, dpi=180); plt.close(fig)


# ============================================================
# 9. 主程序
# ============================================================

def run(project_root: Path, output_dir: Path, alns_iters: int, seed: int) -> None:
    t0 = time.time()
    output_dir.mkdir(parents=True, exist_ok=True)

    # 自动定位正式数据；兼容当前聊天附件的 (3)/(4) 文件名
    node_file = resolve_input_file(project_root, "调度中心与服务区", ".xlsx")
    box_file = resolve_input_file(project_root, "物资需求与配送时限", ".xlsx")
    drone_file = resolve_input_file(project_root, "运输无人机数据", ".xlsx")
    dem_file = resolve_input_file(project_root, "镇龙乡及周边30米DEM", ".mat")
    template_file = resolve_input_file(project_root, "结果提交模板", ".xlsx")

    print("[1/8] 读取数据...")
    print("  节点:", node_file)
    print("  货箱:", box_file)
    print("  无人机:", drone_file)
    print("  DEM:", dem_file)
    print("  模板:", template_file)
    nodes = load_nodes(node_file)
    boxes = load_boxes(box_file)
    drone_types, drone_units, batteries = load_drone_data(drone_file)

    print("[2/8] 构建 Arc Library（120 个无向地形航段 -> 240 个有向航段）...")
    arcs = ArcLibrary(nodes, drone_types, dem_file)
    arcs.export_csv(output_dir / "arc_library.csv")

    print("[3/8] 初始化 Route Evaluator...")
    evaluator = RouteEvaluator(boxes, drone_types, arcs)

    print("[4/8] 构造初始完整路线解...")
    initial_routes = construct_initial_routes(boxes, evaluator)
    print(f"  初始路线数: {len(initial_routes)}")

    decoder = ResourceDecoder(evaluator, drone_units, batteries, boxes)
    initial_dec = decoder.decode(initial_routes)
    print("  初始 Decoder:", initial_dec.feasible, initial_dec.reason)

    print(f"[5/8] ALNS 优化（{alns_iters} 次迭代）...")
    best_routes, best_dec, best_metrics, history = alns_optimize(
        initial_routes, boxes, nodes, evaluator, decoder, alns_iters, seed
    )
    print(f"  最终路线/架次: {len(best_routes)}")
    print(f"  加权迟到: {best_metrics.weighted_lateness:.6f}")
    print(f"  完成时间: {best_metrics.makespan:.2f} s")
    print(f"  总能耗: {best_metrics.energy:.6f} kWh")

    print("[6/8] 最终方案审计...")
    validate_final_solution(best_routes, best_dec, boxes, drone_types)
    print("  VALID = True")

    print("[7/8] 输出 Q2 结果文件...")
    write_csvs(best_dec, boxes, output_dir)
    fill_template(template_file, best_dec, boxes, output_dir / "Q2_结果.xlsx")
    write_detailed_xlsx(best_dec, boxes, output_dir / "Q2_详细结果.xlsx", best_metrics)

    print("[8/8] 生成论文图表...")
    plot_routes(best_dec, nodes, output_dir / "fig1_routes.png")
    plot_drone_gantt(best_dec, output_dir / "fig2_gantt.png")
    plot_delivery(best_dec, boxes, output_dir / "fig3_delivery.png")
    plot_battery(best_dec, output_dir / "fig4_battery.png")
    plot_alns(history, output_dir / "fig5_alns.png")

    # 简单文本摘要，便于直接复制结果数据到论文草稿
    with (output_dir / "Q2_summary.txt").open("w", encoding="utf-8") as f:
        f.write("问题二求解结果摘要\n")
        f.write("=" * 40 + "\n")
        f.write(f"总架次数: {best_metrics.trips}\n")
        f.write(f"全部任务完成时间: {best_metrics.makespan:.3f} s\n")
        f.write(f"总运输能耗: {best_metrics.energy:.6f} kWh\n")
        f.write(f"加权迟到指标: {best_metrics.weighted_lateness:.6f}\n")
        f.write(f"硬时限违反数: {best_metrics.hard_violations}\n")
        f.write(f"ALNS迭代次数: {alns_iters}\n")
        f.write(f"随机种子: {seed}\n")

    print(f"完成，用时 {time.time() - t0:.2f} s")
    print("输出目录：", output_dir)


def parse_args() -> argparse.Namespace:
    p = argparse.ArgumentParser(description="Q2 异构无人机多点多架次运输调度")
    p.add_argument("--root", type=str, default=str(PROJECT_ROOT), help="HUAWEI-CUP 项目根目录")
    p.add_argument("--output", type=str, default=str(DEFAULT_OUTPUT_DIR), help="输出目录")
    p.add_argument("--iters", type=int, default=DEFAULT_ALNS_ITERS, help="ALNS 迭代次数；先调试可设 0/50，正式可 800~3000")
    p.add_argument("--seed", type=int, default=RANDOM_SEED, help="随机种子，保证复现")
    return p.parse_args()


if __name__ == "__main__":
    args = parse_args()
    run(Path(args.root).resolve(), Path(args.output).resolve(), max(0, args.iters), args.seed)
