# -*- coding: utf-8 -*-
"""
问题三：通信约束下的运输与中继联合调度方案
================================================

设计原则
--------
1. 复用问题二的 RouteEvaluator / ResourceDecoder / ALNS，得到运输骨架；
2. 将每个运输架次展开为“爬升-巡航-下降-投送”的连续三维轨迹；
3. 按题目附录 3 计算地形遮挡、自由空间传播损耗、双向链路预算；
4. G01 直连不可用的连续时间窗转化为中继需求；
5. 为中继需求搜索 DEM 内、离地不超过 300 m 的悬停点，并合并可共同保障的需求；
6. 调度 2 架中继无人机与 6 组能源组件，校验返航 SOC、充电周转与连续通信；
7. 自动填写结果提交模板中的 Q3_中继架次、Q3_通信保障，并同步写入本次 Q2 运输结果；
8. 输出通信覆盖图、通信保障时序图、中继甘特图和链路裕量图。

推荐目录：
    HUAWEI-CUP/
      code/2/q2_solver.py        # 问题二最终程序（本程序会自动查找）
      code/3/q3_solver.py        # 本程序
      code/3/output/             # 自动创建
      数据/无人机应急物资运输基础数据/*.xlsx
      数据/镇龙乡地理空间数据/.../镇龙乡及周边30米DEM.mat
      结果提交模板.xlsx

若目录不同，可运行：
    python q3_solver.py --root "D:/.../HUAWEI-CUP"

依赖：numpy, scipy, openpyxl, matplotlib
建议 Python 3.10+
"""

from __future__ import annotations

import argparse
import copy
import csv
import importlib.util
import math
import random
import shutil
import sys
import time
from collections import defaultdict
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence, Set, Tuple

import matplotlib.pyplot as plt
import numpy as np
from openpyxl import Workbook, load_workbook
from scipy.io import loadmat


# ============================================================
# 0. 全局参数
# ============================================================

SCRIPT_PATH = Path(__file__).resolve()
DEFAULT_PROJECT_ROOT = SCRIPT_PATH.parents[2] if len(SCRIPT_PATH.parents) >= 3 else Path.cwd()
DEFAULT_OUTPUT_DIR = SCRIPT_PATH.parent / "output"

RANDOM_SEED = 20260924
DEFAULT_Q2_ITERS = 800
DEFAULT_Q2_CANDIDATES = 1

# 连续通信数值判定：最大运输巡航速度 15 m/s，2 s 对应 30 m，和 DEM 分辨率一致。
COMM_SAMPLE_SEC = 2.0
# 中继候选覆盖验证可稍疏，但最终审计仍使用 COMM_SAMPLE_SEC。
RELAY_SAMPLE_SEC = 5.0
TRANSITION_TOL_SEC = 0.10

# 中继候选点搜索参数
RELAY_AGL_LEVELS = (120.0, 180.0, 240.0, 300.0)
RELAY_LINE_FRACTIONS = (0.30, 0.45, 0.60, 0.75, 0.90, 1.00)
RELAY_OPTION_KEEP = 12
MIN_SPLIT_WINDOW_SEC = 12.0

# 合并中继任务时的启发式代价：先压缩中继架次，再兼顾能耗和长时间空等。
RELAY_SORTIE_FIXED_COST = 1.00
RELAY_ENERGY_COST_W = 0.30
RELAY_IDLE_HOUR_COST_W = 0.10

# 多个问题二候选解之间的联合目标权重（归一化后比较）。
JOINT_WEIGHTS = {
    "lateness": 0.40,
    "makespan": 0.25,
    "energy": 0.20,
    "trips": 0.15,
}

G = 9.81
EARTH_RADIUS_M = 6_371_008.8


# ============================================================
# 1. Q2 模块定位与载入
# ============================================================


def find_q2_solver(root: Path) -> Path:
    """自动寻找问题二最终程序，优先 q2_solver*.py，其次扫描包含关键标记的 .py。"""
    candidates: List[Path] = []
    for pattern in ("q2_solver.py", "q2_solver*.py", "*q2*.py"):
        candidates.extend(root.rglob(pattern))
    uniq = []
    seen = set()
    for p in candidates:
        rp = p.resolve()
        if rp == SCRIPT_PATH or "output" in {x.lower() for x in p.parts}:
            continue
        if rp not in seen:
            seen.add(rp); uniq.append(p)
    if uniq:
        uniq.sort(key=lambda p: (0 if p.name == "q2_solver.py" else 1, len(str(p))))
        return uniq[0]

    marker = "问题二：异构无人机多点多架次运输调度"
    for p in root.rglob("*.py"):
        if p.resolve() == SCRIPT_PATH:
            continue
        try:
            head = p.read_text(encoding="utf-8", errors="ignore")[:12000]
        except Exception:
            continue
        if marker in head and "ResourceDecoder" in head and "RouteEvaluator" in head:
            return p
    raise FileNotFoundError(
        "未找到问题二最终程序。请把 q2_solver.py 放在项目 code/2 目录，"
        "或使用 --q2-solver 显式指定。"
    )


def import_q2_module(path: Path):
    name = "q2_solver_for_q3"
    spec = importlib.util.spec_from_file_location(name, str(path))
    if spec is None or spec.loader is None:
        raise ImportError(f"无法载入 Q2 模块：{path}")
    mod = importlib.util.module_from_spec(spec)
    sys.modules[name] = mod
    spec.loader.exec_module(mod)
    return mod


# ============================================================
# 2. Q3 数据对象
# ============================================================


@dataclass(frozen=True)
class RelaySpec:
    type_id: str
    name: str
    empty_mass: float
    comm_module_mass: float
    takeoff_mass: float
    cruise_speed: float
    cruise_power: float
    energy_kwh: float
    reserve_ratio: float
    prep_time: float
    link_time: float
    turnaround_time: float
    climb_speed: float
    descent_speed: float
    climb_efficiency: float
    descent_efficiency: float
    hover_power: float
    comm_power: float
    max_agl: float


@dataclass(frozen=True)
class RelayUnit:
    relay_id: str
    type_id: str


@dataclass(frozen=True)
class RelayEnergySpec:
    type_id: str
    count: int
    full_charge_time: float


@dataclass(frozen=True)
class CommParams:
    frequency_mhz: float
    system_loss_db: float
    obstruction_loss_db: float
    sensitivity_dbm: float
    fade_margin_db: float
    gateway_height_m: float
    tx_power_dbm: Dict[str, float]
    antenna_gain_dbi: Dict[str, float]


@dataclass(frozen=True)
class Position3D:
    lon: float
    lat: float
    alt: float


@dataclass
class TrajectorySegment:
    trip_id: str
    label: str
    phase: str
    start: float
    end: float
    p0: Position3D
    p1: Position3D

    def position(self, t: float) -> Position3D:
        if self.end <= self.start + 1e-12:
            return self.p0
        u = min(max((t - self.start) / (self.end - self.start), 0.0), 1.0)
        return Position3D(
            self.p0.lon + (self.p1.lon - self.p0.lon) * u,
            self.p0.lat + (self.p1.lat - self.p0.lat) * u,
            self.p0.alt + (self.p1.alt - self.p0.alt) * u,
        )


@dataclass
class CommStage:
    trip_id: str
    label: str
    start: float
    end: float
    mode: str                    # "直连" / "中继"
    relay_sortie_id: str = ""
    min_margin_db: float = math.inf


@dataclass
class RelayCandidate:
    key: Tuple[int, int, int]
    lon: float
    lat: float
    ground_alt: float
    hover_alt: float
    agl: float
    backhaul_margin_db: float
    out_time: float
    back_time: float
    out_energy: float
    back_energy: float
    cruise_alt: float


@dataclass
class RelayDemand:
    demand_id: str
    trip_id: str
    label: str
    start: float
    end: float
    segment: TrajectorySegment
    options: Dict[Tuple[int, int, int], RelayCandidate] = field(default_factory=dict)
    sortie_id: str = ""

    @property
    def duration(self) -> float:
        return max(0.0, self.end - self.start)


@dataclass
class RelayGroup:
    demands: List[RelayDemand]
    candidate_keys: Set[Tuple[int, int, int]]
    candidate: RelayCandidate
    service_start: float
    service_end: float
    energy_kwh: float
    mission_start_latest: float
    return_time: float


@dataclass
class RelaySortie:
    sortie_id: str
    relay_id: str
    energy_id: str
    start_time: float
    lon: float
    lat: float
    hover_alt: float
    agl: float
    link_complete_time: float
    service_end_time: float
    return_time: float
    energy_kwh: float
    end_soc: float
    charge_end_time: float
    turnaround_end_time: float
    candidate: RelayCandidate
    demand_ids: List[str]


@dataclass
class Q3Metrics:
    weighted_lateness: float
    joint_makespan: float
    transport_energy: float
    relay_energy: float
    total_energy: float
    transport_trips: int
    relay_trips: int
    total_trips: int
    hard_violations: int
    min_comm_margin_db: float


@dataclass
class Q3Solution:
    routes: List[Any]
    decoded: Any
    q2_metrics: Any
    segments: Dict[str, List[TrajectorySegment]]
    stages: List[CommStage]
    demands: List[RelayDemand]
    relay_sorties: List[RelaySortie]
    metrics: Q3Metrics
    seed: int


# ============================================================
# 3. 文件读取
# ============================================================


def resolve_input_file(root: Path, base_name: str, suffix: str) -> Path:
    exact = list(root.rglob(base_name + suffix))
    if exact:
        return exact[0]
    candidates = [
        p for p in root.rglob(f"{base_name}*{suffix}")
        if "output" not in {x.lower() for x in p.parts}
    ]
    if not candidates:
        raise FileNotFoundError(f"未找到 {base_name}{suffix}，搜索根目录：{root}")
    candidates.sort(key=lambda p: (len(p.name), len(str(p))))
    return candidates[0]


def load_relay_data(path: Path) -> Tuple[RelaySpec, List[RelayUnit], RelayEnergySpec]:
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    v = [ws.cell(3, c).value for c in range(1, 20)]
    spec = RelaySpec(
        type_id=str(v[0]).strip(), name=str(v[1]), empty_mass=float(v[2]),
        comm_module_mass=float(v[3]), takeoff_mass=float(v[4]), cruise_speed=float(v[5]),
        cruise_power=float(v[6]), energy_kwh=float(v[7]), reserve_ratio=float(v[8]) / 100.0,
        prep_time=float(v[9]), link_time=float(v[10]), turnaround_time=float(v[11]),
        climb_speed=float(v[12]), descent_speed=float(v[13]), climb_efficiency=float(v[14]),
        descent_efficiency=float(v[15]), hover_power=float(v[16]), comm_power=float(v[17]),
        max_agl=float(v[18]),
    )
    units = []
    for r in range(7, 20):
        rid = ws.cell(r, 1).value
        if rid and str(rid).startswith("R"):
            units.append(RelayUnit(str(rid).strip(), str(ws.cell(r, 2).value).strip()))
    erow = None
    for r in range(1, ws.max_row + 1):
        if str(ws.cell(r, 1).value).strip() == spec.type_id and r > 8:
            c2, c3 = ws.cell(r, 2).value, ws.cell(r, 3).value
            if isinstance(c2, (int, float)) and isinstance(c3, (int, float)):
                erow = r
    if erow is None:
        raise ValueError("中继能源组件库存读取失败")
    energy = RelayEnergySpec(spec.type_id, int(ws.cell(erow, 2).value), float(ws.cell(erow, 3).value))
    wb.close()
    return spec, units, energy


def load_comm_params(path: Path) -> CommParams:
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = []
    for r in range(3, ws.max_row + 1):
        cat = ws.cell(r, 1).value
        name = ws.cell(r, 2).value
        val = ws.cell(r, 5).value
        if cat is not None and name is not None and val is not None:
            rows.append((str(cat).strip(), str(name).strip(), float(val)))
    wb.close()
    mp = {(c, n): v for c, n, v in rows}
    tx = {
        "T": mp[("运输无人机", "发射功率（dBm）")],
        "RA": mp[("中继接入端", "发射功率（dBm）")],
        "RB": mp[("中继回传端", "发射功率（dBm）")],
        "G": mp[("固定网关 G01", "发射功率（dBm）")],
    }
    gain = {
        "T": mp[("运输无人机", "天线增益（dBi）")],
        "RA": mp[("中继接入端", "天线增益（dBi）")],
        "RB": mp[("中继回传端", "天线增益（dBi）")],
        "G": mp[("固定网关 G01", "天线增益（dBi）")],
    }
    return CommParams(
        frequency_mhz=mp[("传播参数", "载波频率（MHz）")],
        system_loss_db=mp[("传播参数", "系统损耗（dB）")],
        obstruction_loss_db=mp[("传播参数", "地形遮挡附加损耗（dB）")],
        sensitivity_dbm=mp[("接收参数", "接收灵敏度（dBm）")],
        fade_margin_db=mp[("接收参数", "衰落裕量（dB）")],
        gateway_height_m=mp[("固定网关 G01", "天线离地高度（m）")],
        tx_power_dbm=tx,
        antenna_gain_dbi=gain,
    )


# ============================================================
# 4. DEM 与通信链路模型
# ============================================================


def haversine_m(lon1: float, lat1: float, lon2: float, lat2: float) -> float:
    p1, p2 = math.radians(lat1), math.radians(lat2)
    dphi = math.radians(lat2 - lat1)
    dlambda = math.radians(lon2 - lon1)
    a = math.sin(dphi / 2) ** 2 + math.cos(p1) * math.cos(p2) * math.sin(dlambda / 2) ** 2
    return 2 * EARTH_RADIUS_M * math.asin(min(1.0, math.sqrt(a)))


class DEMHelper:
    def __init__(self, dem_path: Path):
        mat = loadmat(dem_path)
        self.dem = np.asarray(mat["dem"], dtype=float)
        self.lat = np.asarray(mat["latitude"], dtype=float).reshape(-1)
        self.lon = np.asarray(mat["longitude"], dtype=float).reshape(-1)
        self.nodata = float(np.asarray(mat.get("nodata", [[-32767.0]])).reshape(-1)[0])
        self.lat0 = float(self.lat[0]); self.lon0 = float(self.lon[0])
        self.dlat = float(self.lat[1] - self.lat[0])
        self.dlon = float(self.lon[1] - self.lon[0])
        self.lat_min, self.lat_max = float(self.lat.min()), float(self.lat.max())
        self.lon_min, self.lon_max = float(self.lon.min()), float(self.lon.max())

    def in_bounds(self, lon: float, lat: float) -> bool:
        return self.lon_min <= lon <= self.lon_max and self.lat_min <= lat <= self.lat_max

    def _rc_vec(self, lon_arr: np.ndarray, lat_arr: np.ndarray) -> Tuple[np.ndarray, np.ndarray]:
        c = np.rint((lon_arr - self.lon0) / self.dlon).astype(int)
        r = np.rint((lat_arr - self.lat0) / self.dlat).astype(int)
        r = np.clip(r, 0, self.dem.shape[0] - 1)
        c = np.clip(c, 0, self.dem.shape[1] - 1)
        return r, c

    def ground(self, lon: float, lat: float) -> float:
        if not self.in_bounds(lon, lat):
            return math.nan
        r, c = self._rc_vec(np.array([lon]), np.array([lat]))
        v = float(self.dem[int(r[0]), int(c[0])])
        return v if np.isfinite(v) and abs(v - self.nodata) > 1e-6 else math.nan

    def line_dem(self, p1: Position3D, p2: Position3D, spacing_m: float = 30.0) -> Tuple[np.ndarray, np.ndarray]:
        horiz = haversine_m(p1.lon, p1.lat, p2.lon, p2.lat)
        if horiz <= 1e-9:
            return np.array([], dtype=float), np.array([], dtype=float)
        n = max(2, int(math.ceil(horiz / spacing_m)))
        u = np.arange(1, n, dtype=float) / n
        lon = p1.lon + (p2.lon - p1.lon) * u
        lat = p1.lat + (p2.lat - p1.lat) * u
        z = p1.alt + (p2.alt - p1.alt) * u
        r, c = self._rc_vec(lon, lat)
        ter = self.dem[r, c]
        valid = np.isfinite(ter) & (np.abs(ter - self.nodata) > 1e-6)
        return ter[valid].astype(float), z[valid].astype(float)

    def obstructed(self, p1: Position3D, p2: Position3D) -> bool:
        ter, los = self.line_dem(p1, p2)
        if ter.size == 0:
            return False
        return bool(np.any(ter >= los - 1e-9))

    def max_ground_along(self, lon1: float, lat1: float, lon2: float, lat2: float) -> float:
        p1 = Position3D(lon1, lat1, 0.0); p2 = Position3D(lon2, lat2, 0.0)
        horiz = haversine_m(lon1, lat1, lon2, lat2)
        n = max(2, int(math.ceil(horiz / 30.0)))
        u = np.linspace(0.0, 1.0, n + 1)
        lon = lon1 + (lon2 - lon1) * u
        lat = lat1 + (lat2 - lat1) * u
        r, c = self._rc_vec(lon, lat)
        ter = self.dem[r, c]
        valid = np.isfinite(ter) & (np.abs(ter - self.nodata) > 1e-6)
        if not np.any(valid):
            raise ValueError("中继航线没有有效 DEM 像元")
        return float(np.max(ter[valid]))


class CommunicationModel:
    def __init__(self, params: CommParams, dem: DEMHelper, gateway: Position3D):
        self.params = params
        self.dem = dem
        self.gateway = gateway
        self._cache: Dict[Tuple, Tuple[bool, float, bool, float]] = {}

    def threshold(self, a: str, b: str) -> float:
        p = self.params
        rx_eff = p.sensitivity_dbm + p.fade_margin_db
        ab = p.tx_power_dbm[a] + p.antenna_gain_dbi[a] + p.antenna_gain_dbi[b] - p.system_loss_db - rx_eff
        ba = p.tx_power_dbm[b] + p.antenna_gain_dbi[b] + p.antenna_gain_dbi[a] - p.system_loss_db - rx_eff
        return min(ab, ba)

    @staticmethod
    def _pkey(p: Position3D) -> Tuple[int, int, int]:
        return (round(p.lon * 1e6), round(p.lat * 1e6), round(p.alt * 10))

    def link(self, p1: Position3D, p2: Position3D, a: str, b: str) -> Tuple[bool, float, bool, float]:
        key = (self._pkey(p1), self._pkey(p2), a, b)
        if key in self._cache:
            return self._cache[key]
        horiz = haversine_m(p1.lon, p1.lat, p2.lon, p2.lat)
        dist_km = math.sqrt(horiz ** 2 + (p1.alt - p2.alt) ** 2) / 1000.0
        fspl = 32.44 + 20.0 * math.log10(self.params.frequency_mhz) + 20.0 * math.log10(max(dist_km, 1e-9))
        thr = self.threshold(a, b)
        # 即使无遮挡损耗已经超过门限，也无需再做 DEM 遮挡扫描。
        if fspl > thr + 1e-12:
            out = (False, fspl, False, thr - fspl)
            self._cache[key] = out
            return out
        obs = self.dem.obstructed(p1, p2)
        total = fspl + (self.params.obstruction_loss_db if obs else 0.0)
        margin = thr - total
        out = (margin >= -1e-9, total, obs, margin)
        self._cache[key] = out
        return out


# ============================================================
# 5. 运输轨迹展开与直连区间识别
# ============================================================


def build_transport_segments(q2, decoded, nodes, drone_types, arcs) -> Dict[str, List[TrajectorySegment]]:
    all_segments: Dict[str, List[TrajectorySegment]] = {}
    for trip in decoded.trips:
        ev = trip.route_eval
        dt = drone_types[trip.drone_type]
        t = trip.start_time + ev.pre_operation_time
        stop_map = {s.service_id: s for s in ev.stops}
        segs: List[TrajectorySegment] = []
        for leg in ev.legs:
            arc = arcs[(leg.from_node, leg.to_node)]
            na, nb = nodes[leg.from_node], nodes[leg.to_node]
            tc = arc.climb_m / dt.climb_speed
            th = arc.distance_m / dt.cruise_speed
            td = arc.descent_m / dt.descent_speed
            if tc > 1e-9:
                segs.append(TrajectorySegment(
                    trip.trip_id, f"爬升 {leg.from_node}→{leg.to_node}", "爬升", t, t + tc,
                    Position3D(na.lon, na.lat, na.work_alt),
                    Position3D(na.lon, na.lat, arc.cruise_alt_m),
                )); t += tc
            if th > 1e-9:
                segs.append(TrajectorySegment(
                    trip.trip_id, f"巡航 {leg.from_node}→{leg.to_node}", "巡航", t, t + th,
                    Position3D(na.lon, na.lat, arc.cruise_alt_m),
                    Position3D(nb.lon, nb.lat, arc.cruise_alt_m),
                )); t += th
            if td > 1e-9:
                segs.append(TrajectorySegment(
                    trip.trip_id, f"下降 {leg.from_node}→{leg.to_node}", "下降", t, t + td,
                    Position3D(nb.lon, nb.lat, arc.cruise_alt_m),
                    Position3D(nb.lon, nb.lat, nb.work_alt),
                )); t += td
            if leg.to_node != "O01":
                se = stop_map[leg.to_node]
                segs.append(TrajectorySegment(
                    trip.trip_id, f"投送 {leg.to_node}", "投送", t, t + se.service_duration,
                    Position3D(nb.lon, nb.lat, nb.work_alt),
                    Position3D(nb.lon, nb.lat, nb.work_alt),
                )); t += se.service_duration
        if abs(t - trip.return_time) > 1e-4:
            raise AssertionError(f"{trip.trip_id} 轨迹时间展开与 Q2 返回时刻不一致：{t} vs {trip.return_time}")
        all_segments[trip.trip_id] = segs
    return all_segments


def time_grid(a: float, b: float, step: float) -> List[float]:
    if b <= a + 1e-12:
        return [a]
    vals = list(np.arange(a, b, step, dtype=float))
    if not vals or abs(vals[0] - a) > 1e-9:
        vals.insert(0, a)
    if vals[-1] < b - 1e-9:
        vals.append(b)
    else:
        vals[-1] = b
    return vals


def find_transition(seg: TrajectorySegment, t0: float, t1: float, state0: bool, comm: CommunicationModel) -> float:
    """二分逼近直连可用性切换时刻。"""
    lo, hi = t0, t1
    for _ in range(20):
        if hi - lo <= TRANSITION_TOL_SEC:
            break
        mid = 0.5 * (lo + hi)
        state_mid = comm.link(seg.position(mid), comm.gateway, "T", "G")[0]
        if state_mid == state0:
            lo = mid
        else:
            hi = mid
    return 0.5 * (lo + hi)


def split_segment_by_direct(seg: TrajectorySegment, comm: CommunicationModel) -> List[Tuple[float, float, bool, float]]:
    ts = time_grid(seg.start, seg.end, COMM_SAMPLE_SEC)
    states = []
    margins = []
    for t in ts:
        ok, _, _, margin = comm.link(seg.position(t), comm.gateway, "T", "G")
        states.append(ok); margins.append(margin)
    pieces: List[Tuple[float, float, bool, float]] = []
    cur_state = states[0]
    cur_start = ts[0]
    cur_margins = [margins[0]]
    for i in range(1, len(ts)):
        if states[i] == cur_state:
            cur_margins.append(margins[i])
            continue
        tr = find_transition(seg, ts[i - 1], ts[i], cur_state, comm)
        pieces.append((cur_start, tr, cur_state, min(cur_margins)))
        cur_start = tr
        cur_state = states[i]
        cur_margins = [margins[i]]
    if seg.end > cur_start + 1e-9:
        pieces.append((cur_start, seg.end, cur_state, min(cur_margins)))
    elif pieces:
        pieces[-1] = (pieces[-1][0], seg.end, pieces[-1][2], pieces[-1][3])
    return pieces


def build_direct_stages_and_demands(
    segments: Dict[str, List[TrajectorySegment]], comm: CommunicationModel
) -> Tuple[List[CommStage], List[RelayDemand]]:
    stages: List[CommStage] = []
    demands: List[RelayDemand] = []
    did = 1
    for trip_id in sorted(segments):
        for seg in segments[trip_id]:
            for a, b, direct, margin in split_segment_by_direct(seg, comm):
                if b <= a + 1e-8:
                    continue
                if direct:
                    stages.append(CommStage(trip_id, seg.label, a, b, "直连", "", margin))
                else:
                    d = RelayDemand(f"D{did:04d}", trip_id, seg.label, a, b, seg)
                    did += 1
                    demands.append(d)
                    stages.append(CommStage(trip_id, seg.label, a, b, "中继", "", math.inf))
    return stages, demands


# ============================================================
# 6. 中继悬停候选与能量模型
# ============================================================


def relay_flight_candidate(
    lon: float, lat: float, agl: float, nodes, dem: DEMHelper,
    relay: RelaySpec, comm: CommunicationModel,
) -> Optional[RelayCandidate]:
    if agl > relay.max_agl + 1e-9 or not dem.in_bounds(lon, lat):
        return None
    ground = dem.ground(lon, lat)
    if not np.isfinite(ground):
        return None
    hover_alt = ground + agl
    hp = Position3D(lon, lat, hover_alt)
    back_ok, _, _, back_margin = comm.link(hp, comm.gateway, "RB", "G")
    if not back_ok:
        return None

    o = nodes["O01"]
    dist = haversine_m(o.lon, o.lat, lon, lat)
    dem_max = dem.max_ground_along(o.lon, o.lat, lon, lat)
    cruise_alt = max(dem_max + 50.0, hover_alt, o.work_alt)

    out_climb = max(0.0, cruise_alt - o.work_alt)
    out_desc = max(0.0, cruise_alt - hover_alt)
    back_climb = max(0.0, cruise_alt - hover_alt)
    back_desc = max(0.0, cruise_alt - o.work_alt)

    out_time = out_climb / relay.climb_speed + dist / relay.cruise_speed + out_desc / relay.descent_speed
    back_time = back_climb / relay.climb_speed + dist / relay.cruise_speed + back_desc / relay.descent_speed
    hor_energy = relay.cruise_power * (dist / relay.cruise_speed) / 3600.0
    out_up = relay.takeoff_mass * G * out_climb / max(relay.climb_efficiency, 1e-9) / 3.6e6
    back_up = relay.takeoff_mass * G * back_climb / max(relay.climb_efficiency, 1e-9) / 3.6e6
    out_energy = hor_energy + out_up
    back_energy = hor_energy + back_up
    key = (round(lon * 1e6), round(lat * 1e6), round(agl))
    return RelayCandidate(
        key, lon, lat, ground, hover_alt, agl, back_margin,
        out_time, back_time, out_energy, back_energy, cruise_alt,
    )


def demand_sample_points(d: RelayDemand, step: float = RELAY_SAMPLE_SEC) -> List[Position3D]:
    ts = time_grid(d.start, d.end, step)
    # 再加入区间中点，避免短区间只取到端点。
    if d.end - d.start > 1e-6:
        ts.append(0.5 * (d.start + d.end))
    ts = sorted(set(round(x, 6) for x in ts))
    return [d.segment.position(t) for t in ts]


def candidate_covers_demand(c: RelayCandidate, d: RelayDemand, comm: CommunicationModel) -> Tuple[bool, float]:
    hp = Position3D(c.lon, c.lat, c.hover_alt)
    worst = c.backhaul_margin_db
    for p in demand_sample_points(d):
        ok, _, _, margin = comm.link(p, hp, "T", "RA")
        worst = min(worst, margin)
        if not ok:
            return False, worst
    return True, worst


def group_energy(c: RelayCandidate, service_start: float, service_end: float, relay: RelaySpec) -> Tuple[float, float, float]:
    service_dur = max(0.0, service_end - service_start)
    e_service = (relay.hover_power + relay.comm_power) * (relay.link_time + service_dur) / 3600.0
    total = c.out_energy + c.back_energy + e_service
    start_latest = service_start - relay.prep_time - c.out_time - relay.link_time
    ret = service_end + c.back_time
    return total, start_latest, ret


def candidate_score_for_demand(c: RelayCandidate, d: RelayDemand, relay: RelaySpec) -> float:
    e, start, ret = group_energy(c, d.start, d.end, relay)
    if start < -1e-9 or e > (1.0 - relay.reserve_ratio) * relay.energy_kwh + 1e-9:
        return math.inf
    return e + 0.00002 * ret - 0.01 * c.backhaul_margin_db


def candidate_xy_for_demand(d: RelayDemand, gateway: Position3D) -> List[Tuple[float, float]]:
    pts = demand_sample_points(d)
    center_lon = float(np.mean([p.lon for p in pts])); center_lat = float(np.mean([p.lat for p in pts]))
    out: List[Tuple[float, float]] = []
    for f in RELAY_LINE_FRACTIONS:
        out.append((gateway.lon + f * (center_lon - gateway.lon), gateway.lat + f * (center_lat - gateway.lat)))
    # 在需求中心周围增加小幅横向扰动，帮助绕过局部山脊。
    base_offsets = [(-0.004, 0), (0.004, 0), (0, -0.004), (0, 0.004)]
    out.append((center_lon, center_lat))
    for dx, dy in base_offsets:
        out.append((center_lon + dx, center_lat + dy))
    # 去重
    uniq, seen = [], set()
    for lon, lat in out:
        k = (round(lon, 6), round(lat, 6))
        if k not in seen:
            seen.add(k); uniq.append((lon, lat))
    return uniq


def search_demand_options(
    d: RelayDemand, nodes, dem: DEMHelper, relay: RelaySpec, comm: CommunicationModel
) -> Dict[Tuple[int, int, int], RelayCandidate]:
    found: List[Tuple[float, RelayCandidate]] = []
    for lon, lat in candidate_xy_for_demand(d, comm.gateway):
        for agl in RELAY_AGL_LEVELS:
            c = relay_flight_candidate(lon, lat, agl, nodes, dem, relay, comm)
            if c is None:
                continue
            ok, _ = candidate_covers_demand(c, d, comm)
            if not ok:
                continue
            score = candidate_score_for_demand(c, d, relay)
            if math.isfinite(score):
                found.append((score, c))
    found.sort(key=lambda x: x[0])
    out: Dict[Tuple[int, int, int], RelayCandidate] = {}
    for _, c in found:
        if c.key not in out:
            out[c.key] = c
        if len(out) >= RELAY_OPTION_KEEP:
            break
    return out


def split_demand(d: RelayDemand, serial_a: str, serial_b: str) -> Tuple[RelayDemand, RelayDemand]:
    mid = 0.5 * (d.start + d.end)
    return (
        RelayDemand(serial_a, d.trip_id, d.label, d.start, mid, d.segment),
        RelayDemand(serial_b, d.trip_id, d.label, mid, d.end, d.segment),
    )


def ensure_demand_options(
    demands: List[RelayDemand], nodes, dem: DEMHelper, relay: RelaySpec, comm: CommunicationModel
) -> List[RelayDemand]:
    queue = list(demands)
    out: List[RelayDemand] = []
    serial = 1
    while queue:
        d = queue.pop(0)
        d.options = search_demand_options(d, nodes, dem, relay, comm)
        if d.options:
            out.append(d)
            continue
        if d.duration <= MIN_SPLIT_WINDOW_SEC:
            raise RuntimeError(
                f"中继候选搜索失败：{d.trip_id} {d.label} {d.start:.1f}-{d.end:.1f}s。"
                "可尝试增大候选搜索范围或改动运输方案。"
            )
        a, b = split_demand(d, f"S{serial:05d}a", f"S{serial:05d}b")
        serial += 1
        queue.insert(0, b); queue.insert(0, a)
    # 重新统一编号，便于输出。
    out.sort(key=lambda x: (x.start, x.trip_id, x.label))
    for i, d in enumerate(out, 1):
        d.demand_id = f"D{i:04d}"
    return out


# ============================================================
# 7. 中继需求合并与资源调度
# ============================================================


def choose_group_candidate(
    demands: Sequence[RelayDemand], candidate_keys: Set[Tuple[int, int, int]], relay: RelaySpec
) -> Optional[Tuple[RelayCandidate, float, float, float]]:
    if not candidate_keys:
        return None
    s0 = min(d.start for d in demands); s1 = max(d.end for d in demands)
    best = None
    ref = demands[0]
    for key in candidate_keys:
        c = ref.options.get(key)
        if c is None:
            for d in demands:
                if key in d.options:
                    c = d.options[key]; break
        if c is None:
            continue
        e, start_latest, ret = group_energy(c, s0, s1, relay)
        if start_latest < -1e-9:
            continue
        if e > (1.0 - relay.reserve_ratio) * relay.energy_kwh + 1e-9:
            continue
        span_hr = (s1 - s0) / 3600.0
        cost = RELAY_SORTIE_FIXED_COST + RELAY_ENERGY_COST_W * e + RELAY_IDLE_HOUR_COST_W * span_hr
        item = (cost, c, e, start_latest, ret)
        if best is None or item[0] < best[0]:
            best = item
    if best is None:
        return None
    return best[1], best[2], best[3], best[4]


def make_single_group(d: RelayDemand, relay: RelaySpec) -> RelayGroup:
    keys = set(d.options)
    picked = choose_group_candidate([d], keys, relay)
    if picked is None:
        raise RuntimeError(f"需求 {d.demand_id} 虽有候选点，但能源/提前量不可行")
    c, e, start_latest, ret = picked
    return RelayGroup([d], keys, c, d.start, d.end, e, start_latest, ret)


def group_cost(g: RelayGroup) -> float:
    span_hr = (g.service_end - g.service_start) / 3600.0
    return RELAY_SORTIE_FIXED_COST + RELAY_ENERGY_COST_W * g.energy_kwh + RELAY_IDLE_HOUR_COST_W * span_hr


def merge_groups(g1: RelayGroup, g2: RelayGroup, relay: RelaySpec) -> Optional[RelayGroup]:
    keys = g1.candidate_keys & g2.candidate_keys
    if not keys:
        return None
    demands = g1.demands + g2.demands
    picked = choose_group_candidate(demands, keys, relay)
    if picked is None:
        return None
    c, e, start_latest, ret = picked
    return RelayGroup(
        demands=demands,
        candidate_keys=keys,
        candidate=c,
        service_start=min(g1.service_start, g2.service_start),
        service_end=max(g1.service_end, g2.service_end),
        energy_kwh=e,
        mission_start_latest=start_latest,
        return_time=ret,
    )


def compress_relay_groups(demands: List[RelayDemand], relay: RelaySpec) -> List[RelayGroup]:
    groups = [make_single_group(d, relay) for d in demands]
    # 贪心两两合并：每次选“原两架次代价 - 合并后代价”最大的正收益。
    while True:
        best = None
        n = len(groups)
        for i in range(n):
            for j in range(i + 1, n):
                mg = merge_groups(groups[i], groups[j], relay)
                if mg is None:
                    continue
                saving = group_cost(groups[i]) + group_cost(groups[j]) - group_cost(mg)
                if saving > 1e-10 and (best is None or saving > best[0]):
                    best = (saving, i, j, mg)
        if best is None:
            break
        _, i, j, mg = best
        groups = [g for k, g in enumerate(groups) if k not in (i, j)] + [mg]
    groups.sort(key=lambda g: (g.service_start, g.service_end))
    return groups


def charge_time_seconds(soc: float, full_charge_time: float) -> float:
    soc = min(max(soc, 0.0), 1.0)
    if soc >= 1.0 - 1e-12:
        return 0.0
    if soc >= 0.9:
        return full_charge_time * 0.35 * (1.0 - soc) / 0.1
    return full_charge_time * (0.65 * (0.9 - soc) / 0.9 + 0.35)


def schedule_relay_resources(
    groups: List[RelayGroup], relay: RelaySpec, units: List[RelayUnit], energy_spec: RelayEnergySpec
) -> List[RelaySortie]:
    drone_avail = {u.relay_id: 0.0 for u in units}
    energy_avail = {f"REN-{i:02d}": 0.0 for i in range(1, energy_spec.count + 1)}
    sorties: List[RelaySortie] = []

    for idx, g in enumerate(sorted(groups, key=lambda x: (x.service_start, x.mission_start_latest)), 1):
        latest = g.mission_start_latest
        drone_choices = [(t, rid) for rid, t in drone_avail.items() if t <= latest + 1e-9]
        energy_choices = [(t, eid) for eid, t in energy_avail.items() if t <= latest + 1e-9]
        if not drone_choices:
            raise RuntimeError(
                f"中继无人机资源冲突：服务窗 {g.service_start:.1f}-{g.service_end:.1f}s "
                f"要求最迟 {latest:.1f}s 开始准备，但 2 架中继均不可用。"
            )
        if not energy_choices:
            raise RuntimeError("中继能源组件周转不足")
        # 选择最早可用资源；任务开始仍取 latest，使中继恰好在需求开始前完成建链。
        _, rid = min(drone_choices)
        _, eid = min(energy_choices)
        start = max(0.0, latest)
        ret = g.return_time
        end_soc = 1.0 - g.energy_kwh / relay.energy_kwh
        if end_soc + 1e-12 < relay.reserve_ratio:
            raise RuntimeError("中继架次返航 SOC 低于下限")
        charge_end = ret + charge_time_seconds(end_soc, energy_spec.full_charge_time)
        turnaround_end = ret + relay.turnaround_time
        sid = f"Q3R{idx:03d}"
        s = RelaySortie(
            sid, rid, eid, start,
            g.candidate.lon, g.candidate.lat, g.candidate.hover_alt, g.candidate.agl,
            g.service_start, g.service_end, ret, g.energy_kwh, end_soc,
            charge_end, turnaround_end, g.candidate,
            [d.demand_id for d in g.demands],
        )
        sorties.append(s)
        drone_avail[rid] = turnaround_end
        energy_avail[eid] = charge_end
        for d in g.demands:
            d.sortie_id = sid
    return sorties


def assign_sorties_to_stages(stages: List[CommStage], demands: List[RelayDemand]) -> None:
    # 以 trip/label/start/end 对齐，考虑 demand 可能因候选不可行而被再次切分。
    relay_stages = [s for s in stages if s.mode == "中继"]
    new_stages: List[CommStage] = [s for s in stages if s.mode == "直连"]
    by_key = defaultdict(list)
    for d in demands:
        by_key[(d.trip_id, d.label)].append(d)
    for s in relay_stages:
        ds = sorted(by_key[(s.trip_id, s.label)], key=lambda d: d.start)
        for d in ds:
            a = max(s.start, d.start); b = min(s.end, d.end)
            if b > a + 1e-8:
                new_stages.append(CommStage(s.trip_id, s.label, a, b, "中继", d.sortie_id, math.inf))
    new_stages.sort(key=lambda x: (x.trip_id, x.start, x.end))
    stages[:] = new_stages


# ============================================================
# 8. 最终通信审计与指标
# ============================================================


def audit_communication(
    stages: List[CommStage], segments: Dict[str, List[TrajectorySegment]],
    sorties: List[RelaySortie], comm: CommunicationModel
) -> float:
    sortie_map = {s.sortie_id: s for s in sorties}
    seg_map = defaultdict(list)
    for tid, segs in segments.items():
        seg_map[tid].extend(segs)

    min_margin = math.inf
    # 每个运输架次必须从起飞到返回无缝覆盖。
    by_trip = defaultdict(list)
    for s in stages:
        by_trip[s.trip_id].append(s)

    for tid, segs in segments.items():
        t0 = min(s.start for s in segs); t1 = max(s.end for s in segs)
        sts = sorted(by_trip[tid], key=lambda s: s.start)
        if not sts:
            raise AssertionError(f"{tid} 没有通信保障阶段")
        if abs(sts[0].start - t0) > 0.2 or abs(sts[-1].end - t1) > 0.2:
            raise AssertionError(f"{tid} 通信阶段未覆盖完整飞行时间")
        for a, b in zip(sts, sts[1:]):
            if abs(a.end - b.start) > 0.2:
                raise AssertionError(f"{tid} 通信阶段存在时间空洞：{a.end} -> {b.start}")

    for stage in stages:
        # 找到该 stage 所属的原始轨迹段（标签唯一对应一个 phase）。
        matches = [sg for sg in seg_map[stage.trip_id] if sg.label == stage.label and sg.start <= stage.start + 0.2 and sg.end >= stage.end - 0.2]
        if not matches:
            raise AssertionError(f"无法定位通信阶段轨迹：{stage.trip_id} {stage.label}")
        seg = matches[0]
        worst = math.inf
        for t in time_grid(stage.start, stage.end, COMM_SAMPLE_SEC):
            p = seg.position(t)
            if stage.mode == "直连":
                ok, _, _, margin = comm.link(p, comm.gateway, "T", "G")
                if not ok and margin < -0.05:
                    raise AssertionError(f"直连审计失败：{stage.trip_id} {stage.label} t={t:.2f}, margin={margin:.3f}dB")
                worst = min(worst, margin)
            else:
                rs = sortie_map.get(stage.relay_sortie_id)
                if rs is None:
                    raise AssertionError(f"中继阶段缺少架次：{stage.relay_sortie_id}")
                if t < rs.link_complete_time - 0.2 or t > rs.service_end_time + 0.2:
                    raise AssertionError("中继服务时段未覆盖运输需求")
                hp = Position3D(rs.lon, rs.lat, rs.hover_alt)
                ok1, _, _, m1 = comm.link(p, hp, "T", "RA")
                ok2, _, _, m2 = comm.link(hp, comm.gateway, "RB", "G")
                if not (ok1 and ok2) and min(m1, m2) < -0.05:
                    raise AssertionError(
                        f"中继通信审计失败：{stage.trip_id} {stage.label} t={t:.2f}, "
                        f"access={m1:.3f}dB, backhaul={m2:.3f}dB"
                    )
                worst = min(worst, m1, m2)
        stage.min_margin_db = worst
        min_margin = min(min_margin, worst)
    return min_margin


def audit_relay_resources(sorties: List[RelaySortie], relay: RelaySpec) -> None:
    # 高度、SOC、无人机冲突、能源组件冲突
    for s in sorties:
        if s.agl > relay.max_agl + 1e-9:
            raise AssertionError(f"{s.sortie_id} 悬停离地高度超限")
        if s.end_soc + 1e-12 < relay.reserve_ratio:
            raise AssertionError(f"{s.sortie_id} 返航 SOC 超限")
    for attr, end_attr in [("relay_id", "turnaround_end_time"), ("energy_id", "charge_end_time")]:
        groups = defaultdict(list)
        for s in sorties:
            groups[getattr(s, attr)].append(s)
        for rid, arr in groups.items():
            arr.sort(key=lambda x: x.start_time)
            for a, b in zip(arr, arr[1:]):
                if getattr(a, end_attr) > b.start_time + 1e-9:
                    raise AssertionError(f"资源 {rid} 存在重叠：{a.sortie_id} / {b.sortie_id}")


def compute_q3_metrics(q2, decoded, boxes, q2_metrics, sorties: List[RelaySortie], min_margin: float) -> Q3Metrics:
    relay_energy = sum(s.energy_kwh for s in sorties)
    transport_makespan = max((t.return_time for t in decoded.trips), default=0.0)
    relay_makespan = max((s.return_time for s in sorties), default=0.0)
    return Q3Metrics(
        weighted_lateness=q2_metrics.weighted_lateness,
        joint_makespan=max(transport_makespan, relay_makespan),
        transport_energy=q2_metrics.energy,
        relay_energy=relay_energy,
        total_energy=q2_metrics.energy + relay_energy,
        transport_trips=q2_metrics.trips,
        relay_trips=len(sorties),
        total_trips=q2_metrics.trips + len(sorties),
        hard_violations=q2_metrics.hard_violations,
        min_comm_margin_db=min_margin,
    )


# ============================================================
# 9. 单个 Q2 候选 -> Q3 联合可行解
# ============================================================


def solve_one_candidate(
    q2, seed: int, q2_iters: int,
    nodes, boxes, drone_types, drone_units, batteries, arcs, evaluator,
    relay_spec, relay_units, relay_energy_spec, dem, comm,
) -> Q3Solution:
    initial_routes = q2.construct_initial_routes(boxes, evaluator)
    decoder = q2.ResourceDecoder(evaluator, drone_units, batteries, boxes)
    routes, decoded, q2_metrics, _ = q2.alns_optimize(
        initial_routes, boxes, nodes, evaluator, decoder, q2_iters, seed
    )
    q2.validate_final_solution(routes, decoded, boxes, drone_types)

    segments = build_transport_segments(q2, decoded, nodes, drone_types, arcs)
    stages, raw_demands = build_direct_stages_and_demands(segments, comm)
    demands = ensure_demand_options(raw_demands, nodes, dem, relay_spec, comm)
    groups = compress_relay_groups(demands, relay_spec)
    sorties = schedule_relay_resources(groups, relay_spec, relay_units, relay_energy_spec)
    assign_sorties_to_stages(stages, demands)
    audit_relay_resources(sorties, relay_spec)
    min_margin = audit_communication(stages, segments, sorties, comm)
    metrics = compute_q3_metrics(q2, decoded, boxes, q2_metrics, sorties, min_margin)
    return Q3Solution(routes, decoded, q2_metrics, segments, stages, demands, sorties, metrics, seed)


def joint_score(m: Q3Metrics, base: Q3Metrics) -> float:
    def ratio(x: float, b: float) -> float:
        if b <= 1e-12:
            return 0.0 if x <= 1e-12 else x
        return x / b
    return (
        JOINT_WEIGHTS["lateness"] * ratio(m.weighted_lateness, base.weighted_lateness)
        + JOINT_WEIGHTS["makespan"] * ratio(m.joint_makespan, base.joint_makespan)
        + JOINT_WEIGHTS["energy"] * ratio(m.total_energy, base.total_energy)
        + JOINT_WEIGHTS["trips"] * ratio(m.total_trips, base.total_trips)
    )


# ============================================================
# 10. 输出：模板、详细文件、CSV、图表
# ============================================================


def copy_row_style(ws, src_row: int, dst_row: int, max_col: int) -> None:
    for c in range(1, max_col + 1):
        src, dst = ws.cell(src_row, c), ws.cell(dst_row, c)
        if src.has_style:
            dst._style = copy.copy(src._style)
        if src.number_format:
            dst.number_format = src.number_format
        if src.alignment:
            dst.alignment = copy.copy(src.alignment)
        if src.font:
            dst.font = copy.copy(src.font)
        if src.fill:
            dst.fill = copy.copy(src.fill)
        if src.border:
            dst.border = copy.copy(src.border)


def fill_result_template(q2, template_path: Path, solution: Q3Solution, boxes, output_path: Path) -> None:
    # 先使用 Q2 官方输出逻辑填 Q2 两张表，保证 Q3 引用的运输架次编号完全一致。
    temp_q2 = output_path.with_suffix(".q2tmp.xlsx")
    q2.fill_template(template_path, solution.decoded, boxes, temp_q2)
    wb = load_workbook(temp_q2)
    ws_r = wb["Q3_中继架次"]
    ws_c = wb["Q3_通信保障"]

    for ws, cols in ((ws_r, 11), (ws_c, 6)):
        for row in ws.iter_rows(min_row=2, max_row=max(ws.max_row, 2), max_col=cols):
            for cell in row:
                cell.value = None

    for r, s in enumerate(solution.relay_sorties, 2):
        if r > 2:
            copy_row_style(ws_r, 2, r, 11)
        vals = [
            s.sortie_id, s.relay_id, s.energy_id, round(s.start_time, 3),
            round(s.lon, 7), round(s.lat, 7), round(s.hover_alt, 3),
            round(s.link_complete_time, 3), round(s.service_end_time, 3),
            round(s.return_time, 3), round(s.energy_kwh, 6),
        ]
        for c, v in enumerate(vals, 1):
            ws_r.cell(r, c, v)

    for r, s in enumerate(sorted(solution.stages, key=lambda x: (x.trip_id, x.start)), 2):
        if r > 2:
            copy_row_style(ws_c, 2, r, 6)
        vals = [
            s.trip_id, s.label, round(s.start, 3), round(s.end, 3),
            s.mode, s.relay_sortie_id if s.mode == "中继" else "",
        ]
        for c, v in enumerate(vals, 1):
            ws_c.cell(r, c, v)

    wb.save(output_path)
    wb.close()
    try:
        temp_q2.unlink()
    except OSError:
        pass


def write_csvs(solution: Q3Solution, output_dir: Path) -> None:
    with (output_dir / "Q3_中继架次.csv").open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["中继架次编号","中继无人机编号","能源组件编号","开始时刻（s）","悬停经度（°）","悬停纬度（°）","悬停海拔（m）","建链完成时刻（s）","服务结束时刻（s）","返回O01时刻（s）","架次能耗（kWh）","返航SOC"])
        for s in solution.relay_sorties:
            w.writerow([s.sortie_id,s.relay_id,s.energy_id,s.start_time,s.lon,s.lat,s.hover_alt,s.link_complete_time,s.service_end_time,s.return_time,s.energy_kwh,s.end_soc])
    with (output_dir / "Q3_通信保障.csv").open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["运输架次编号","通信阶段","开始时刻（s）","结束时刻（s）","保障方式","中继架次编号","最小链路裕量（dB）"])
        for s in sorted(solution.stages, key=lambda x:(x.trip_id,x.start)):
            w.writerow([s.trip_id,s.label,s.start,s.end,s.mode,s.relay_sortie_id,s.min_margin_db])


def write_detailed_xlsx(solution: Q3Solution, output_path: Path) -> None:
    wb = Workbook()
    ws = wb.active; ws.title = "指标汇总"
    m = solution.metrics
    ws.append(["指标","数值"])
    for k,v in [
        ("配送加权迟到",m.weighted_lateness),("联合任务完成时间(s)",m.joint_makespan),
        ("运输能耗(kWh)",m.transport_energy),("中继能耗(kWh)",m.relay_energy),
        ("总能耗(kWh)",m.total_energy),("运输架次数",m.transport_trips),
        ("中继架次数",m.relay_trips),("两类无人机总架次数",m.total_trips),
        ("硬时限违反数",m.hard_violations),("最小通信链路裕量(dB)",m.min_comm_margin_db),
        ("Q2随机种子",solution.seed),
    ]: ws.append([k,v])

    ws2 = wb.create_sheet("中继架次")
    ws2.append(["架次","中继机","能源组件","开始","经度","纬度","地面高程","悬停海拔","AGL","建链完成","服务结束","返回","能耗","返航SOC","充电完成","周转完成","保障需求"])
    for s in solution.relay_sorties:
        ws2.append([s.sortie_id,s.relay_id,s.energy_id,s.start_time,s.lon,s.lat,s.candidate.ground_alt,s.hover_alt,s.agl,s.link_complete_time,s.service_end_time,s.return_time,s.energy_kwh,s.end_soc,s.charge_end_time,s.turnaround_end_time,",".join(s.demand_ids)])

    ws3 = wb.create_sheet("通信阶段")
    ws3.append(["运输架次","阶段","开始","结束","方式","中继架次","最小链路裕量dB"])
    for s in sorted(solution.stages,key=lambda x:(x.trip_id,x.start)):
        ws3.append([s.trip_id,s.label,s.start,s.end,s.mode,s.relay_sortie_id,s.min_margin_db])

    ws4 = wb.create_sheet("中继需求")
    ws4.append(["需求ID","运输架次","阶段","开始","结束","候选悬停点数","最终中继架次"])
    for d in solution.demands:
        ws4.append([d.demand_id,d.trip_id,d.label,d.start,d.end,len(d.options),d.sortie_id])
    wb.save(output_path)


def plot_coverage_map(solution: Q3Solution, nodes, out: Path) -> None:
    fig, ax = plt.subplots(figsize=(10, 8))
    # 运输路线
    for t in solution.decoded.trips:
        seq = ["O01"] + t.service_sequence + ["O01"]
        ax.plot([nodes[n].lon for n in seq], [nodes[n].lat for n in seq], linewidth=1.0, alpha=0.35)
    for nid,n in nodes.items():
        ax.scatter(n.lon,n.lat,s=46 if nid=="O01" else 22)
        ax.text(n.lon,n.lat,nid,fontsize=8,ha="left",va="bottom")
    # 中继悬停点
    for s in solution.relay_sorties:
        ax.scatter(s.lon,s.lat,marker="^",s=75)
        ax.text(s.lon,s.lat,s.sortie_id,fontsize=8,ha="right",va="top")
    ax.set_xlabel("Longitude"); ax.set_ylabel("Latitude")
    ax.set_title("Q3 Transport Routes and Relay Hover Points")
    ax.grid(alpha=0.25)
    fig.tight_layout(); fig.savefig(out,dpi=180); plt.close(fig)


def plot_comm_timeline(solution: Q3Solution, out: Path) -> None:
    trips = [t.trip_id for t in solution.decoded.trips]
    y = {tid:i for i,tid in enumerate(trips)}
    fig, ax = plt.subplots(figsize=(12, max(6, 0.34*len(trips)+2)))
    seen = set()
    for s in sorted(solution.stages,key=lambda x:(x.trip_id,x.start)):
        label = s.mode if s.mode not in seen else None
        seen.add(s.mode)
        ax.barh(y[s.trip_id], s.end-s.start, left=s.start, height=0.62, label=label)
    ax.set_yticks(list(y.values()),list(y.keys()))
    ax.set_xlabel("Time (s)"); ax.set_title("Q3 Continuous Communication Guarantee Timeline")
    ax.grid(axis="x",alpha=0.25); ax.legend()
    fig.tight_layout(); fig.savefig(out,dpi=180); plt.close(fig)


def plot_relay_gantt(solution: Q3Solution, out: Path) -> None:
    ids = sorted({s.relay_id for s in solution.relay_sorties})
    y = {rid:i for i,rid in enumerate(ids)}
    fig, ax = plt.subplots(figsize=(11, 4.8))
    for s in solution.relay_sorties:
        ax.barh(y[s.relay_id], s.return_time-s.start_time, left=s.start_time, height=0.55)
        ax.text((s.start_time+s.return_time)/2,y[s.relay_id],s.sortie_id,ha="center",va="center",fontsize=7)
    ax.set_yticks(list(y.values()),list(y.keys()))
    ax.set_xlabel("Time (s)"); ax.set_title("Q3 Relay UAV Schedule Gantt")
    ax.grid(axis="x",alpha=0.25)
    fig.tight_layout(); fig.savefig(out,dpi=180); plt.close(fig)


def plot_link_margin(solution: Q3Solution, out: Path) -> None:
    arr = sorted(solution.stages,key=lambda x:(x.start,x.trip_id))
    x = np.arange(len(arr)); y = [s.min_margin_db for s in arr]
    fig, ax = plt.subplots(figsize=(12,4.8))
    ax.plot(x,y,marker=".",linewidth=0.9)
    ax.axhline(0.0,linestyle="--",linewidth=1.0)
    ax.set_xlabel("Communication stages sorted by start time")
    ax.set_ylabel("Minimum link margin (dB)")
    ax.set_title("Q3 Minimum Bidirectional Link Margin by Stage")
    ax.grid(alpha=0.25)
    fig.tight_layout(); fig.savefig(out,dpi=180); plt.close(fig)


def write_summary(solution: Q3Solution, output_path: Path, comm: CommunicationModel) -> None:
    m = solution.metrics
    relay_seconds = sum(s.end-s.start for s in solution.stages if s.mode=="中继")
    direct_seconds = sum(s.end-s.start for s in solution.stages if s.mode=="直连")
    with output_path.open("w",encoding="utf-8") as f:
        f.write("问题三求解结果摘要\n" + "="*50 + "\n")
        f.write(f"Q2候选随机种子: {solution.seed}\n")
        f.write(f"运输架次数: {m.transport_trips}\n")
        f.write(f"中继架次数: {m.relay_trips}\n")
        f.write(f"两类无人机总架次数: {m.total_trips}\n")
        f.write(f"配送加权迟到指标: {m.weighted_lateness:.6f}\n")
        f.write(f"联合任务完成时间: {m.joint_makespan:.3f} s\n")
        f.write(f"运输能耗: {m.transport_energy:.6f} kWh\n")
        f.write(f"中继能耗: {m.relay_energy:.6f} kWh\n")
        f.write(f"总能耗: {m.total_energy:.6f} kWh\n")
        f.write(f"硬时限违反数: {m.hard_violations}\n")
        f.write(f"最小通信链路裕量: {m.min_comm_margin_db:.6f} dB\n")
        f.write(f"直连保障累计时长(各运输机累加): {direct_seconds:.3f} s\n")
        f.write(f"中继保障累计时长(各运输机累加): {relay_seconds:.3f} s\n")
        f.write(f"T-G01双向损耗门限: {comm.threshold('T','G'):.3f} dB\n")
        f.write(f"T-Relay双向损耗门限: {comm.threshold('T','RA'):.3f} dB\n")
        f.write(f"Relay-G01双向损耗门限: {comm.threshold('RB','G'):.3f} dB\n")


# ============================================================
# 11. 主程序
# ============================================================


def run(project_root: Path, output_dir: Path, q2_solver_path: Optional[Path], q2_iters: int, seed: int, q2_candidates: int) -> Q3Solution:
    t0 = time.time(); output_dir.mkdir(parents=True, exist_ok=True)
    if q2_solver_path is None:
        q2_solver_path = find_q2_solver(project_root)
    q2 = import_q2_module(q2_solver_path)

    node_file = resolve_input_file(project_root,"调度中心与服务区",".xlsx")
    box_file = resolve_input_file(project_root,"物资需求与配送时限",".xlsx")
    drone_file = resolve_input_file(project_root,"运输无人机数据",".xlsx")
    relay_file = resolve_input_file(project_root,"中继无人机数据",".xlsx")
    comm_file = resolve_input_file(project_root,"通信链路参数",".xlsx")
    dem_file = resolve_input_file(project_root,"镇龙乡及周边30米DEM",".mat")
    template_file = resolve_input_file(project_root,"结果提交模板",".xlsx")

    print("[1/9] 读取 Q2/Q3 数据...")
    print("  Q2程序:",q2_solver_path)
    print("  节点:",node_file); print("  货箱:",box_file); print("  运输无人机:",drone_file)
    print("  中继无人机:",relay_file); print("  通信参数:",comm_file); print("  DEM:",dem_file)
    nodes = q2.load_nodes(node_file)
    boxes = q2.load_boxes(box_file)
    drone_types, drone_units, batteries = q2.load_drone_data(drone_file)
    relay_spec, relay_units, relay_energy_spec = load_relay_data(relay_file)
    comm_params = load_comm_params(comm_file)

    print("[2/9] 构建运输 Arc Library 与 DEM 通信模型...")
    arcs = q2.ArcLibrary(nodes, drone_types, dem_file)
    evaluator = q2.RouteEvaluator(boxes, drone_types, arcs)
    dem = DEMHelper(dem_file)
    gateway = Position3D(nodes["O01"].lon,nodes["O01"].lat,nodes["O01"].ground_alt+comm_params.gateway_height_m)
    comm = CommunicationModel(comm_params,dem,gateway)
    print(f"  双向门限: T-G={comm.threshold('T','G'):.2f} dB, T-R={comm.threshold('T','RA'):.2f} dB, R-G={comm.threshold('RB','G'):.2f} dB")

    print(f"[3/9] 生成 {q2_candidates} 个 Q2 运输候选并加入通信/中继解码...")
    sols: List[Q3Solution] = []
    for k in range(q2_candidates):
        s = seed + 97*k
        print(f"  候选 {k+1}/{q2_candidates}: seed={s}, Q2 ALNS iters={q2_iters}")
        try:
            sol = solve_one_candidate(
                q2,s,q2_iters,nodes,boxes,drone_types,drone_units,batteries,arcs,evaluator,
                relay_spec,relay_units,relay_energy_spec,dem,comm,
            )
            sols.append(sol)
            m=sol.metrics
            print(f"    可行：运输{m.transport_trips}架次 + 中继{m.relay_trips}架次, 联合完成{m.joint_makespan:.1f}s, 总能耗{m.total_energy:.3f}kWh, min margin={m.min_comm_margin_db:.3f}dB")
        except Exception as e:
            print("    不可行：",repr(e))

    if not sols:
        raise RuntimeError("所有运输候选均未得到通信可行的 Q3 方案。可增加 --q2-candidates 或调整候选点搜索参数。")

    base=sols[0].metrics
    best=min(sols,key=lambda s:joint_score(s.metrics,base))
    print("[4/9] 选择联合目标最优的可行候选...")
    print(f"  选中 seed={best.seed}, joint score={joint_score(best.metrics,base):.6f}")

    print("[5/9] 最终运输与中继资源审计...")
    q2.validate_final_solution(best.routes,best.decoded,boxes,drone_types)
    audit_relay_resources(best.relay_sorties,relay_spec)
    min_margin=audit_communication(best.stages,best.segments,best.relay_sorties,comm)
    best.metrics.min_comm_margin_db=min_margin
    print("  VALID = True")

    print("[6/9] 回填结果提交模板 Q2/Q3 Sheets...")
    fill_result_template(q2,template_file,best,boxes,output_dir/"Q3_结果.xlsx")

    print("[7/9] 输出详细表与 CSV...")
    write_csvs(best,output_dir)
    write_detailed_xlsx(best,output_dir/"Q3_详细结果.xlsx")
    write_summary(best,output_dir/"Q3_summary.txt",comm)

    print("[8/9] 生成图表...")
    plot_coverage_map(best,nodes,output_dir/"fig_q3_1_coverage_map.png")
    plot_comm_timeline(best,output_dir/"fig_q3_2_comm_timeline.png")
    plot_relay_gantt(best,output_dir/"fig_q3_3_relay_gantt.png")
    plot_link_margin(best,output_dir/"fig_q3_4_link_margin.png")

    print("[9/9] 完成")
    print(f"  用时: {time.time()-t0:.2f}s")
    print("  输出目录:",output_dir)
    return best


def parse_args() -> argparse.Namespace:
    p=argparse.ArgumentParser(description="Q3 通信约束下的运输与中继联合调度")
    p.add_argument("--root",type=str,default=str(DEFAULT_PROJECT_ROOT),help="HUAWEI-CUP 项目根目录")
    p.add_argument("--output",type=str,default=str(DEFAULT_OUTPUT_DIR),help="输出目录")
    p.add_argument("--q2-solver",type=str,default="",help="问题二最终程序路径；留空自动查找")
    p.add_argument("--q2-iters",type=int,default=DEFAULT_Q2_ITERS,help="每个 Q2 候选的 ALNS 迭代次数")
    p.add_argument("--q2-candidates",type=int,default=DEFAULT_Q2_CANDIDATES,help="运输候选数；正式可设 2~3，计算更慢")
    p.add_argument("--seed",type=int,default=RANDOM_SEED,help="随机种子")
    return p.parse_args()


if __name__=="__main__":
    args=parse_args()
    run(
        Path(args.root).resolve(),Path(args.output).resolve(),
        Path(args.q2_solver).resolve() if args.q2_solver else None,
        max(0,args.q2_iters),args.seed,max(1,args.q2_candidates),
    )