# -*- coding: utf-8 -*-
"""
问题四：救援任务分区与独立资源配置优化
========================================

基于问题三最终联合调度结果，完成：
1. 自动识别同一运输架次导致的不可拆分服务区单元；
2. 用实际 Q3_通信保障关系构造“运输单元—中继架次”超图；
3. 在严格继承口径下检查 K=2、3 是否可分；
4. 在“中继任务按组完整复制”的扩展口径下，完整枚举 K=2、3 的所有无标签分区；
5. 采用区间划分（Interval Partitioning）精确核算 A/B/C 运输无人机、共享电池、
   中继无人机和中继能源组件的最少独立配置数量；
6. 输出理论中继下界方案、Pareto 候选和“运输侧零增配优先”的主提交方案；
7. 自动回填结果提交模板 Q4_分区配置；
8. 输出大量诊断图/论文备选图，并自动生成 Q4_图表说明.tex；
9. 输出 Q4_结果宏.tex、Q4_参考文献.tex 及完整 CSV/JSON/Excel 审计结果。

重要口径：
- 运输任务完全继承问题三：机型、路线、时刻、能耗均不改变；
- 实际运输—中继关联以 Q3_通信保障 中“中继架次编号”的实际使用记录为准；
- 严格情景：原中继架次不可复制，则同一超边涉及的运输单元必须同组；
- 扩展情景：若同一原中继架次涉及多个组，则各组建立一个完整任务副本；
  副本保持原中继位置、海拔、开始、建链、服务结束、返回、充电完成与能耗不变；
- 中继无人机占用：[开始, 返回+300s)；中继能源组件占用：[开始, 充电完成)；
- 运输无人机占用：[开始, 返回)；运输共享电池占用：[开始, 充电完成)；
- 同型资源在组内允许重新编号复用，但任务期间不得跨组调配；
- 通信链路预算完全继承问题三：载波频率、发射功率、天线增益、系统损耗、地形遮挡损耗、
  接收灵敏度与衰落裕量已在问题三逐轨迹采样点判定；Q4复制中继任务不改变位置/海拔/时段，
  因而不改变单链路几何与链路裕量，只改变资源归属和需要执行的中继任务副本数。

依赖：numpy, scipy, openpyxl, matplotlib
建议 Python 3.10+

典型运行：
    python q4_solver.py --root "D:/.../HUAWEI-CUP"

也可只用问题三最终输出和模板运行：
    python q4_solver.py --q3-zip output.zip --template 结果提交模板.xlsx \
        --nodes-xlsx 调度中心与服务区.xlsx --dem-mat 镇龙乡及周边30米DEM.mat \
        --output-dir output_q4

若项目根目录中存在调度中心与服务区、DEM 等原始附件，程序会额外生成真实地理分区图与 DEM 图；
若没有这些附件，核心枚举、资源核算和绝大多数图表仍可完整生成。
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import heapq
import itertools
import json
import math
import os
import re
import shutil
import statistics
import sys
import tempfile
import textwrap
import zipfile
from collections import defaultdict
from dataclasses import dataclass, field, asdict
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence, Set, Tuple

import numpy as np
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import Rectangle, FancyBboxPatch
from matplotlib.lines import Line2D
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

try:
    from scipy.io import loadmat
except Exception:
    loadmat = None


# ============================================================
# 0. 常量与默认设置
# ============================================================

SCRIPT_PATH = Path(__file__).resolve()
DEFAULT_ROOT = SCRIPT_PATH.parents[2] if len(SCRIPT_PATH.parents) >= 3 else Path.cwd()

RELAY_TURNAROUND_S = 300.0
DEFAULT_INVENTORY: Dict[str, int] = {
    "A_drone": 4,
    "B_drone": 2,
    "C_drone": 2,
    "A_battery": 6,
    "B_battery": 4,
    "C_battery": 4,
    "relay_drone": 2,
    "relay_energy": 6,
}

RESOURCE_ORDER = [
    "A_drone", "B_drone", "C_drone",
    "A_battery", "B_battery", "C_battery",
    "relay_drone", "relay_energy",
]
RESOURCE_CN = {
    "A_drone": "A型运输无人机",
    "B_drone": "B型运输无人机",
    "C_drone": "C型运输无人机",
    "A_battery": "A型共享电池",
    "B_battery": "B型共享电池",
    "C_battery": "C型共享电池",
    "relay_drone": "中继无人机",
    "relay_energy": "中继能源组件",
}
TRANSPORT_RESOURCE_KEYS = [
    "A_drone", "B_drone", "C_drone",
    "A_battery", "B_battery", "C_battery",
]

# 学术蓝绿配色，保持与前几问图形风格相近
COLORS = {
    "blue": "#2F6B8A",
    "teal": "#2A9D8F",
    "cyan": "#64B5CD",
    "green": "#5F9E6E",
    "orange": "#E9A23B",
    "red": "#C95C54",
    "purple": "#8064A2",
    "gray": "#7A8793",
    "light": "#EAF2F4",
    "dark": "#233746",
}
GROUP_COLORS = ["#2F6B8A", "#2A9D8F", "#E9A23B"]

# 尝试中文字体；若本机没有则回退
plt.rcParams["font.sans-serif"] = [
    "Microsoft YaHei", "SimHei", "Noto Sans CJK SC", "Noto Sans CJK JP", "AR PL UMing CN", "Source Han Sans CN", "Arial Unicode MS", "DejaVu Sans"
]
plt.rcParams["axes.unicode_minus"] = False
plt.rcParams["figure.dpi"] = 120
plt.rcParams["savefig.bbox"] = "tight"


# ============================================================
# 1. 数据类
# ============================================================

@dataclass(frozen=True)
class TransportTrip:
    trip_id: str
    drone_id: str
    drone_type: str
    battery_id: str
    start: float
    takeoff: float
    services: Tuple[str, ...]
    ret: float
    energy_kwh: float
    end_soc: float
    charge_end: float

    @property
    def duration(self) -> float:
        return self.ret - self.start


@dataclass(frozen=True)
class RelayTrip:
    relay_id: str
    drone_id: str
    energy_id: str
    start: float
    lon: float
    lat: float
    hover_alt: float
    link_ready: float
    service_end: float
    ret: float
    energy_kwh: float
    end_soc: float
    charge_end: float
    coverage_hint: str = ""

    @property
    def duration(self) -> float:
        return self.ret - self.start


@dataclass(frozen=True)
class CommRecord:
    transport_id: str
    phase: str
    start: float
    end: float
    mode: str
    relay_id: Optional[str]


@dataclass(frozen=True)
class DeliveryRecord:
    box_id: str
    transport_id: str
    service_id: str
    delivery_time: float
    expected_time: Optional[float]
    hard_deadline: Optional[float]


@dataclass
class Q3Data:
    transports: Dict[str, TransportTrip]
    relays: Dict[str, RelayTrip]
    comm: List[CommRecord]
    deliveries: List[DeliveryRecord]
    q3_detail_path: Path


@dataclass
class GroupEval:
    units: Tuple[int, ...]
    services: Tuple[str, ...]
    transport_ids: Tuple[str, ...]
    relay_ids: Tuple[str, ...]
    resources: Dict[str, int]
    transport_workload: float
    relay_workload: float
    joint_workload: float
    transport_energy: float
    relay_energy: float
    box_count: int
    service_count: int
    transport_trip_count: int
    relay_trip_count: int
    geo_compactness_km: Optional[float] = None


@dataclass
class PartitionEval:
    K: int
    groups: Tuple[Tuple[int, ...], ...]
    group_evals: Tuple[GroupEval, ...]
    totals: Dict[str, int]
    shortage: Dict[str, int]
    redundancy_vs_pool: Dict[str, int]
    relay_copies: int
    relay_delta_energy: float
    cv_transport: float
    cv_joint: float
    maxmin_transport_ratio: float
    transport_zero_augmentation: bool
    all_inventory_feasible: bool
    total_shortage: int
    weighted_shortage_types: int
    geo_compactness_km: Optional[float]
    canonical: str
    pareto: bool = False
    selected: bool = False
    theoretical_relay_min: bool = False


@dataclass
class FigureInfo:
    filename_png: str
    filename_pdf: Optional[str]
    title: str
    description: str
    meaning: str
    recommended: str


# ============================================================
# 2. 通用工具
# ============================================================

class DSU:
    def __init__(self, items: Iterable[Any]):
        self.p = {x: x for x in items}
        self.r = {x: 0 for x in items}

    def find(self, x: Any) -> Any:
        if self.p[x] != x:
            self.p[x] = self.find(self.p[x])
        return self.p[x]

    def union(self, a: Any, b: Any) -> None:
        ra, rb = self.find(a), self.find(b)
        if ra == rb:
            return
        if self.r[ra] < self.r[rb]:
            ra, rb = rb, ra
        self.p[rb] = ra
        if self.r[ra] == self.r[rb]:
            self.r[ra] += 1

    def components(self) -> List[List[Any]]:
        d: Dict[Any, List[Any]] = defaultdict(list)
        for x in self.p:
            d[self.find(x)].append(x)
        return list(d.values())


def fnum(x: Any, default: float = 0.0) -> float:
    if x is None or x == "":
        return default
    return float(x)


def opt_float(x: Any) -> Optional[float]:
    if x is None or x == "":
        return None
    try:
        return float(x)
    except Exception:
        return None


def split_services(x: Any) -> Tuple[str, ...]:
    s = str(x or "").strip()
    if not s:
        return tuple()
    s = s.replace("->", "→").replace("—>", "→")
    return tuple(v.strip() for v in s.split("→") if v.strip())


def cv(values: Sequence[float]) -> float:
    if not values:
        return 0.0
    mean = sum(values) / len(values)
    if abs(mean) < 1e-12:
        return 0.0
    var = sum((v - mean) ** 2 for v in values) / len(values)
    return math.sqrt(var) / mean


def maxmin_ratio(values: Sequence[float]) -> float:
    vals = [float(v) for v in values if v is not None]
    if not vals:
        return 1.0
    mn, mx = min(vals), max(vals)
    if mn <= 1e-12:
        return math.inf if mx > 1e-12 else 1.0
    return mx / mn


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def write_csv(path: Path, header: Sequence[str], rows: Iterable[Sequence[Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(list(header))
        for r in rows:
            w.writerow(list(r))


def safe_name(s: str) -> str:
    return re.sub(r"[^0-9A-Za-z_\-]+", "_", s).strip("_")


def service_sort_key(s: str) -> Tuple[int, str]:
    m = re.search(r"(\d+)", s)
    return (int(m.group(1)) if m else 10**9, s)


def group_label(i: int) -> str:
    return f"G{i+1}"


def km_distance(lon1: float, lat1: float, lon2: float, lat2: float) -> float:
    r = 6371.0088
    p1, p2 = math.radians(lat1), math.radians(lat2)
    dp = math.radians(lat2 - lat1)
    dl = math.radians(lon2 - lon1)
    a = math.sin(dp/2)**2 + math.cos(p1)*math.cos(p2)*math.sin(dl/2)**2
    return 2*r*math.asin(math.sqrt(max(0.0, min(1.0, a))))


def canonical_partition(groups: Sequence[Sequence[int]]) -> str:
    return "|".join(",".join(str(v+1) for v in g) for g in groups)


# ============================================================
# 3. 输入文件定位与读取
# ============================================================

def find_files(root: Path, patterns: Sequence[str]) -> List[Path]:
    out: List[Path] = []
    if not root.exists():
        return out
    for ptn in patterns:
        out.extend(root.rglob(ptn))
    # 去重且稳定排序
    return sorted(set(p.resolve() for p in out if p.is_file()))


def workbook_has_sheets(path: Path, required: Set[str]) -> bool:
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
        ok = required.issubset(set(wb.sheetnames))
        wb.close()
        return ok
    except Exception:
        return False


def find_q3_detail_xlsx(search_dir: Path) -> Optional[Path]:
    required = {"Q3_运输架次", "Q3_中继架次", "Q3_通信保障"}
    for p in sorted(search_dir.rglob("*.xlsx")):
        if workbook_has_sheets(p, required):
            return p
    return None


def extract_q3_zip(zip_path: Path, target: Path) -> Path:
    target.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(zip_path, "r") as z:
        z.extractall(target)
    return target


def read_sheet_records(ws) -> List[Dict[str, Any]]:
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(v).strip() if v is not None else "" for v in rows[0]]
    out = []
    for row in rows[1:]:
        if all(v is None for v in row):
            continue
        d = {headers[i]: row[i] if i < len(row) else None for i in range(len(headers))}
        out.append(d)
    return out


def load_q3_detail(path: Path) -> Q3Data:
    wb = load_workbook(path, read_only=True, data_only=True)
    needed = ["Q3_运输架次", "Q3_中继架次", "Q3_通信保障"]
    for s in needed:
        if s not in wb.sheetnames:
            raise ValueError(f"Q3详细结果缺少工作表：{s}")

    trows = read_sheet_records(wb["Q3_运输架次"])
    rrows = read_sheet_records(wb["Q3_中继架次"])
    crows = read_sheet_records(wb["Q3_通信保障"])
    drows = read_sheet_records(wb["Q3_逐箱交付"]) if "Q3_逐箱交付" in wb.sheetnames else []
    wb.close()

    transports: Dict[str, TransportTrip] = {}
    for r in trows:
        tid = str(r.get("运输架次编号", "")).strip()
        if not tid:
            continue
        transports[tid] = TransportTrip(
            trip_id=tid,
            drone_id=str(r.get("运输无人机", "") or ""),
            drone_type=str(r.get("机型", "") or "").strip(),
            battery_id=str(r.get("共享电池", "") or ""),
            start=fnum(r.get("开始时刻")),
            takeoff=fnum(r.get("起飞时刻")),
            services=split_services(r.get("服务区序列")),
            ret=fnum(r.get("返回O01")),
            energy_kwh=fnum(r.get("运输能耗kWh")),
            end_soc=fnum(r.get("返航SOC")),
            charge_end=fnum(r.get("充电完成")),
        )

    relays: Dict[str, RelayTrip] = {}
    for r in rrows:
        rid = str(r.get("中继架次编号", "")).strip()
        if not rid:
            continue
        relays[rid] = RelayTrip(
            relay_id=rid,
            drone_id=str(r.get("中继无人机", "") or ""),
            energy_id=str(r.get("能源组件", "") or ""),
            start=fnum(r.get("开始")),
            lon=fnum(r.get("悬停经度")),
            lat=fnum(r.get("悬停纬度")),
            hover_alt=fnum(r.get("悬停海拔")),
            link_ready=fnum(r.get("建链完成")),
            service_end=fnum(r.get("服务结束")),
            ret=fnum(r.get("返回")),
            energy_kwh=fnum(r.get("能耗kWh")),
            end_soc=fnum(r.get("返航SOC")),
            charge_end=fnum(r.get("充电完成")),
            coverage_hint=str(r.get("覆盖需求", "") or ""),
        )

    comm: List[CommRecord] = []
    for r in crows:
        tid = str(r.get("运输架次编号", "") or "").strip()
        if not tid:
            continue
        rr = str(r.get("中继架次编号", "") or "").strip()
        comm.append(CommRecord(
            transport_id=tid,
            phase=str(r.get("通信阶段", "") or ""),
            start=fnum(r.get("开始时刻")),
            end=fnum(r.get("结束时刻")),
            mode=str(r.get("保障方式", "") or ""),
            relay_id=rr if rr else None,
        ))

    deliveries: List[DeliveryRecord] = []
    for r in drows:
        bid = str(r.get("货箱编号", "") or "").strip()
        if not bid:
            continue
        deliveries.append(DeliveryRecord(
            box_id=bid,
            transport_id=str(r.get("运输架次编号", "") or ""),
            service_id=str(r.get("服务区", "") or ""),
            delivery_time=fnum(r.get("交付时刻")),
            expected_time=opt_float(r.get("期望时刻")),
            hard_deadline=opt_float(r.get("硬截止")),
        ))

    if not transports or not relays or not comm:
        raise ValueError("Q3详细结果读取为空，请确认文件版本。")
    return Q3Data(transports, relays, comm, deliveries, path)


def find_template(root: Path, explicit: Optional[Path]) -> Optional[Path]:
    if explicit and explicit.exists():
        return explicit
    candidates = find_files(root, ["*结果提交模板*.xlsx", "结果提交模板.xlsx"])
    for p in candidates:
        if workbook_has_sheets(p, {"Q4_分区配置"}):
            return p
    return None


def find_node_xlsx(root: Path, explicit: Optional[Path]) -> Optional[Path]:
    if explicit and explicit.exists():
        return explicit
    preferred = find_files(root, ["*调度中心*服务区*.xlsx", "*服务区*.xlsx", "调度中心与服务区.xlsx"])
    for p in preferred:
        if "结果" in p.name or "模板" in p.name:
            continue
        try:
            wb = load_workbook(p, read_only=True, data_only=True)
            found = False
            for ws in wb.worksheets:
                for row in ws.iter_rows(min_row=1, max_row=min(30, ws.max_row), values_only=True):
                    vals = {str(v).strip() for v in row if v is not None}
                    if "O01" in vals and any(v.startswith("S00") for v in vals):
                        found = True
                        break
                if found:
                    break
            wb.close()
            if found:
                return p
        except Exception:
            pass
    return None


def load_node_coordinates(path: Optional[Path]) -> Dict[str, Tuple[float, float]]:
    """兼容本题分段表头，识别 O01/Sxxx 的经纬度。"""
    if path is None or not path.exists():
        return {}

    def norm_header(v: Any) -> str:
        if v is None:
            return ""
        x = str(v).strip().lower()
        return re.sub(r"[\s（）()°_\-/]+", "", x)

    wb = load_workbook(path, read_only=True, data_only=True)
    result: Dict[str, Tuple[float, float]] = {}
    id_keys = {"节点编号", "节点", "编号", "nodeid", "id", "服务区编号", "调度中心编号"}
    id_norm = {norm_header(k) for k in id_keys}
    try:
        for ws in wb.worksheets:
            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                continue
            for hi in range(min(20, len(rows))):
                headers = [norm_header(v) for v in rows[hi]]
                idx_id = idx_lon = idx_lat = None
                for j, h in enumerate(headers):
                    if h in id_norm:
                        idx_id = j
                    if h in {"经度", "longitude", "lon", "lng"} or h.startswith("经度"):
                        idx_lon = j
                    if h in {"纬度", "latitude", "lat"} or h.startswith("纬度"):
                        idx_lat = j
                if idx_id is None or idx_lon is None or idx_lat is None:
                    continue
                for row in rows[hi + 1:]:
                    if max(idx_id, idx_lon, idx_lat) >= len(row):
                        continue
                    nid = str(row[idx_id] or "").strip()
                    if nid == "O01" or re.fullmatch(r"S\d{3}", nid):
                        try:
                            result[nid] = (float(row[idx_lon]), float(row[idx_lat]))
                        except (TypeError, ValueError):
                            pass
    finally:
        wb.close()
    return result


def find_dem_mat(root: Path, explicit: Optional[Path]) -> Optional[Path]:
    if explicit and explicit.exists():
        return explicit
    cands = find_files(root, ["*DEM*.mat", "*dem*.mat", "*.mat"])
    return cands[0] if cands else None


def load_dem_grid(path: Optional[Path]) -> Tuple[Optional[np.ndarray], Optional[np.ndarray], Optional[np.ndarray]]:
    if path is None or loadmat is None or not path.exists():
        return None, None, None
    try:
        d = loadmat(path, squeeze_me=True)
    except Exception:
        return None, None, None
    arrays = {k: np.asarray(v) for k, v in d.items() if not k.startswith("__") and isinstance(v, np.ndarray)}
    dem = None
    lat = None
    lon = None
    for k, a in arrays.items():
        kl = k.lower()
        if a.ndim == 2 and min(a.shape) > 10 and ("dem" in kl or "elev" in kl or "height" in kl):
            dem = a.astype(float)
            break
    if dem is None:
        mats = [a for a in arrays.values() if a.ndim == 2 and min(a.shape) > 10 and np.issubdtype(a.dtype, np.number)]
        if mats:
            dem = max(mats, key=lambda x: x.size).astype(float)
    if dem is None:
        return None, None, None
    for k, a in arrays.items():
        kl = k.lower()
        flat = np.ravel(a).astype(float) if np.issubdtype(a.dtype, np.number) else None
        if flat is None:
            continue
        if ("lat" in kl or "latitude" in kl) and flat.size in dem.shape:
            lat = flat
        if ("lon" in kl or "lng" in kl or "longitude" in kl) and flat.size in dem.shape:
            lon = flat
    # 若方向长度相反，绘图时再兼容
    return dem, lat, lon


def discover_named_inputs(root: Path) -> Dict[str, Optional[Path]]:
    patterns = {
        "problem_docx": ["*山区洪涝灾害下无人机运输与通信协同优化*.docx"],
        "geo_pdf": ["*镇龙乡地理空间数据说明*.pdf"],
        "comm_xlsx": ["*通信链路参数*.xlsx"],
        "relay_xlsx": ["*中继无人机数据*.xlsx"],
        "transport_xlsx": ["*运输无人机数据*.xlsx"],
        "node_xlsx": ["*调度中心*服务区*.xlsx"],
        "demand_xlsx": ["*物资需求*配送时限*.xlsx"],
        "q2_solver": ["q2_solver*.py"],
        "q3_solver": ["q3_solver_joint_v2_1.py"],
        "dem_mat": ["*DEM*.mat", "*dem*.mat"],
    }
    out: Dict[str, Optional[Path]] = {}
    for key, pats in patterns.items():
        fs = find_files(root, pats)
        out[key] = fs[0] if fs else None
    return out


# ============================================================
# 4. 问题三一致性与关联构造
# ============================================================

def q3_integrity_checks(data: Q3Data) -> Dict[str, Any]:
    checks: Dict[str, Any] = {}
    checks["transport_count"] = len(data.transports)
    checks["relay_count"] = len(data.relays)
    checks["delivery_count"] = len(data.deliveries)
    checks["unique_box_count"] = len({d.box_id for d in data.deliveries})
    checks["hard_deadline_violations"] = sum(
        1 for d in data.deliveries
        if d.hard_deadline is not None and d.delivery_time > d.hard_deadline + 1e-8
    )
    checks["transport_min_soc"] = min(t.end_soc for t in data.transports.values())
    checks["relay_min_soc"] = min(r.end_soc for r in data.relays.values())

    # 通信导出记录越过中继服务结束的审计，仅作提示，不用于 Q4 资源时段
    relay_overrun = []
    for c in data.comm:
        if c.relay_id and c.relay_id in data.relays:
            r = data.relays[c.relay_id]
            if c.end > r.service_end + 1e-8:
                relay_overrun.append((c.transport_id, c.relay_id, c.end, r.service_end, c.end-r.service_end))
    checks["comm_interval_overrun_count"] = len(relay_overrun)
    checks["comm_interval_overrun_records"] = relay_overrun
    return checks


def build_transport_units(data: Q3Data) -> Tuple[List[Tuple[str, ...]], Dict[str, int], Dict[str, int]]:
    services: Set[str] = set()
    for t in data.transports.values():
        services.update(t.services)
    dsu = DSU(sorted(services, key=service_sort_key))
    for t in data.transports.values():
        if len(t.services) >= 2:
            first = t.services[0]
            for s in t.services[1:]:
                dsu.union(first, s)
    comps = [tuple(sorted(c, key=service_sort_key)) for c in dsu.components()]
    comps.sort(key=lambda c: service_sort_key(c[0]))
    service_to_unit: Dict[str, int] = {}
    for i, comp in enumerate(comps):
        for s in comp:
            service_to_unit[s] = i
    trip_to_unit: Dict[str, int] = {}
    for tid, t in data.transports.items():
        us = {service_to_unit[s] for s in t.services}
        if len(us) != 1:
            raise AssertionError(f"运输架次 {tid} 的服务区未被压缩成同一单元：{t.services}")
        trip_to_unit[tid] = next(iter(us))
    return comps, service_to_unit, trip_to_unit


def build_actual_relay_relation(
    data: Q3Data,
    trip_to_unit: Dict[str, int],
) -> Tuple[Dict[str, Tuple[str, ...]], Dict[str, Tuple[int, ...]], Dict[str, Tuple[str, ...]]]:
    relay_to_trips: Dict[str, Set[str]] = defaultdict(set)
    trip_to_relays: Dict[str, Set[str]] = defaultdict(set)
    for c in data.comm:
        if c.relay_id:
            relay_to_trips[c.relay_id].add(c.transport_id)
            trip_to_relays[c.transport_id].add(c.relay_id)
    # 保留所有实际出现的中继；若中继表有但保障表没出现也记录空超边
    relay_to_units: Dict[str, Tuple[int, ...]] = {}
    for rid in sorted(data.relays):
        tids = relay_to_trips.get(rid, set())
        units = sorted({trip_to_unit[t] for t in tids if t in trip_to_unit})
        relay_to_units[rid] = tuple(units)
    return (
        {rid: tuple(sorted(ts)) for rid, ts in relay_to_trips.items()},
        relay_to_units,
        {tid: tuple(sorted(rs)) for tid, rs in trip_to_relays.items()},
    )


def strict_components(n_units: int, relay_to_units: Dict[str, Tuple[int, ...]]) -> List[Tuple[int, ...]]:
    dsu = DSU(range(n_units))
    for us in relay_to_units.values():
        if len(us) >= 2:
            for u in us[1:]:
                dsu.union(us[0], u)
    comps = [tuple(sorted(c)) for c in dsu.components()]
    comps.sort(key=lambda c: c[0])
    return comps


# ============================================================
# 5. 区间划分与资源配置
# ============================================================

def max_overlap(intervals: Sequence[Tuple[float, float]]) -> int:
    events: List[Tuple[float, int]] = []
    for s, e in intervals:
        if e < s - 1e-9:
            raise ValueError(f"非法区间 [{s},{e})")
        events.append((s, +1))
        events.append((e, -1))
    # 半开区间：同一时刻先释放(-1)再占用(+1)
    events.sort(key=lambda x: (x[0], x[1]))
    cur = 0
    best = 0
    for _, d in events:
        cur += d
        best = max(best, cur)
    return best


def interval_partition(
    jobs: Sequence[Tuple[str, float, float]],
    prefix: str,
) -> Tuple[int, Dict[str, str], List[Tuple[str, str, float, float]]]:
    """按开始时刻贪心，为区间任务分配最少同质资源。"""
    if not jobs:
        return 0, {}, []
    ordered = sorted(jobs, key=lambda x: (x[1], x[2], x[0]))
    busy: List[Tuple[float, int]] = []  # (释放时刻, resource index)
    free: List[int] = []
    next_idx = 1
    assign: Dict[str, str] = {}
    detailed: List[Tuple[str, str, float, float]] = []
    for jid, s, e in ordered:
        while busy and busy[0][0] <= s + 1e-10:
            _, idx = heapq.heappop(busy)
            heapq.heappush(free, idx)
        if free:
            idx = heapq.heappop(free)
        else:
            idx = next_idx
            next_idx += 1
        rid = f"{prefix}{idx:02d}"
        assign[jid] = rid
        detailed.append((jid, rid, s, e))
        heapq.heappush(busy, (e, idx))
    return next_idx - 1, assign, detailed


def set_partitions_k(n: int, k: int) -> Iterable[Tuple[Tuple[int, ...], ...]]:
    """Restricted-growth string：枚举无标签的 k 组集合分区。"""
    if k < 1 or k > n:
        return
    a = [0] * n

    def rec(i: int, max_label: int):
        if i == n:
            if max_label == k - 1:
                groups = [[] for _ in range(k)]
                for idx, lab in enumerate(a):
                    groups[lab].append(idx)
                yield tuple(tuple(g) for g in groups)
            return
        upper = min(max_label + 1, k - 1)
        for lab in range(upper + 1):
            a[i] = lab
            yield from rec(i + 1, max(max_label, lab))

    a[0] = 0
    yield from rec(1, 0)


def group_geo_compactness(services: Sequence[str], node_xy: Dict[str, Tuple[float, float]]) -> Optional[float]:
    pts = [node_xy[s] for s in services if s in node_xy]
    if len(pts) <= 1:
        return 0.0 if pts else None
    lonc = sum(p[0] for p in pts) / len(pts)
    latc = sum(p[1] for p in pts) / len(pts)
    return sum(km_distance(lonc, latc, p[0], p[1]) for p in pts)


def build_group_eval(
    unit_set: Sequence[int],
    units: Sequence[Tuple[str, ...]],
    data: Q3Data,
    trip_to_unit: Dict[str, int],
    relay_to_units: Dict[str, Tuple[int, ...]],
    node_xy: Dict[str, Tuple[float, float]],
) -> GroupEval:
    us = tuple(sorted(unit_set))
    u_set = set(us)
    services = tuple(sorted({s for u in us for s in units[u]}, key=service_sort_key))
    tids = tuple(sorted(tid for tid, u in trip_to_unit.items() if u in u_set))
    relay_ids = tuple(sorted(rid for rid, edge in relay_to_units.items() if u_set.intersection(edge)))

    resources: Dict[str, int] = {}
    for typ in ("A", "B", "C"):
        tt = [data.transports[tid] for tid in tids if data.transports[tid].drone_type == typ]
        resources[f"{typ}_drone"] = max_overlap([(t.start, t.ret) for t in tt])
        resources[f"{typ}_battery"] = max_overlap([(t.start, t.charge_end) for t in tt])
    rr = [data.relays[rid] for rid in relay_ids]
    resources["relay_drone"] = max_overlap([(r.start, r.ret + RELAY_TURNAROUND_S) for r in rr])
    resources["relay_energy"] = max_overlap([(r.start, r.charge_end) for r in rr])

    twork = sum(data.transports[tid].duration for tid in tids)
    rwork = sum(data.relays[rid].duration for rid in relay_ids)
    tenergy = sum(data.transports[tid].energy_kwh for tid in tids)
    renergy = sum(data.relays[rid].energy_kwh for rid in relay_ids)
    service_set = set(services)
    box_count = sum(1 for d in data.deliveries if d.service_id in service_set)
    compact = group_geo_compactness(services, node_xy)

    return GroupEval(
        units=us,
        services=services,
        transport_ids=tids,
        relay_ids=relay_ids,
        resources=resources,
        transport_workload=twork,
        relay_workload=rwork,
        joint_workload=twork+rwork,
        transport_energy=tenergy,
        relay_energy=renergy,
        box_count=box_count,
        service_count=len(services),
        transport_trip_count=len(tids),
        relay_trip_count=len(relay_ids),
        geo_compactness_km=compact,
    )


def evaluate_partition(
    groups: Tuple[Tuple[int, ...], ...],
    units: Sequence[Tuple[str, ...]],
    data: Q3Data,
    trip_to_unit: Dict[str, int],
    relay_to_units: Dict[str, Tuple[int, ...]],
    inventory: Dict[str, int],
    pool_resources: Dict[str, int],
    node_xy: Dict[str, Tuple[float, float]],
) -> PartitionEval:
    group_evals = tuple(build_group_eval(g, units, data, trip_to_unit, relay_to_units, node_xy) for g in groups)
    totals = {r: sum(g.resources[r] for g in group_evals) for r in RESOURCE_ORDER}
    shortage = {r: max(0, totals[r] - inventory[r]) for r in RESOURCE_ORDER}
    redundancy = {r: totals[r] - pool_resources[r] for r in RESOURCE_ORDER}

    relay_copies = 0
    delta_energy = 0.0
    for rid, edge in relay_to_units.items():
        if not edge:
            continue
        lam = sum(1 for g in groups if set(g).intersection(edge))
        relay_copies += max(0, lam - 1)
        delta_energy += max(0, lam - 1) * data.relays[rid].energy_kwh

    tw = [g.transport_workload for g in group_evals]
    jw = [g.joint_workload for g in group_evals]
    compact_vals = [g.geo_compactness_km for g in group_evals]
    compact = None if any(v is None for v in compact_vals) else float(sum(v or 0.0 for v in compact_vals))
    zero_transport = all(totals[k] <= inventory[k] for k in TRANSPORT_RESOURCE_KEYS)
    all_ok = all(totals[k] <= inventory[k] for k in RESOURCE_ORDER)
    return PartitionEval(
        K=len(groups),
        groups=groups,
        group_evals=group_evals,
        totals=totals,
        shortage=shortage,
        redundancy_vs_pool=redundancy,
        relay_copies=relay_copies,
        relay_delta_energy=delta_energy,
        cv_transport=cv(tw),
        cv_joint=cv(jw),
        maxmin_transport_ratio=maxmin_ratio(tw),
        transport_zero_augmentation=zero_transport,
        all_inventory_feasible=all_ok,
        total_shortage=sum(shortage.values()),
        weighted_shortage_types=sum(1 for v in shortage.values() if v > 0),
        geo_compactness_km=compact,
        canonical=canonical_partition(groups),
    )


def pareto_front(results: Sequence[PartitionEval]) -> List[PartitionEval]:
    """在缺口、中继资源、复制代价、运输工作量均衡上做非支配筛选。"""
    vals = []
    for r in results:
        vals.append((
            r.total_shortage,
            r.weighted_shortage_types,
            r.totals["relay_drone"],
            r.totals["relay_energy"],
            r.relay_copies,
            r.cv_transport,
        ))
    keep = [True] * len(results)
    for i, vi in enumerate(vals):
        if not keep[i]:
            continue
        for j, vj in enumerate(vals):
            if i == j:
                continue
            no_worse = all(a <= b + 1e-12 for a, b in zip(vj, vi))
            strictly = any(a < b - 1e-12 for a, b in zip(vj, vi))
            if no_worse and strictly:
                keep[i] = False
                break
    front = []
    for r, k in zip(results, keep):
        r.pareto = bool(k)
        if k:
            front.append(r)
    return front


def select_main_solution(results: Sequence[PartitionEval]) -> PartitionEval:
    """
    主提交规则（无价格权重）：
    1) 优先运输侧零增配；若不存在，则先最小运输资源缺口；
    2) 最小化中继无人机总配置；
    3) 最小化中继能源组件总配置；
    4) 最小化中继任务复制数；
    5) 最小化运输工作量 CV；
    6) 若有坐标，再最小化地理离散度；
    7) 用 canonical 保证确定性。
    """
    zero = [r for r in results if r.transport_zero_augmentation]
    if zero:
        pool = zero
        prefix = ()
    else:
        # 按运输侧缺口总数和缺口类型数优先
        best_transport_short = min(
            (sum(r.shortage[k] for k in TRANSPORT_RESOURCE_KEYS),
             sum(1 for k in TRANSPORT_RESOURCE_KEYS if r.shortage[k] > 0))
            for r in results
        )
        pool = [r for r in results if (
            sum(r.shortage[k] for k in TRANSPORT_RESOURCE_KEYS),
            sum(1 for k in TRANSPORT_RESOURCE_KEYS if r.shortage[k] > 0)
        ) == best_transport_short]
        prefix = best_transport_short

    def key(r: PartitionEval):
        geo = r.geo_compactness_km if r.geo_compactness_km is not None else 1e100
        return (
            r.totals["relay_drone"],
            r.totals["relay_energy"],
            r.relay_copies,
            r.cv_transport,
            geo,
            r.canonical,
        )
    ans = min(pool, key=key)
    ans.selected = True
    return ans


def select_theoretical_relay_min(results: Sequence[PartitionEval]) -> PartitionEval:
    ans = min(results, key=lambda r: (
        r.totals["relay_drone"],
        r.totals["relay_energy"],
        r.relay_copies,
        r.total_shortage,
        r.cv_transport,
        r.canonical,
    ))
    ans.theoretical_relay_min = True
    return ans


# ============================================================
# 6. 实体资源重新编号与审计
# ============================================================

def allocate_selected_resources(
    sol: PartitionEval,
    data: Q3Data,
) -> Tuple[List[Dict[str, Any]], List[Dict[str, Any]]]:
    transport_rows: List[Dict[str, Any]] = []
    relay_rows: List[Dict[str, Any]] = []
    for gi, g in enumerate(sol.group_evals, start=1):
        # 运输无人机与电池分别按机型分配
        for typ in ("A", "B", "C"):
            tt = [data.transports[tid] for tid in g.transport_ids if data.transports[tid].drone_type == typ]
            drone_jobs = [(t.trip_id, t.start, t.ret) for t in tt]
            bat_jobs = [(t.trip_id, t.start, t.charge_end) for t in tt]
            _, d_assign, _ = interval_partition(drone_jobs, f"G{gi}-{typ}U-")
            _, b_assign, _ = interval_partition(bat_jobs, f"G{gi}-{typ}B-")
            for t in sorted(tt, key=lambda x: (x.start, x.trip_id)):
                transport_rows.append({
                    "K": sol.K,
                    "group": gi,
                    "trip_id": t.trip_id,
                    "type": typ,
                    "new_drone": d_assign[t.trip_id],
                    "new_battery": b_assign[t.trip_id],
                    "start": t.start,
                    "return": t.ret,
                    "charge_end": t.charge_end,
                    "services": "→".join(t.services),
                })

        # 每组需要的原中继任务各复制一次；组内同一原中继只出现一次
        rr = [data.relays[rid] for rid in g.relay_ids]
        drone_jobs = [(r.relay_id, r.start, r.ret + RELAY_TURNAROUND_S) for r in rr]
        energy_jobs = [(r.relay_id, r.start, r.charge_end) for r in rr]
        _, d_assign, _ = interval_partition(drone_jobs, f"G{gi}-RU-")
        _, e_assign, _ = interval_partition(energy_jobs, f"G{gi}-RE-")
        for r in sorted(rr, key=lambda x: (x.start, x.relay_id)):
            relay_rows.append({
                "K": sol.K,
                "group": gi,
                "source_relay_trip": r.relay_id,
                "copy_id": f"{r.relay_id}-G{gi}",
                "new_relay_drone": d_assign[r.relay_id],
                "new_energy": e_assign[r.relay_id],
                "start": r.start,
                "link_ready": r.link_ready,
                "service_end": r.service_end,
                "return": r.ret,
                "drone_release": r.ret + RELAY_TURNAROUND_S,
                "charge_end": r.charge_end,
                "energy_kwh": r.energy_kwh,
                "lon": r.lon,
                "lat": r.lat,
                "hover_alt": r.hover_alt,
            })
    return transport_rows, relay_rows


def check_no_resource_conflict(rows: Sequence[Dict[str, Any]], resource_field: str, s_field: str, e_field: str) -> List[str]:
    issues = []
    by: Dict[str, List[Tuple[float, float, str]]] = defaultdict(list)
    for r in rows:
        by[str(r[resource_field])].append((float(r[s_field]), float(r[e_field]), str(r.get("trip_id") or r.get("copy_id"))))
    for rid, arr in by.items():
        arr.sort()
        for a, b in zip(arr, arr[1:]):
            if b[0] < a[1] - 1e-8:
                issues.append(f"{rid}: {a[2]} [{a[0]:.3f},{a[1]:.3f}) 与 {b[2]} [{b[0]:.3f},{b[1]:.3f}) 重叠")
    return issues


# ============================================================
# 7. 结果表与 Excel 输出
# ============================================================

def partition_to_flat_row(r: PartitionEval) -> List[Any]:
    row: List[Any] = [
        r.K, r.canonical, r.transport_zero_augmentation, r.all_inventory_feasible,
        r.total_shortage, r.weighted_shortage_types,
        r.relay_copies, r.relay_delta_energy,
        r.cv_transport, r.cv_joint, r.maxmin_transport_ratio,
        r.geo_compactness_km if r.geo_compactness_km is not None else "",
        int(r.pareto), int(r.selected), int(r.theoretical_relay_min),
    ]
    for k in RESOURCE_ORDER:
        row.append(r.totals[k])
    for k in RESOURCE_ORDER:
        row.append(r.shortage[k])
    for g in r.group_evals:
        row.append(";".join(g.services))
    return row


def partition_csv_header(K: int) -> List[str]:
    h = [
        "K", "canonical_units", "transport_zero_augmentation", "all_inventory_feasible",
        "total_shortage", "shortage_resource_types", "relay_copy_count", "relay_delta_energy_kWh",
        "CV_transport_workload", "CV_joint_workload", "max_min_transport_workload_ratio",
        "geo_compactness_km", "pareto", "selected", "theoretical_relay_min",
    ]
    h += [f"total_{k}" for k in RESOURCE_ORDER]
    h += [f"shortage_{k}" for k in RESOURCE_ORDER]
    h += [f"group_{i}_services" for i in range(1, K+1)]
    return h


def write_template_q4(template: Path, out_path: Path, selected: Dict[int, PartitionEval]) -> None:
    shutil.copy2(template, out_path)
    wb = load_workbook(out_path)
    if "Q4_分区配置" not in wb.sheetnames:
        raise ValueError("结果提交模板缺少 Q4_分区配置 工作表")
    ws = wb["Q4_分区配置"]
    # 只清空数据区，不改表头与模板结构
    for row in ws.iter_rows(min_row=2, max_row=max(ws.max_row, 20), min_col=1, max_col=11):
        for cell in row:
            cell.value = None
    rr = 2
    for K in (2, 3):
        sol = selected[K]
        for gi, g in enumerate(sol.group_evals, start=1):
            vals = [
                K,
                gi,
                ",".join(g.services),
                g.resources["A_drone"],
                g.resources["B_drone"],
                g.resources["C_drone"],
                g.resources["A_battery"],
                g.resources["B_battery"],
                g.resources["C_battery"],
                g.resources["relay_drone"],
                g.resources["relay_energy"],
            ]
            for c, v in enumerate(vals, start=1):
                ws.cell(rr, c, v)
            rr += 1
    wb.save(out_path)


def style_detail_workbook(wb: Workbook) -> None:
    header_fill = PatternFill("solid", fgColor="2F6B8A")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="D7E0E5")
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = Border(bottom=thin)
        for col in range(1, ws.max_column + 1):
            maxlen = 0
            for row in range(1, min(ws.max_row, 200) + 1):
                v = ws.cell(row, col).value
                if v is not None:
                    maxlen = max(maxlen, len(str(v)))
            ws.column_dimensions[get_column_letter(col)].width = min(max(maxlen + 2, 10), 32)
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="center", wrap_text=True)


def write_detail_xlsx(
    out_path: Path,
    units: Sequence[Tuple[str, ...]],
    relay_to_units: Dict[str, Tuple[int, ...]],
    relay_to_trips: Dict[str, Tuple[str, ...]],
    pool_resources: Dict[str, int],
    inventory: Dict[str, int],
    selected: Dict[int, PartitionEval],
    theoretical: Dict[int, PartitionEval],
    audits: Dict[str, Any],
) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "运输不可拆分单元"
    ws.append(["单元编号", "服务区", "服务区数量"])
    for i, u in enumerate(units, start=1):
        ws.append([f"U{i}", ",".join(u), len(u)])

    ws = wb.create_sheet("中继超边")
    ws.append(["中继架次", "运输架次", "涉及单元", "超边规模"])
    for rid in sorted(relay_to_units):
        us = relay_to_units[rid]
        ws.append([rid, ",".join(relay_to_trips.get(rid, ())), ",".join(f"U{x+1}" for x in us), len(us)])

    ws = wb.create_sheet("共享基准与库存")
    ws.append(["资源", "统一共享最少需求", "现有库存", "库存冗余"])
    for r in RESOURCE_ORDER:
        ws.append([RESOURCE_CN[r], pool_resources[r], inventory[r], inventory[r]-pool_resources[r]])

    ws = wb.create_sheet("主方案汇总")
    ws.append(["K", "组", "服务区", *[RESOURCE_CN[r] for r in RESOURCE_ORDER], "运输工作量/s", "联合工作量/s"])
    for K in (2, 3):
        for gi, g in enumerate(selected[K].group_evals, start=1):
            ws.append([K, gi, ",".join(g.services), *[g.resources[r] for r in RESOURCE_ORDER], g.transport_workload, g.joint_workload])

    ws = wb.create_sheet("主方案指标")
    ws.append(["K", "中继复制架次", "新增中继能耗/kWh", "运输工作量CV", "联合工作量CV", "理论最少中继机", "主方案中继机", "主方案中继能源", "总资源缺口"])
    for K in (2, 3):
        s, t = selected[K], theoretical[K]
        ws.append([K, s.relay_copies, s.relay_delta_energy, s.cv_transport, s.cv_joint,
                   t.totals["relay_drone"], s.totals["relay_drone"], s.totals["relay_energy"], s.total_shortage])

    ws = wb.create_sheet("Q3审计")
    ws.append(["检查项", "结果"])
    for k, v in audits.items():
        if k == "comm_interval_overrun_records":
            continue
        ws.append([k, v])
    if audits.get("comm_interval_overrun_records"):
        ws.append([])
        ws.append(["通信越界记录", "中继", "导出结束", "中继服务结束", "越界秒数"])
        for x in audits["comm_interval_overrun_records"]:
            ws.append(list(x))

    style_detail_workbook(wb)
    wb.save(out_path)


# ============================================================
# 8. 图表注册与绘图函数
# ============================================================

class FigureRegistry:
    def __init__(self, outdir: Path, save_pdf: bool = True):
        self.outdir = outdir
        self.outdir.mkdir(parents=True, exist_ok=True)
        self.save_pdf = save_pdf
        self.items: List[FigureInfo] = []

    def save(self, fig, stem: str, title: str, description: str, meaning: str, recommended: str = "备选") -> None:
        png = self.outdir / f"{stem}.png"
        pdf = self.outdir / f"{stem}.pdf" if self.save_pdf else None
        fig.tight_layout()
        fig.savefig(png, dpi=240)
        if pdf is not None:
            fig.savefig(pdf)
        plt.close(fig)
        # 只有真实成功落盘的图片才登记到图表清单，避免“说明中有、文件夹里没有”。
        if not png.exists():
            raise RuntimeError(f"图表保存失败：{png}")
        if pdf is not None and not pdf.exists():
            raise RuntimeError(f"矢量图保存失败：{pdf}")
        self.items.append(FigureInfo(png.name, pdf.name if pdf else None, title, description, meaning, recommended))


def bar_chart(reg: FigureRegistry, stem: str, labels: Sequence[str], values: Sequence[float], title: str,
              ylabel: str, description: str, meaning: str, recommended: str = "备选", horizontal: bool = False) -> None:
    fig, ax = plt.subplots(figsize=(9.2, 5.4))
    if horizontal:
        y = np.arange(len(labels))
        ax.barh(y, values, color=COLORS["blue"], alpha=0.88)
        ax.set_yticks(y); ax.set_yticklabels(labels)
        ax.set_xlabel(ylabel)
        for yi, v in zip(y, values):
            ax.text(v, yi, f" {v:.3g}", va="center", fontsize=8)
    else:
        x = np.arange(len(labels))
        ax.bar(x, values, color=COLORS["blue"], alpha=0.88)
        ax.set_xticks(x); ax.set_xticklabels(labels, rotation=25, ha="right")
        ax.set_ylabel(ylabel)
        for xi, v in zip(x, values):
            ax.text(xi, v, f"{v:.3g}", ha="center", va="bottom", fontsize=8)
    ax.set_title(title)
    ax.grid(axis="y" if not horizontal else "x", alpha=0.18)
    reg.save(fig, stem, title, description, meaning, recommended)


def hist_chart(reg: FigureRegistry, stem: str, values: Sequence[float], title: str, xlabel: str,
               description: str, meaning: str, recommended: str = "备选", integer_bins: bool = False) -> None:
    fig, ax = plt.subplots(figsize=(8.8, 5.2))
    vals = np.asarray(values, dtype=float)
    if integer_bins and len(vals):
        mn, mx = int(vals.min()), int(vals.max())
        bins = np.arange(mn - 0.5, mx + 1.5, 1)
    else:
        bins = min(30, max(8, int(math.sqrt(max(1, len(vals))))))
    ax.hist(vals, bins=bins, color=COLORS["teal"], alpha=0.86, edgecolor="white")
    ax.set_title(title); ax.set_xlabel(xlabel); ax.set_ylabel("方案数量")
    ax.grid(axis="y", alpha=0.18)
    reg.save(fig, stem, title, description, meaning, recommended)


def scatter_chart(reg: FigureRegistry, stem: str, xs: Sequence[float], ys: Sequence[float], title: str,
                  xlabel: str, ylabel: str, description: str, meaning: str, recommended: str = "备选",
                  selected: Optional[Tuple[float, float]] = None, pareto_xy: Optional[List[Tuple[float,float]]] = None) -> None:
    fig, ax = plt.subplots(figsize=(8.4, 5.8))
    ax.scatter(xs, ys, s=22, alpha=0.45, color=COLORS["blue"], label="全部分区")
    if pareto_xy:
        px, py = zip(*pareto_xy)
        ax.scatter(px, py, s=42, alpha=0.9, color=COLORS["teal"], label="非支配方案")
    if selected is not None:
        ax.scatter([selected[0]], [selected[1]], s=110, marker="*", color=COLORS["red"], label="主提交方案", zorder=10)
    ax.set_title(title); ax.set_xlabel(xlabel); ax.set_ylabel(ylabel)
    ax.grid(alpha=0.18)
    if selected is not None or pareto_xy:
        ax.legend(frameon=False)
    reg.save(fig, stem, title, description, meaning, recommended)


def plot_incidence_matrix(reg: FigureRegistry, units: Sequence[Tuple[str, ...]], relay_to_units: Dict[str, Tuple[int, ...]]) -> None:
    rids = sorted(relay_to_units)
    mat = np.zeros((len(rids), len(units)))
    for i, rid in enumerate(rids):
        for u in relay_to_units[rid]:
            mat[i, u] = 1
    fig, ax = plt.subplots(figsize=(10.5, 7.0))
    ax.imshow(mat, aspect="auto", cmap="Blues", vmin=0, vmax=1)
    ax.set_xticks(range(len(units))); ax.set_xticklabels([f"U{i+1}" for i in range(len(units))])
    ax.set_yticks(range(len(rids))); ax.set_yticklabels(rids)
    ax.set_xlabel("运输不可拆分单元"); ax.set_ylabel("中继架次")
    ax.set_title("中继架次—运输单元关联矩阵")
    for i in range(mat.shape[0]):
        for j in range(mat.shape[1]):
            if mat[i, j] > 0:
                ax.text(j, i, "●", ha="center", va="center", color=COLORS["dark"], fontsize=9)
    reg.save(fig, "fig_q4_06_hypergraph_incidence", "中继架次—运输单元关联矩阵",
             "以13个中继架次为行、9个运输不可拆分单元为列，标出实际Q3通信保障关系。",
             "矩阵中的多列非零行就是超边；它揭示共享中继把多个运输单元绑定在一起，是严格分区冲突的直接证据。", "正文强烈推荐")


def plot_bipartite_hypergraph(reg: FigureRegistry, units: Sequence[Tuple[str, ...]], relay_to_units: Dict[str, Tuple[int, ...]], strict: bool = False) -> None:
    rids = sorted(relay_to_units)
    fig, ax = plt.subplots(figsize=(11.5, 7.6))
    uy = np.linspace(0.95, 0.05, len(units))
    ry = np.linspace(0.95, 0.05, len(rids))
    ux, rx = 0.22, 0.78
    for i, u in enumerate(units):
        ax.scatter([ux], [uy[i]], s=760, marker="o", color=COLORS["blue"], edgecolor="white", linewidth=1.5, zorder=3)
        ax.text(ux, uy[i], f"U{i+1}", ha="center", va="center", fontsize=8.5, weight="bold", color="white", zorder=4)
        ax.text(ux-0.045, uy[i], ",".join(u), ha="right", va="center", fontsize=7.6, color="#303030", zorder=4)
    for j, rid in enumerate(rids):
        ax.scatter([rx], [ry[j]], s=430, marker="s", color=COLORS["teal"], edgecolor="white", linewidth=1.2, zorder=3)
        ax.text(rx, ry[j], rid.replace("Q3", ""), ha="center", va="center", fontsize=7.5, color="white", zorder=4)
        for u in relay_to_units[rid]:
            ax.plot([ux, rx], [uy[u], ry[j]], color=COLORS["gray"], alpha=0.42, linewidth=1.0, zorder=1)
    ax.text(ux, 1.01, "运输不可拆分单元", ha="center", fontsize=11, weight="bold")
    ax.text(rx, 1.01, "实际使用中继架次", ha="center", fontsize=11, weight="bold")
    ax.set_xlim(0.02, 0.98); ax.set_ylim(0.0, 1.06); ax.axis("off")
    title = "运输单元—中继架次二部关联图"
    ax.set_title(title, pad=15)
    reg.save(fig, "fig_q4_07_unit_relay_bipartite", title,
             "将运输不可拆分单元与问题三实际使用的中继架次置于两侧，连线表示该中继实际保障了该单元中的运输任务。",
             "比普通两两网络更直观地展示高阶共享关系；可观察到关联链条贯通全部单元，因此严格继承时不能切成2组或3组。", "正文强烈推荐")


def plot_strict_component_graph(reg: FigureRegistry, n_units: int, relay_to_units: Dict[str, Tuple[int, ...]]) -> None:
    fig, ax = plt.subplots(figsize=(7.2, 7.2))
    theta = np.linspace(0, 2*np.pi, n_units, endpoint=False)
    xy = {i: (math.cos(theta[i]), math.sin(theta[i])) for i in range(n_units)}
    drawn = set()
    for rid, us in relay_to_units.items():
        if len(us) < 2:
            continue
        for a, b in itertools.combinations(us, 2):
            key = tuple(sorted((a, b)))
            if key not in drawn:
                xa, ya = xy[a]; xb, yb = xy[b]
                ax.plot([xa, xb], [ya, yb], color=COLORS["gray"], alpha=0.45, linewidth=1.2)
                drawn.add(key)
    for i in range(n_units):
        x, y = xy[i]
        ax.scatter([x], [y], s=850, color=COLORS["blue"], edgecolor="white", linewidth=1.5, zorder=3)
        ax.text(x, y, f"U{i+1}", color="white", ha="center", va="center", weight="bold")
    ax.set_title("严格继承口径下的运输单元连通图")
    ax.axis("equal"); ax.axis("off")
    reg.save(fig, "fig_q4_08_strict_connectivity", "严格继承口径下的运输单元连通图",
             "把同一实际中继架次连接的运输单元两两展开，仅用于显示连通性。",
             "9个运输单元处于同一个连通分量，说明若原中继架次不可复制且资源不得跨组，则K=2、3均结构性不可行。", "正文强烈推荐")


def plot_partition_counts(reg: FigureRegistry, counts: Dict[int, int]) -> None:
    bar_chart(reg, "fig_q4_09_partition_counts", ["K=2", "K=3"], [counts[2], counts[3]],
              "完整枚举的无标签分区数量", "分区数",
              "展示9个不可拆分单元分别划分为2组和3组的无标签集合分区数量。",
              "K=2仅255个、K=3仅3025个，规模很小，因此完整枚举比使用随机启发式更可靠、可复现。", "正文推荐")


def plot_resource_compare(reg: FigureRegistry, sol: PartitionEval, inventory: Dict[str, int], stem: str) -> None:
    labels = [RESOURCE_CN[k] for k in RESOURCE_ORDER]
    demand = [sol.totals[k] for k in RESOURCE_ORDER]
    stock = [inventory[k] for k in RESOURCE_ORDER]
    x = np.arange(len(labels)); w = 0.38
    fig, ax = plt.subplots(figsize=(11.0, 5.8))
    ax.bar(x-w/2, stock, width=w, label="现有库存", color=COLORS["gray"], alpha=0.65)
    ax.bar(x+w/2, demand, width=w, label=f"K={sol.K}主方案需求", color=COLORS["blue"], alpha=0.9)
    ax.set_xticks(x); ax.set_xticklabels(labels, rotation=25, ha="right")
    ax.set_ylabel("数量"); ax.set_title(f"K={sol.K} 主方案资源需求与现有库存对比")
    ax.legend(frameon=False); ax.grid(axis="y", alpha=0.18)
    reg.save(fig, stem, f"K={sol.K}主方案资源需求与库存对比",
             "对8类资源分别比较任务组独立执行后的总需求与题目现有库存。",
             "直接识别具体资源缺口，避免用不同类型资源之间的无依据加权成本相互抵消。", "正文强烈推荐")


def plot_group_resource_heatmap(reg: FigureRegistry, sol: PartitionEval, stem: str) -> None:
    mat = np.array([[g.resources[k] for k in RESOURCE_ORDER] for g in sol.group_evals], dtype=float)
    fig, ax = plt.subplots(figsize=(11.0, 4.8 + 0.5*sol.K))
    im = ax.imshow(mat, aspect="auto", cmap="YlGnBu")
    ax.set_yticks(range(sol.K)); ax.set_yticklabels([f"任务组{g+1}" for g in range(sol.K)])
    ax.set_xticks(range(len(RESOURCE_ORDER))); ax.set_xticklabels([RESOURCE_CN[k] for k in RESOURCE_ORDER], rotation=25, ha="right")
    ax.set_title(f"K={sol.K} 各任务组独立资源配置矩阵")
    for i in range(mat.shape[0]):
        for j in range(mat.shape[1]):
            ax.text(j, i, f"{int(mat[i,j])}", ha="center", va="center", fontsize=9)
    fig.colorbar(im, ax=ax, shrink=0.75, label="配置数量")
    reg.save(fig, stem, f"K={sol.K}各任务组独立资源配置矩阵",
             "按任务组展示A/B/C运输机、电池、中继机和中继能源组件的独立最少配置。",
             "用于解释提交模板中每一行资源数量来自何处，并观察不同组之间的资源压力差异。", "正文推荐")


def plot_group_workload(reg: FigureRegistry, sol: PartitionEval, stem: str) -> None:
    labels = [f"组{i+1}" for i in range(sol.K)]
    t = np.array([g.transport_workload/3600 for g in sol.group_evals])
    r = np.array([g.relay_workload/3600 for g in sol.group_evals])
    x = np.arange(sol.K)
    fig, ax = plt.subplots(figsize=(8.4, 5.2))
    ax.bar(x, t, label="运输任务累计作业时间", color=COLORS["blue"])
    ax.bar(x, r, bottom=t, label="中继副本累计作业时间", color=COLORS["teal"], alpha=0.85)
    ax.set_xticks(x); ax.set_xticklabels(labels)
    ax.set_ylabel("累计作业时间/h"); ax.set_title(f"K={sol.K} 主方案组间工作量")
    ax.legend(frameon=False); ax.grid(axis="y", alpha=0.18)
    reg.save(fig, stem, f"K={sol.K}主方案组间工作量",
             "运输任务和中继副本的累计作业时间采用堆叠柱展示；运输任务部分不受中继复制口径影响。",
             f"运输工作量CV={sol.cv_transport:.3f}，可用于评价分区的工作量均衡性，同时避免只按服务区数量判断均衡。", "正文强烈推荐")


def plot_selected_membership(reg: FigureRegistry, sol: PartitionEval, units: Sequence[Tuple[str,...]], stem: str) -> None:
    services = sorted({s for u in units for s in u}, key=service_sort_key)
    svc_group = {}
    for gi, g in enumerate(sol.group_evals, start=1):
        for s in g.services:
            svc_group[s] = gi
    fig, ax = plt.subplots(figsize=(11.0, 2.6))
    for i, s in enumerate(services):
        g = svc_group[s]
        ax.add_patch(Rectangle((i,0), 1, 1, facecolor=GROUP_COLORS[g-1], edgecolor="white"))
        ax.text(i+0.5, 0.5, s, ha="center", va="center", color="white", fontsize=9, rotation=90 if len(services)>12 else 0)
    ax.set_xlim(0,len(services)); ax.set_ylim(0,1); ax.axis("off")
    ax.set_title(f"K={sol.K} 主方案15个服务区任务组归属")
    handles=[Line2D([0],[0],color=GROUP_COLORS[i],lw=7,label=f"任务组{i+1}") for i in range(sol.K)]
    ax.legend(handles=handles, loc="upper center", bbox_to_anchor=(0.5,-0.08), ncol=sol.K, frameon=False)
    reg.save(fig, stem, f"K={sol.K}主方案服务区归属条带图",
             "按服务区编号顺序，用颜色标识15个服务区的任务组归属。",
             "作为任务分区的紧凑总览图，适合与地图配合使用；即使没有DEM也能清楚复核提交结果。", "正文备选")


def plot_replication_matrix(reg: FigureRegistry, sol: PartitionEval, relay_to_units: Dict[str, Tuple[int,...]], stem: str) -> None:
    rids=sorted(relay_to_units)
    mat=np.zeros((len(rids),sol.K))
    for i,rid in enumerate(rids):
        edge=set(relay_to_units[rid])
        for j,g in enumerate(sol.groups):
            if edge.intersection(g): mat[i,j]=1
    fig, ax=plt.subplots(figsize=(7.5,7.2))
    ax.imshow(mat,aspect="auto",cmap="BuGn",vmin=0,vmax=1)
    ax.set_xticks(range(sol.K)); ax.set_xticklabels([f"组{j+1}" for j in range(sol.K)])
    ax.set_yticks(range(len(rids))); ax.set_yticklabels(rids)
    ax.set_title(f"K={sol.K} 中继任务按组复制矩阵")
    for i in range(mat.shape[0]):
        for j in range(mat.shape[1]):
            ax.text(j,i,"复制" if mat[i,j] else "",ha="center",va="center",fontsize=8,color=COLORS["dark"])
    reg.save(fig,stem,f"K={sol.K}中继任务按组复制矩阵",
             "行表示问题三的原中继架次，列表示任务组；单元格非零表示该组需要一个保持原参数不变的中继任务副本。",
             "可逐项解释总复制架次和新增中继能耗，并验证同一原中继在一个组内最多复制一次。", "正文推荐")


def plot_relay_copy_energy(reg: FigureRegistry, sol: PartitionEval, data: Q3Data, relay_to_units: Dict[str, Tuple[int,...]], stem: str) -> None:
    labels=[]; vals=[]
    for rid in sorted(relay_to_units):
        lam=sum(1 for g in sol.groups if set(g).intersection(relay_to_units[rid]))
        extra=max(0,lam-1)
        if extra>0:
            labels.append(rid)
            vals.append(extra*data.relays[rid].energy_kwh)
    if not labels:
        return
    bar_chart(reg,stem,labels,vals,f"K={sol.K} 各跨组中继任务的新增复制能耗","新增能耗/kWh",
              "仅展示跨组超边；柱高等于(涉及组数-1)×原中继架次能耗。",
              f"全部柱高之和为{sol.relay_delta_energy:.3f} kWh，是完整复制假设下可直接计算的中继能耗增量。", "正文推荐")


def plot_gantt_transport(reg: FigureRegistry, sol: PartitionEval, data: Q3Data, stem: str) -> None:
    fig, ax=plt.subplots(figsize=(12.0,6.5))
    y=0; yt=[]; yl=[]
    for gi,g in enumerate(sol.group_evals, start=1):
        for tid in sorted(g.transport_ids,key=lambda x:data.transports[x].start):
            t=data.transports[tid]
            ax.barh(y,t.ret-t.start,left=t.start/3600,height=0.72,color=GROUP_COLORS[gi-1],alpha=0.82)
            ax.text(t.start/3600+(t.ret-t.start)/7200,y,tid.replace("Q3",""),ha="center",va="center",fontsize=6,color="white")
            yt.append(y); yl.append(f"G{gi}-{tid}"); y+=1
        y+=0.6
    ax.set_yticks(yt); ax.set_yticklabels(yl,fontsize=7)
    ax.set_xlabel("时刻/h"); ax.set_title(f"K={sol.K} 固定运输任务按组时间分布")
    ax.grid(axis="x",alpha=0.18); ax.invert_yaxis()
    reg.save(fig,stem,f"K={sol.K}固定运输任务按组甘特图",
             "不改变问题三的任何运输时刻，只按最终任务组重新排列展示25个运输架次。",
             "用于复核分区没有改变原运输调度，并观察各组的运输高峰与并发结构。", "附录/备选")


def plot_gantt_relay(reg: FigureRegistry, sol: PartitionEval, data: Q3Data, stem: str) -> None:
    fig, ax=plt.subplots(figsize=(12.0,6.5))
    y=0; yt=[]; yl=[]
    for gi,g in enumerate(sol.group_evals, start=1):
        for rid in sorted(g.relay_ids,key=lambda x:data.relays[x].start):
            r=data.relays[rid]
            ax.barh(y,(r.ret+RELAY_TURNAROUND_S-r.start)/3600,left=r.start/3600,height=0.72,color=GROUP_COLORS[gi-1],alpha=0.82)
            ax.axvline(r.ret/3600,color=COLORS["gray"],alpha=0.10,linewidth=0.5)
            ax.text(r.start/3600+(r.ret+RELAY_TURNAROUND_S-r.start)/7200,y,rid.replace("Q3",""),ha="center",va="center",fontsize=6,color="white")
            yt.append(y); yl.append(f"G{gi}-{rid}"); y+=1
        y+=0.6
    ax.set_yticks(yt); ax.set_yticklabels(yl,fontsize=7)
    ax.set_xlabel("时刻/h"); ax.set_title(f"K={sol.K} 中继任务副本占用时段（含返航后300s周转）")
    ax.grid(axis="x",alpha=0.18); ax.invert_yaxis()
    reg.save(fig,stem,f"K={sol.K}中继副本甘特图",
             "每组列出其所需的完整中继任务副本，中继无人机占用时段延续到返回O01后300秒。",
             "直观说明为什么“新增中继架次数”不等于“新增中继无人机数量”：不同时段副本可由同一组内实体中继机顺序执行。", "正文推荐")


def plot_resource_redundancy_compare(reg: FigureRegistry, selected: Dict[int,PartitionEval], pool: Dict[str,int]) -> None:
    labels=[RESOURCE_CN[k] for k in RESOURCE_ORDER]
    x=np.arange(len(labels)); w=0.28
    p=[pool[k] for k in RESOURCE_ORDER]
    k2=[selected[2].totals[k] for k in RESOURCE_ORDER]
    k3=[selected[3].totals[k] for k in RESOURCE_ORDER]
    fig,ax=plt.subplots(figsize=(11.2,5.8))
    ax.bar(x-w,p,w,label="统一共享最少需求",color=COLORS["gray"],alpha=0.7)
    ax.bar(x,k2,w,label="K=2独立配置",color=COLORS["blue"],alpha=0.9)
    ax.bar(x+w,k3,w,label="K=3独立配置",color=COLORS["teal"],alpha=0.9)
    ax.set_xticks(x); ax.set_xticklabels(labels,rotation=25,ha="right")
    ax.set_ylabel("资源数量"); ax.set_title("统一共享、2组独立与3组独立的资源规模比较")
    ax.legend(frameon=False); ax.grid(axis="y",alpha=0.18)
    reg.save(fig,"fig_q4_32_resource_redundancy_compare","三种资源共享口径的配置规模比较",
             "将问题三固定任务统一共享时的理论最少资源与2组、3组主方案独立配置需求并列。",
             "直接展示取消跨组资源复用造成的配置冗余，并可看出分组越细通信资源冗余越明显。", "正文强烈推荐")


def plot_pooled_vs_original_ids(reg: FigureRegistry, data: Q3Data, pool: Dict[str,int]) -> None:
    original={
        "A_drone":len({t.drone_id for t in data.transports.values() if t.drone_type=="A"}),
        "B_drone":len({t.drone_id for t in data.transports.values() if t.drone_type=="B"}),
        "C_drone":len({t.drone_id for t in data.transports.values() if t.drone_type=="C"}),
        "A_battery":len({t.battery_id for t in data.transports.values() if t.drone_type=="A"}),
        "B_battery":len({t.battery_id for t in data.transports.values() if t.drone_type=="B"}),
        "C_battery":len({t.battery_id for t in data.transports.values() if t.drone_type=="C"}),
        "relay_drone":len({r.drone_id for r in data.relays.values()}),
        "relay_energy":len({r.energy_id for r in data.relays.values()}),
    }
    labels=[RESOURCE_CN[k] for k in RESOURCE_ORDER]
    x=np.arange(len(labels)); w=0.36
    fig,ax=plt.subplots(figsize=(11.0,5.8))
    ax.bar(x-w/2,[original[k] for k in RESOURCE_ORDER],w,label="问题三原编号使用数",color=COLORS["gray"],alpha=0.7)
    ax.bar(x+w/2,[pool[k] for k in RESOURCE_ORDER],w,label="固定时刻表最少需求",color=COLORS["blue"],alpha=0.9)
    ax.set_xticks(x);ax.set_xticklabels(labels,rotation=25,ha="right")
    ax.set_ylabel("数量");ax.set_title("问题三原资源编号使用数与最少共享需求")
    ax.legend(frameon=False);ax.grid(axis="y",alpha=0.18)
    reg.save(fig,"fig_q4_33_original_ids_vs_minimum","问题三资源编号与最少共享需求对比",
             "在保持问题三所有任务时刻不变的情况下，重新做区间划分，比较原方案实际使用的资源编号数与理论最少数量。",
             "说明“出现过多少个编号”并不等于“最少需要多少资源”；B型电池和中继能源组件存在进一步复用空间。", "正文推荐")


def plot_balance_sensitivity(reg: FigureRegistry, results: Sequence[PartitionEval], K: int, stem: str) -> None:
    eps_list=np.linspace(0.05,0.80,16)
    ys=[]
    for eps in eps_list:
        feasible=[r for r in results if r.cv_transport<=eps+1e-12]
        ys.append(min((r.totals["relay_drone"] for r in feasible),default=np.nan))
    fig,ax=plt.subplots(figsize=(8.6,5.2))
    ax.plot(eps_list,ys,marker="o",color=COLORS["blue"],linewidth=2)
    ax.set_xlabel(r"运输工作量均衡阈值  $CV(W^T)\leq\epsilon$")
    ax.set_ylabel("可达到的最少中继无人机总配置")
    ax.set_title(f"K={K} 均衡要求对中继资源下界的敏感性")
    ax.grid(alpha=0.18)
    reg.save(fig,stem,f"K={K}均衡阈值敏感性",
             "逐步收紧运输工作量CV上限，重新统计满足该均衡要求的分区中最小中继无人机需求。",
             "用于说明过强的均衡要求可能迫使更多共享中继被切开，从而提高通信资源配置下界。", "正文备选")


def plot_top_candidates(reg: FigureRegistry, results: Sequence[PartitionEval], K: int, stem: str) -> None:
    ranked=sorted(results,key=lambda r:(
        0 if r.transport_zero_augmentation else 1,
        sum(r.shortage[k] for k in TRANSPORT_RESOURCE_KEYS),
        r.totals["relay_drone"],r.totals["relay_energy"],r.relay_copies,r.cv_transport,r.canonical))[:25]
    x=np.arange(len(ranked))
    relay=[r.totals["relay_drone"] for r in ranked]
    cvv=[r.cv_transport for r in ranked]
    fig,ax=plt.subplots(figsize=(11.2,5.4))
    sc=ax.scatter(x,relay,s=70,c=cvv,cmap="viridis",edgecolor="white")
    ax.set_xlabel("按主筛选规则排序的前25个候选")
    ax.set_ylabel("中继无人机总配置")
    ax.set_title(f"K={K} 前25个候选方案：中继配置与均衡性")
    fig.colorbar(sc,ax=ax,label="运输工作量CV")
    ax.grid(alpha=0.18)
    reg.save(fig,stem,f"K={K}前25候选方案对比",
             "按主筛选规则取前25个候选，纵轴为中继无人机需求，颜色表示运输工作量CV。",
             "展示最终方案附近仍存在的资源—均衡权衡，避免只报告一个方案而看不到邻近候选结构。", "附录/备选")


def plot_geo_partition(reg: FigureRegistry, sol: PartitionEval, node_xy: Dict[str,Tuple[float,float]], stem: str,
                       dem: Optional[np.ndarray]=None, latv: Optional[np.ndarray]=None, lonv: Optional[np.ndarray]=None) -> None:
    if not node_xy:
        return
    fig,ax=plt.subplots(figsize=(9.2,7.2))
    _draw_dem_background(ax,dem,latv,lonv,0.23)
    if "O01" in node_xy:
        ax.scatter(*node_xy["O01"],marker="*",s=220,color=COLORS["red"],label="O01",zorder=5)
        ax.text(*node_xy["O01"]," O01",fontsize=9,weight="bold")
    for gi,g in enumerate(sol.group_evals,start=1):
        pts=[node_xy[s] for s in g.services if s in node_xy]
        if not pts: continue
        xs=[p[0] for p in pts];ys=[p[1] for p in pts]
        ax.scatter(xs,ys,s=70,color=GROUP_COLORS[gi-1],label=f"任务组{gi}",edgecolor="white",linewidth=0.8,zorder=4)
        for s in g.services:
            if s in node_xy:
                x,y=node_xy[s];ax.text(x,y,f" {s}",fontsize=8)
    ax.set_xlabel("经度");ax.set_ylabel("纬度");ax.set_title(f"K={sol.K} 主方案服务区空间分区")
    ax.legend(frameon=False);ax.grid(alpha=0.12)
    reg.save(fig,stem,f"K={sol.K}主方案空间分区图",
             "按真实服务区经纬度绘制最终任务组；若可读取DEM则以地形灰度作为背景。",
             "用于判断任务组的空间分布和紧凑程度。地理距离只作为辅助解释，不覆盖固定运输与中继关联约束。", "正文强烈推荐")


def plot_relay_position_map(reg: FigureRegistry, sol: PartitionEval, data: Q3Data, node_xy: Dict[str,Tuple[float,float]], stem: str) -> None:
    if not node_xy:
        return
    fig,ax=plt.subplots(figsize=(9.2,7.2))
    for gi,g in enumerate(sol.group_evals,start=1):
        for s in g.services:
            if s in node_xy:
                x,y=node_xy[s];ax.scatter(x,y,s=42,color=GROUP_COLORS[gi-1],alpha=0.8);ax.text(x,y,f" {s}",fontsize=7)
        for rid in g.relay_ids:
            r=data.relays[rid]
            ax.scatter(r.lon,r.lat,marker="^",s=70,color=GROUP_COLORS[gi-1],edgecolor="black",linewidth=0.3)
    ax.set_xlabel("经度");ax.set_ylabel("纬度");ax.set_title(f"K={sol.K} 任务组与所需中继悬停点")
    ax.grid(alpha=0.12)
    reg.save(fig,stem,f"K={sol.K}任务组与中继悬停点",
             "同一颜色同时标记任务组服务区和该组需要复制的原中继悬停位置。",
             "揭示分区后通信保障的空间复制范围，并可解释某些组虽然服务区较少仍需要较多中继资源。", "正文备选")



def _draw_dem_background(ax, dem: Optional[np.ndarray], latv: Optional[np.ndarray], lonv: Optional[np.ndarray], alpha: float = 0.26) -> None:
    """在经纬度坐标轴上叠加DEM灰度背景。"""
    if dem is None or latv is None or lonv is None:
        return
    try:
        extent=[float(np.nanmin(lonv)),float(np.nanmax(lonv)),float(np.nanmin(latv)),float(np.nanmax(latv))]
        # DEM只用于可视化背景，不改变任何物理计算；绘图时下采样可显著减少批量出图耗时。
        rs=max(1,int(math.ceil(dem.shape[0]/650)))
        cs=max(1,int(math.ceil(dem.shape[1]/650)))
        dem_show=dem[::rs,::cs]
        ax.imshow(dem_show,extent=extent,origin="lower",cmap="Greys",alpha=alpha,aspect="auto",zorder=0)
    except Exception:
        pass


def plot_geo_routes(reg: FigureRegistry, sol: PartitionEval, data: Q3Data,
                    node_xy: Dict[str,Tuple[float,float]], stem: str,
                    dem: Optional[np.ndarray]=None, latv: Optional[np.ndarray]=None,
                    lonv: Optional[np.ndarray]=None) -> None:
    """按任务组绘制问题三固定运输路线；分区只改变归属，不改变路线本身。"""
    if not node_xy or "O01" not in node_xy:
        return
    fig,ax=plt.subplots(figsize=(9.6,7.6))
    _draw_dem_background(ax,dem,latv,lonv,0.24)
    trip_group={tid:gi for gi,g in enumerate(sol.group_evals,start=1) for tid in g.transport_ids}
    for tid,t in data.transports.items():
        gi=trip_group.get(tid)
        if gi is None: continue
        seq=["O01"]+list(t.services)+["O01"]
        if not all(n in node_xy for n in seq): continue
        xs=[node_xy[n][0] for n in seq]; ys=[node_xy[n][1] for n in seq]
        ax.plot(xs,ys,color=GROUP_COLORS[gi-1],alpha=0.34,linewidth=1.15,zorder=2)
    ax.scatter(*node_xy["O01"],marker="*",s=240,color=COLORS["red"],edgecolor="white",linewidth=0.8,zorder=6,label="O01")
    ax.text(*node_xy["O01"]," O01",fontsize=9,weight="bold",zorder=7)
    for gi,g in enumerate(sol.group_evals,start=1):
        pts=[node_xy[x] for x in g.services if x in node_xy]
        if not pts: continue
        ax.scatter([p[0] for p in pts],[p[1] for p in pts],s=72,color=GROUP_COLORS[gi-1],edgecolor="white",linewidth=0.8,zorder=5,label=f"任务组{gi}")
        for sid in g.services:
            if sid in node_xy:
                x,y=node_xy[sid]; ax.text(x,y,f" {sid}",fontsize=8,zorder=7)
    ax.set_xlabel("经度"); ax.set_ylabel("纬度")
    ax.set_title(f"K={sol.K} 主方案固定运输路线与空间分区")
    ax.legend(frameon=False); ax.grid(alpha=0.12)
    reg.save(fig,stem,f"K={sol.K}固定运输路线与空间分区",
             "在30 m DEM背景上按任务组颜色叠加问题三已经固定的25个运输架次路线；任何路线均未因问题四重新优化。",
             "直观验证同一多点运输架次涉及的服务区位于同一组，同时展示空间分区与既有运输网络的关系。","正文强烈推荐")


def plot_geo_routes_relays(reg: FigureRegistry, sol: PartitionEval, data: Q3Data,
                           node_xy: Dict[str,Tuple[float,float]], stem: str,
                           dem: Optional[np.ndarray]=None, latv: Optional[np.ndarray]=None,
                           lonv: Optional[np.ndarray]=None) -> None:
    """综合绘制固定运输路线、服务区与按组复制的原中继悬停点。"""
    if not node_xy or "O01" not in node_xy:
        return
    fig,ax=plt.subplots(figsize=(10.0,7.8))
    _draw_dem_background(ax,dem,latv,lonv,0.24)
    trip_group={tid:gi for gi,g in enumerate(sol.group_evals,start=1) for tid in g.transport_ids}
    for tid,t in data.transports.items():
        gi=trip_group.get(tid)
        if gi is None: continue
        seq=["O01"]+list(t.services)+["O01"]
        if not all(n in node_xy for n in seq): continue
        xs=[node_xy[n][0] for n in seq]; ys=[node_xy[n][1] for n in seq]
        ax.plot(xs,ys,color=GROUP_COLORS[gi-1],alpha=0.26,linewidth=1.0,zorder=2)
    ax.scatter(*node_xy["O01"],marker="*",s=240,color=COLORS["red"],edgecolor="white",linewidth=0.8,zorder=7,label="O01")
    for gi,g in enumerate(sol.group_evals,start=1):
        pts=[node_xy[x] for x in g.services if x in node_xy]
        if pts:
            ax.scatter([p[0] for p in pts],[p[1] for p in pts],s=62,color=GROUP_COLORS[gi-1],edgecolor="white",linewidth=0.7,zorder=5,label=f"组{gi}服务区")
        for sid in g.services:
            if sid in node_xy:
                x,y=node_xy[sid]; ax.text(x,y,f" {sid}",fontsize=7.5,zorder=8)
        # 同一原中继任务若跨组，则在相同经纬度出现多个组颜色的空心三角，表示按组复制而非移动悬停点。
        for rid in g.relay_ids:
            r=data.relays[rid]
            ax.scatter(r.lon,r.lat,marker="^",s=86,facecolors="none",edgecolors=GROUP_COLORS[gi-1],linewidths=1.7,zorder=6)
    relay_proxy=Line2D([0],[0],marker='^',color='none',markerfacecolor='none',markeredgecolor=COLORS['dark'],markersize=8,label='原中继悬停位置/按组副本')
    handles,labels=ax.get_legend_handles_labels(); handles.append(relay_proxy); labels.append('原中继悬停位置/按组副本')
    ax.legend(handles,labels,frameon=False,loc="best")
    ax.set_xlabel("经度"); ax.set_ylabel("纬度")
    ax.set_title(f"K={sol.K} 运输—分区—中继复制空间综合图")
    ax.grid(alpha=0.12)
    reg.save(fig,stem,f"K={sol.K}运输分区与中继复制空间综合图",
             "在DEM背景上同时叠加固定运输路线、任务组服务区和各组所需的原中继悬停位置；空心三角表示中继任务按组完整复制。",
             "强调复制只增加执行资源，不移动中继位置、不改变海拔和服务时段，因此问题三链路几何与链路预算可直接继承。","正文强烈推荐")


def plot_group_geo_detail(reg: FigureRegistry, sol: PartitionEval, group_index: int, data: Q3Data,
                          node_xy: Dict[str,Tuple[float,float]], stem: str,
                          dem: Optional[np.ndarray]=None, latv: Optional[np.ndarray]=None,
                          lonv: Optional[np.ndarray]=None) -> None:
    """单独放大一个任务组，便于查看其固定运输路线与中继悬停点。"""
    if not node_xy or "O01" not in node_xy or group_index<1 or group_index>sol.K:
        return
    g=sol.group_evals[group_index-1]
    fig,ax=plt.subplots(figsize=(8.8,7.0))
    _draw_dem_background(ax,dem,latv,lonv,0.28)
    for tid in g.transport_ids:
        t=data.transports[tid]
        seq=["O01"]+list(t.services)+["O01"]
        if all(n in node_xy for n in seq):
            xs=[node_xy[n][0] for n in seq]; ys=[node_xy[n][1] for n in seq]
            ax.plot(xs,ys,color=GROUP_COLORS[group_index-1],alpha=0.42,linewidth=1.25,zorder=2)
    ax.scatter(*node_xy["O01"],marker="*",s=230,color=COLORS["red"],edgecolor="white",linewidth=0.8,zorder=6)
    ax.text(*node_xy["O01"]," O01",fontsize=9,weight="bold")
    for sid in g.services:
        if sid in node_xy:
            x,y=node_xy[sid]; ax.scatter(x,y,s=78,color=GROUP_COLORS[group_index-1],edgecolor="white",linewidth=0.8,zorder=5); ax.text(x,y,f" {sid}",fontsize=8)
    for rid in g.relay_ids:
        r=data.relays[rid]
        ax.scatter(r.lon,r.lat,marker="^",s=90,facecolors="none",edgecolors=COLORS["teal"],linewidths=1.6,zorder=6)
        ax.text(r.lon,r.lat,f" {rid}",fontsize=6.5,color=COLORS["dark"])
    ax.set_xlabel("经度"); ax.set_ylabel("纬度")
    ax.set_title(f"K={sol.K} 主方案任务组{group_index}：运输路线与中继任务")
    ax.grid(alpha=0.12)
    reg.save(fig,stem,f"K={sol.K}任务组{group_index}空间执行细节",
             f"单独放大任务组{group_index}，显示其服务区、固定运输路线以及需要保留/复制的原中继悬停位置。",
             "用于逐组解释资源配置来源，尤其适合附录核查某组为何需要特定数量的中继无人机和能源组件。","附录/备选")

def plot_flowchart(reg: FigureRegistry) -> None:
    fig,ax=plt.subplots(figsize=(12.0,3.7));ax.axis("off")
    boxes=[
        (0.03,"固定Q3最终输出"),(0.20,"构造9个运输单元"),(0.38,"建立中继关联超图"),
        (0.56,"严格可分性检查"),(0.72,"2/3组完整枚举"),(0.87,"区间资源核算\n+主方案筛选")]
    for i,(x,txt) in enumerate(boxes):
        w=0.13 if i<5 else 0.11
        rect=FancyBboxPatch((x,0.38),w,0.28,boxstyle="round,pad=0.02",facecolor=COLORS["light"],edgecolor=COLORS["blue"],linewidth=1.6)
        ax.add_patch(rect);ax.text(x+w/2,0.52,txt,ha="center",va="center",fontsize=9)
        if i<len(boxes)-1:
            nx=boxes[i+1][0]
            ax.annotate("",xy=(nx,0.52),xytext=(x+w,0.52),arrowprops=dict(arrowstyle="->",color=COLORS["gray"],lw=1.5))
    ax.set_title("问题四固定任务下的分区—资源配置求解链路",fontsize=13,weight="bold")
    reg.save(fig,"fig_q4_01_solution_flow","问题四总体求解流程",
             "从问题三最终输出出发，依次完成任务压缩、超图建模、严格冲突判定、完整枚举和区间资源配置。",
             "概括问题四与问题三的数据继承关系，并说明本题没有重新优化运输路线或通信位置。", "正文强烈推荐")


def plot_scenario_flow(reg: FigureRegistry) -> None:
    fig,ax=plt.subplots(figsize=(10.5,4.8));ax.axis("off")
    ax.text(0.5,0.90,"同一份问题三最终联合调度",ha="center",va="center",fontsize=12,weight="bold",
            bbox=dict(boxstyle="round,pad=0.4",fc=COLORS["light"],ec=COLORS["blue"]))
    ax.annotate("",xy=(0.28,0.70),xytext=(0.48,0.84),arrowprops=dict(arrowstyle="->",lw=1.5,color=COLORS["gray"]))
    ax.annotate("",xy=(0.72,0.70),xytext=(0.52,0.84),arrowprops=dict(arrowstyle="->",lw=1.5,color=COLORS["gray"]))
    ax.text(0.25,0.60,"严格继承情景\n原中继架次不可复制\n$\\lambda_q=1$",ha="center",va="center",fontsize=10,
            bbox=dict(boxstyle="round,pad=0.5",fc="#F4ECEB",ec=COLORS["red"]))
    ax.text(0.75,0.60,"独立配置扩展情景\n跨组中继按组完整复制\n$\\lambda_q\\geq 1$",ha="center",va="center",fontsize=10,
            bbox=dict(boxstyle="round,pad=0.5",fc="#E9F4F1",ec=COLORS["teal"]))
    ax.text(0.25,0.25,"9个运输单元经中继关联后\n形成1个连通分量\nK=2、3结构性不可行",ha="center",va="center",fontsize=9)
    ax.text(0.75,0.25,"枚举合法运输分区\n复制必要中继任务\n核算独立资源需求与缺口",ha="center",va="center",fontsize=9)
    ax.set_title("问题四两种中继继承口径",fontsize=13,weight="bold")
    reg.save(fig,"fig_q4_02_scenario_definition","严格继承与按组复制两种情景",
             "明确区分题面字面严格继承和用于完成2组/3组资源核算的中继任务完整复制扩展口径。",
             "避免把结构性不可行与库存不足混为一谈，也避免用增加设备数量掩盖原中继架次跨组共享问题。", "正文强烈推荐")


def plot_unit_basic_figures(reg: FigureRegistry, units: Sequence[Tuple[str,...]], data: Q3Data, trip_to_unit: Dict[str,int], relay_to_units: Dict[str,Tuple[int,...]]) -> None:
    labels=[f"U{i+1}" for i in range(len(units))]
    bar_chart(reg,"fig_q4_03_unit_service_count",labels,[len(u) for u in units],"各运输不可拆分单元包含的服务区数量","服务区数",
              "按同一运输架次多服务区同组约束压缩后，统计每个单元包含的服务区数量。",
              "验证15个服务区被压缩为9个不可拆分单元，并识别由多点运输架次形成的成对单元。","正文备选")
    trip_count=[sum(1 for x in trip_to_unit.values() if x==i) for i in range(len(units))]
    bar_chart(reg,"fig_q4_04_unit_trip_count",labels,trip_count,"各运输不可拆分单元承担的运输架次数","运输架次",
              "统计问题三25个运输架次在9个不可拆分单元上的分布。",
              "服务区数量相同的单元可能承担完全不同的运输频次，说明分区均衡不能只看服务区个数。","备选")
    workloads=[]
    for i in range(len(units)):
        workloads.append(sum(t.duration for tid,t in data.transports.items() if trip_to_unit[tid]==i)/3600)
    bar_chart(reg,"fig_q4_05_unit_transport_workload",labels,workloads,"各运输不可拆分单元的运输作业工作量","累计运输作业时间/h",
              "将属于同一单元的固定运输架次持续时间累加。",
              "为组间工作量均衡指标提供最小粒度的解释基础。","正文备选")
    degree=[sum(1 for edge in relay_to_units.values() if i in edge) for i in range(len(units))]
    bar_chart(reg,"fig_q4_10_unit_relay_degree",labels,degree,"各运输单元关联的实际中继架次数","中继关联度",
              "统计每个运输单元实际使用过多少个问题三中继架次。",
              "高关联度单元更容易在分区时引起中继任务复制，是通信资源压力的重要来源。","备选")
    rids=sorted(relay_to_units)
    edge_sizes=[len(relay_to_units[r]) for r in rids]
    bar_chart(reg,"fig_q4_11_relay_hyperedge_size",rids,edge_sizes,"各中继架次连接的运输单元数量","超边规模",
              "以实际通信保障关系统计每个中继架次同时关联多少个运输不可拆分单元。",
              "规模大于1的超边是潜在跨组复制源；规模为1的中继任务不会因该超边自身导致额外复制。","正文备选",horizontal=True)


def plot_partition_distribution_suite(reg: FigureRegistry, results: Dict[int,List[PartitionEval]], selected: Dict[int,PartitionEval]) -> None:
    for K in (2,3):
        arr=results[K];sel=selected[K]
        hist_chart(reg,f"fig_q4_{12 if K==2 else 13:02d}_K{K}_relay_drone_hist",
                   [r.totals["relay_drone"] for r in arr],f"K={K} 全部分区的中继无人机总需求分布","中继无人机总配置",
                   f"对K={K}的全部{len(arr)}个无标签分区统计独立配置所需中继无人机总数量。",
                   "展示中继资源需求不是单一数值，而是由分区方式决定的离散分布；主方案可置于整体分布中理解。","正文备选",True)
        hist_chart(reg,f"fig_q4_{14 if K==2 else 15:02d}_K{K}_relay_energy_hist",
                   [r.totals["relay_energy"] for r in arr],f"K={K} 全部分区的中继能源组件需求分布","中继能源组件总配置",
                   f"对K={K}全部候选统计中继能源组件独立配置规模。",
                   "显示能源组件缺口与中继无人机缺口并非完全同步，应分类核算。","备选",True)
        hist_chart(reg,f"fig_q4_{16 if K==2 else 17:02d}_K{K}_copy_hist",
                   [r.relay_copies for r in arr],f"K={K} 中继任务新增副本数分布","新增中继任务副本数",
                   "按超图connectivity-1口径统计每个分区需要增加的中继任务副本数。",
                   "刻画分区对原中继共享关系的破坏程度，是比单纯设备数量更直接的结构代价。","正文推荐",True)
        pareto=[(r.relay_copies,r.cv_transport) for r in arr if r.pareto]
        scatter_chart(reg,f"fig_q4_{18 if K==2 else 19:02d}_K{K}_copy_balance_scatter",
                      [r.relay_copies for r in arr],[r.cv_transport for r in arr],f"K={K} 中继复制代价—运输工作量均衡关系",
                      "新增中继任务副本数","运输工作量CV",
                      "每个点对应一个完整分区；横轴为中继复制架次数，纵轴为固定运输工作量CV。",
                      "直观显示减少跨组中继切割与提高工作量均衡之间的真实权衡。","正文强烈推荐",
                      (sel.relay_copies,sel.cv_transport),pareto)
        pareto2=[(r.totals["relay_drone"],r.cv_transport) for r in arr if r.pareto]
        scatter_chart(reg,f"fig_q4_{20 if K==2 else 21:02d}_K{K}_relay_balance_scatter",
                      [r.totals["relay_drone"] for r in arr],[r.cv_transport for r in arr],f"K={K} 中继无人机配置—运输工作量均衡关系",
                      "中继无人机总配置","运输工作量CV",
                      "将实体中继资源最少配置与运输工作量均衡指标直接关联。",
                      "说明理论最少中继方案不一定是主提交方案；主方案还需兼顾运输侧零增配与均衡。","正文推荐",
                      (sel.totals["relay_drone"],sel.cv_transport),pareto2)
        scatter_chart(reg,f"fig_q4_{22 if K==2 else 23:02d}_K{K}_shortage_balance_scatter",
                      [r.total_shortage for r in arr],[r.cv_transport for r in arr],f"K={K} 总资源缺口—运输工作量均衡关系",
                      "相对现有库存的分类缺口数量之和","运输工作量CV",
                      "按资源类别分别计算缺口后求和，仅用于散点横轴，不把不同资源成本等价化。",
                      "用于识别库存可行性与工作量均衡之间的关系；正式结论仍应回到分类缺口表。","备选",
                      (sel.total_shortage,sel.cv_transport),None)
        scatter_chart(reg,f"fig_q4_{24 if K==2 else 25:02d}_K{K}_copy_energy_scatter",
                      [r.relay_copies for r in arr],[r.relay_delta_energy for r in arr],f"K={K} 中继副本数量—新增能耗关系",
                      "新增中继副本数","新增中继能耗/kWh",
                      "每个点对应一个分区，新增能耗按被复制原中继架次的实际能耗累加。",
                      "相同复制次数仍可能对应不同新增能耗，说明应使用实际架次能耗而不是统一副本惩罚系数。","备选",
                      (sel.relay_copies,sel.relay_delta_energy),None)


def plot_inventory_utilization(reg: FigureRegistry, sol: PartitionEval, inventory: Dict[str,int], stem: str) -> None:
    labels=[RESOURCE_CN[k] for k in RESOURCE_ORDER]
    vals=[100*sol.totals[k]/inventory[k] if inventory[k]>0 else 0 for k in RESOURCE_ORDER]
    bar_chart(reg,stem,labels,vals,f"K={sol.K} 主方案库存利用率","需求/库存（%）",
              "按8类资源计算主方案总需求与现有库存的比例。",
              "超过100%的资源就是必须增配的类别；同时可观察未形成缺口的资源冗余程度。","正文备选")


def plot_group_box_trip_counts(reg: FigureRegistry, sol: PartitionEval, stem: str) -> None:
    labels=[f"组{i+1}" for i in range(sol.K)]
    x=np.arange(sol.K);w=0.27
    fig,ax=plt.subplots(figsize=(8.6,5.2))
    ax.bar(x-w,[g.service_count for g in sol.group_evals],w,label="服务区数",color=COLORS["gray"])
    ax.bar(x,[g.transport_trip_count for g in sol.group_evals],w,label="运输架次",color=COLORS["blue"])
    ax.bar(x+w,[g.box_count for g in sol.group_evals],w,label="货箱数",color=COLORS["teal"])
    ax.set_xticks(x);ax.set_xticklabels(labels);ax.set_title(f"K={sol.K} 各组任务规模的多口径比较")
    ax.legend(frameon=False);ax.grid(axis="y",alpha=0.18)
    reg.save(fig,stem,f"K={sol.K}各组任务规模比较",
             "同时比较各组服务区数、固定运输架次数和货箱数。",
             "强调“服务区数量均衡”不等于“实际工作量均衡”，为采用累计任务时间CV提供解释。","正文备选")


def plot_redundancy_decomposition(reg: FigureRegistry, selected: Dict[int,PartitionEval], pool: Dict[str,int]) -> None:
    labels=["K=2","K=3"]
    copy_inc=[]; isolation_inc=[]; total_inc=[]
    # 对中继无人机做归因：副本后的假想共享池需求需要单独重建，这里用可计算下界近似不可取。
    # 因此本图只做“总资源冗余”按运输侧/通信侧分类，避免虚构精确分解。
    trans=[]; comm=[]
    for K in (2,3):
        s=selected[K]
        trans.append(sum(s.totals[k]-pool[k] for k in TRANSPORT_RESOURCE_KEYS))
        comm.append(sum(s.totals[k]-pool[k] for k in ["relay_drone","relay_energy"]))
    x=np.arange(2)
    fig,ax=plt.subplots(figsize=(7.6,5.2))
    ax.bar(x,trans,label="运输侧独立配置冗余",color=COLORS["blue"])
    ax.bar(x,comm,bottom=trans,label="通信侧独立配置冗余",color=COLORS["teal"])
    ax.set_xticks(x);ax.set_xticklabels(labels);ax.set_ylabel("相对统一共享基准的资源数量增量")
    ax.set_title("2组与3组主方案的资源冗余来源")
    ax.legend(frameon=False);ax.grid(axis="y",alpha=0.18)
    reg.save(fig,"fig_q4_38_redundancy_source","主方案资源冗余来源",
             "将相对统一共享最少需求增加的资源，按运输侧和通信侧两大类汇总展示。",
             "主方案优先维持运输侧零增配，因此新增配置主要集中在共享中继被拆分后的通信侧。","正文推荐")


def generate_all_figures(
    outdir: Path,
    units: Sequence[Tuple[str,...]],
    data: Q3Data,
    trip_to_unit: Dict[str,int],
    relay_to_units: Dict[str,Tuple[int,...]],
    results: Dict[int,List[PartitionEval]],
    selected: Dict[int,PartitionEval],
    pool: Dict[str,int],
    inventory: Dict[str,int],
    node_xy: Dict[str,Tuple[float,float]],
    dem: Optional[np.ndarray], latv: Optional[np.ndarray], lonv: Optional[np.ndarray],
    save_pdf: bool,
) -> FigureRegistry:
    reg=FigureRegistry(outdir,save_pdf=save_pdf)
    plot_flowchart(reg)
    plot_scenario_flow(reg)
    plot_unit_basic_figures(reg,units,data,trip_to_unit,relay_to_units)
    plot_incidence_matrix(reg,units,relay_to_units)
    plot_bipartite_hypergraph(reg,units,relay_to_units)
    plot_strict_component_graph(reg,len(units),relay_to_units)
    plot_partition_counts(reg,{2:len(results[2]),3:len(results[3])})
    plot_partition_distribution_suite(reg,results,selected)

    # 主方案图
    for K in (2,3):
        sol=selected[K]
        base=26 if K==2 else 27
        plot_resource_compare(reg,sol,inventory,f"fig_q4_{base:02d}_K{K}_resource_vs_stock")
        plot_group_resource_heatmap(reg,sol,f"fig_q4_{28 if K==2 else 29:02d}_K{K}_group_resource_heatmap")
        plot_group_workload(reg,sol,f"fig_q4_{30 if K==2 else 31:02d}_K{K}_group_workload")
        plot_selected_membership(reg,sol,units,f"fig_q4_{34 if K==2 else 35:02d}_K{K}_service_membership")
        plot_replication_matrix(reg,sol,relay_to_units,f"fig_q4_{36 if K==2 else 37:02d}_K{K}_relay_replication")
        plot_relay_copy_energy(reg,sol,data,relay_to_units,f"fig_q4_{39 if K==2 else 40:02d}_K{K}_copy_energy")
        plot_gantt_transport(reg,sol,data,f"fig_q4_{41 if K==2 else 42:02d}_K{K}_transport_gantt")
        plot_gantt_relay(reg,sol,data,f"fig_q4_{43 if K==2 else 44:02d}_K{K}_relay_gantt")
        plot_balance_sensitivity(reg,results[K],K,f"fig_q4_{45 if K==2 else 46:02d}_K{K}_balance_sensitivity")
        plot_top_candidates(reg,results[K],K,f"fig_q4_{47 if K==2 else 48:02d}_K{K}_top_candidates")
        plot_inventory_utilization(reg,sol,inventory,f"fig_q4_{49 if K==2 else 50:02d}_K{K}_inventory_utilization")
        plot_group_box_trip_counts(reg,sol,f"fig_q4_{51 if K==2 else 52:02d}_K{K}_group_task_scale")
        if node_xy:
            plot_geo_partition(reg,sol,node_xy,f"fig_q4_{53 if K==2 else 54:02d}_K{K}_geo_partition",dem,latv,lonv)
            plot_relay_position_map(reg,sol,data,node_xy,f"fig_q4_{55 if K==2 else 56:02d}_K{K}_relay_positions")
            plot_geo_routes(reg,sol,data,node_xy,f"fig_q4_{57 if K==2 else 58:02d}_K{K}_geo_routes",dem,latv,lonv)
            plot_geo_routes_relays(reg,sol,data,node_xy,f"fig_q4_{59 if K==2 else 60:02d}_K{K}_geo_routes_relays",dem,latv,lonv)
            if K==2:
                plot_group_geo_detail(reg,sol,1,data,node_xy,"fig_q4_61_K2_group1_geo_detail",dem,latv,lonv)
                plot_group_geo_detail(reg,sol,2,data,node_xy,"fig_q4_62_K2_group2_geo_detail",dem,latv,lonv)
            else:
                plot_group_geo_detail(reg,sol,1,data,node_xy,"fig_q4_63_K3_group1_geo_detail",dem,latv,lonv)
                plot_group_geo_detail(reg,sol,2,data,node_xy,"fig_q4_64_K3_group2_geo_detail",dem,latv,lonv)
                plot_group_geo_detail(reg,sol,3,data,node_xy,"fig_q4_65_K3_group3_geo_detail",dem,latv,lonv)

    plot_resource_redundancy_compare(reg,selected,pool)
    plot_pooled_vs_original_ids(reg,data,pool)
    plot_redundancy_decomposition(reg,selected,pool)
    return reg


# ============================================================
# 9. TeX/说明文件输出
# ============================================================

def tex_escape(s: str) -> str:
    rep={"\\":"\\textbackslash{}","_":"\\_","%":"\\%","&":"\\&","#":"\\#","$":"\\$","{":"\\{","}":"\\}"}
    return "".join(rep.get(ch,ch) for ch in s)


def write_figure_manifest(path: Path, figs: Sequence[FigureInfo]) -> None:
    """输出实际成功生成的图表清单；数量以磁盘文件为准。"""
    def _num(f: FigureInfo) -> int:
        m = re.search(r"fig_q4_(\d+)", f.filename_png)
        return int(m.group(1)) if m else 10**9
    figs = sorted(figs, key=lambda f: (_num(f), f.filename_png))
    with path.open("w", newline="", encoding="utf-8-sig") as fp:
        w = csv.writer(fp)
        w.writerow(["序号", "PNG文件", "PDF文件", "标题", "正文建议", "PNG存在", "PDF存在"])
        for i, f in enumerate(figs, 1):
            png_ok = (path.parent / "figures" / f.filename_png).exists()
            pdf_ok = True if not f.filename_pdf else (path.parent / "figures" / f.filename_pdf).exists()
            w.writerow([i, f.filename_png, f.filename_pdf or "", f.title, f.recommended, int(png_ok), int(pdf_ok)])


def write_figure_tex(path: Path, figs: Sequence[FigureInfo]) -> None:
    lines=[
        "% 自动生成：问题四全部图表说明。",
        "% 本文件不是正文强制内容，而是用于从大量备选图中挑选论文插图。",
        "\\section*{问题四图表说明与选图建议}",
        f"本次程序实际成功生成\\textbf{{{len(figs)}}}幅不同图表。每幅图均给出图意、可支持的结论以及是否建议放入正文；未被程序实际生成的可选地理图不会列入本文件。",
        "图表数量以\\texttt{Q4\\_图表清单.csv}和磁盘中实际存在的文件为准，不再使用预先写死的数量。",
        "",
    ]
    def _fig_num(f):
        m = re.search(r"fig_q4_(\d+)", f.filename_png)
        return int(m.group(1)) if m else 10**9
    figs = sorted(figs, key=lambda f: (_fig_num(f), f.filename_png))
    for i,f in enumerate(figs, start=1):
        lines += [
            f"\\subsection*{{图 {i}：{tex_escape(f.title)}}}",
            f"\\textbf{{文件：}}\\texttt{{{tex_escape(f.filename_png)}}}" + (f"；矢量版：\\texttt{{{tex_escape(f.filename_pdf)}}}" if f.filename_pdf else "") + "\\\\",
            f"\\textbf{{图中内容：}}{tex_escape(f.description)}\\\\",
            f"\\textbf{{主要意义：}}{tex_escape(f.meaning)}\\\\",
            f"\\textbf{{正文建议：}}{tex_escape(f.recommended)}",
            "",
        ]
    path.write_text("\n".join(lines), encoding="utf-8")


def write_reference_tex(path: Path) -> None:
    txt=r"""% 问题四建议参考文献条目（题名、作者、卷页与 DOI 已核对）
% 可直接并入主文档 thebibliography；正文使用下列 cite key。

\bibitem{q4-eyubov2025}
K. Eyubov, M. Fonseca Faraj, and C. Schulz,
``FREIGHT: Fast Streaming Hypergraph Partitioning,''
\emph{Algorithmica}, vol. 87, pp. 405--428, 2025,
doi: 10.1007/s00453-024-01291-8.
% 本文用途：超图表示、高阶共享关系及 connectivity / lambda-1 指标。
% 注意：本题只借鉴其超图建模与指标，不采用 FREIGHT 流式求解算法。

\bibitem{q4-krause2025}
R. Krause, L. Gottesb\"uren, and N. Maas,
``Deterministic Parallel High-Quality Hypergraph Partitioning,''
in \emph{2025 Proceedings of the Conference on Applied and Computational Discrete Algorithms (ACDA)},
pp. 222--236, 2025,
doi: 10.1137/1.9781611979084.17.
% 本文用途：平衡超图划分与确定性/可复现求解背景。
% 本题规模仅9个运输单元，实际采用完整枚举，不宣称采用其并行多层算法。

\bibitem{q4-kleinberg2006}
J. Kleinberg and \`E. Tardos,
\emph{Algorithm Design}. Boston, MA, USA: Pearson/Addison-Wesley, 2006, Ch. 4.
% 本文用途：Interval Partitioning；固定时间区间下，同质资源的最少数量等于最大重叠深度，
% 按开始时刻并优先复用最早释放资源的贪心算法达到该下界。

\bibitem{q4-park2025}
J. Park, G. Noh, C. Park, J. Kim, J. Kim, D. Lee, and D. Cho,
``Development of mission allocation based on MILP for multi-UAVs with limited resources,''
\emph{Aerospace Science and Technology}, vol. 166, Art. no. 110598, 2025,
doi: 10.1016/j.ast.2025.110598.
% 本文用途：有限资源条件下多无人机任务分配的研究背景；强调资源约束与任务属性必须显式核算。
% 本题不采用其 MILP 求解器，因为问题三时刻表已固定且Q4只剩9个不可拆分单元。

\bibitem{q4-chu2025}
J. Chu, F. Yan, and Z. Xu,
``Communication-constrained multi-UAV task allocation method for non-independent tasks,''
\emph{Engineering Science and Technology, an International Journal}, vol. 70, Art. no. 102166, 2025,
doi: 10.1016/j.jestch.2025.102166.
% 本文用途：通信约束与任务依赖共同影响多无人机任务分配的研究背景。
% 本题中的依赖关系来自问题三既定运输架次与实际中继保障关系。

\bibitem{q4-zhang2026}
L. Zhang, W. Chen, and J. Chang,
``A search and rescue task allocation algorithm for UAV swarms based on dynamic power management,''
\emph{Journal of Information and Intelligence}, 2026, in press,
doi: 10.1016/j.jiixd.2026.05.003.
% 本文用途：灾后搜救中任务时效、能源与物资资源协同配置的近期研究背景。
"""
    path.write_text(txt, encoding="utf-8")


def macro_name(prefix: str, key: str) -> str:
    s=re.sub(r"[^A-Za-z0-9]","",key)
    return prefix+s


def write_q4_macros(path: Path, units: Sequence[Tuple[str,...]], strict_comps: Sequence[Tuple[int,...]],
                    results: Dict[int,List[PartitionEval]], selected: Dict[int,PartitionEval], theoretical: Dict[int,PartitionEval],
                    pool: Dict[str,int], inventory: Dict[str,int]) -> None:
    lines=["% 自动生成：问题四正文数值宏。后续 body.tex 可直接 \\input 本文件。"]
    lines.append(f"\\newcommand{{\\QFourUnitNum}}{{{len(units)}}}")
    lines.append(f"\\newcommand{{\\QFourStrictComponentNum}}{{{len(strict_comps)}}}")
    lines.append(f"\\newcommand{{\\QFourKTwoPartitionNum}}{{{len(results[2])}}}")
    lines.append(f"\\newcommand{{\\QFourKThreePartitionNum}}{{{len(results[3])}}}")
    for K,tag in [(2,"Two"),(3,"Three")]:
        s=selected[K];t=theoretical[K]
        lines.append(f"\\newcommand{{\\QFour{tag}RelayMin}}{{{t.totals['relay_drone']}}}")
        lines.append(f"\\newcommand{{\\QFour{tag}RelayMain}}{{{s.totals['relay_drone']}}}")
        lines.append(f"\\newcommand{{\\QFour{tag}EnergyMain}}{{{s.totals['relay_energy']}}}")
        lines.append(f"\\newcommand{{\\QFour{tag}Copy}}{{{s.relay_copies}}}")
        lines.append(f"\\newcommand{{\\QFour{tag}DeltaRelayEnergy}}{{{s.relay_delta_energy:.3f}}}")
        lines.append(f"\\newcommand{{\\QFour{tag}CV}}{{{s.cv_transport:.3f}}}")
        lines.append(f"\\newcommand{{\\QFour{tag}TotalShortage}}{{{s.total_shortage}}}")
        lines.append(f"\\newcommand{{\\QFour{tag}ZeroTransportCount}}{{{sum(1 for r in results[K] if r.transport_zero_augmentation)}}}")
        for gi,g in enumerate(s.group_evals,start=1):
            lines.append(f"\\newcommand{{\\QFour{tag}Group{gi}Services}}{{{','.join(g.services)}}}")
    path.write_text("\n".join(lines)+"\n",encoding="utf-8")


# ============================================================
# 10. 结果输出、清单与报告
# ============================================================

def write_input_manifest(path: Path, files: Dict[str,Optional[Path]]) -> None:
    rows=[]
    for role,p in files.items():
        if p is not None and p.exists():
            rows.append([role,str(p),p.stat().st_size,sha256_file(p),"FOUND"])
        else:
            rows.append([role,"","","","NOT_FOUND_OR_OPTIONAL"])
    write_csv(path,["role","path","bytes","sha256","status"],rows)


def write_units_csv(path: Path, units: Sequence[Tuple[str,...]], data: Q3Data, trip_to_unit: Dict[str,int], relay_to_units: Dict[str,Tuple[int,...]]) -> None:
    rows=[]
    for i,u in enumerate(units):
        tids=sorted(t for t,x in trip_to_unit.items() if x==i)
        rids=sorted(r for r,e in relay_to_units.items() if i in e)
        rows.append([f"U{i+1}",",".join(u),len(u),",".join(tids),len(tids),",".join(rids),len(rids),
                     sum(data.transports[t].duration for t in tids),sum(data.transports[t].energy_kwh for t in tids)])
    write_csv(path,["unit","services","service_count","transport_trips","transport_trip_count","relay_trips","relay_degree","transport_workload_s","transport_energy_kWh"],rows)


def write_hyperedges_csv(path: Path, relay_to_units: Dict[str,Tuple[int,...]], relay_to_trips: Dict[str,Tuple[str,...]], data: Q3Data) -> None:
    rows=[]
    for rid in sorted(data.relays):
        edge=relay_to_units.get(rid,())
        rows.append([rid,",".join(relay_to_trips.get(rid,())),",".join(f"U{u+1}" for u in edge),len(edge),data.relays[rid].energy_kwh])
    write_csv(path,["relay_trip","actual_transport_trips","unit_hyperedge","edge_size","relay_energy_kWh"],rows)


def write_selected_config_csv(path: Path, selected: Dict[int,PartitionEval]) -> None:
    rows=[]
    for K in (2,3):
        for gi,g in enumerate(selected[K].group_evals,start=1):
            rows.append([K,gi,",".join(g.services),g.resources["A_drone"],g.resources["B_drone"],g.resources["C_drone"],
                         g.resources["A_battery"],g.resources["B_battery"],g.resources["C_battery"],g.resources["relay_drone"],g.resources["relay_energy"]])
    write_csv(path,["K（2或3）","任务组编号","服务区列表","A型运输无人机数","B型运输无人机数","C型运输无人机数",
                    "A型电池组数","B型电池组数","C型电池组数","中继无人机数","中继能源组件数"],rows)


def write_assignments(outdir: Path, K: int, tr_rows: List[Dict[str,Any]], re_rows: List[Dict[str,Any]]) -> None:
    if tr_rows:
        keys=list(tr_rows[0].keys())
        write_csv(outdir/f"Q4_K{K}_运输资源重新编号.csv",keys,([r[k] for k in keys] for r in tr_rows))
    if re_rows:
        keys=list(re_rows[0].keys())
        write_csv(outdir/f"Q4_K{K}_中继副本与资源重新编号.csv",keys,([r[k] for k in keys] for r in re_rows))


def write_audit_report(path: Path, data: Q3Data, audits: Dict[str,Any], units: Sequence[Tuple[str,...]], strict_comps_list: Sequence[Tuple[int,...]],
                       selected: Dict[int,PartitionEval], theoretical: Dict[int,PartitionEval], inventory: Dict[str,int], pool: Dict[str,int],
                       resource_issues: Dict[int,List[str]]) -> None:
    lines=[]
    lines.append("问题四完整审计报告")
    lines.append("="*72)
    lines.append(f"Q3详细结果: {data.q3_detail_path}")
    lines.append(f"运输架次数: {audits['transport_count']}")
    lines.append(f"中继架次数: {audits['relay_count']}")
    lines.append(f"逐箱记录数: {audits['delivery_count']}，唯一货箱: {audits['unique_box_count']}")
    lines.append(f"硬截止违反数: {audits['hard_deadline_violations']}")
    lines.append(f"运输最低返航SOC: {100*audits['transport_min_soc']:.2f}%")
    lines.append(f"中继最低返航SOC: {100*audits['relay_min_soc']:.2f}%")
    lines.append(f"通信导出区间超过中继服务结束的记录数: {audits['comm_interval_overrun_count']}")
    for x in audits.get("comm_interval_overrun_records",[]):
        lines.append(f"  {x[0]}/{x[1]}: 导出结束={x[2]:.3f}, 中继服务结束={x[3]:.3f}, 越界={x[4]:.3f}s")
    lines.append("")
    lines.append("运输不可拆分单元：")
    for i,u in enumerate(units,start=1):
        lines.append(f"  U{i}: {','.join(u)}")
    lines.append(f"严格继承口径中继关联连通分量数: {len(strict_comps_list)}")
    lines.append("  " + " | ".join("{"+",".join(f"U{u+1}" for u in c)+"}" for c in strict_comps_list))
    if len(strict_comps_list)<2:
        lines.append("结论：严格保持原中继架次不可复制且资源不得跨组时，K=2与K=3均结构性不可行。")
    lines.append("")
    lines.append("固定问题三时刻表统一共享最少资源 vs 库存：")
    for k in RESOURCE_ORDER:
        lines.append(f"  {RESOURCE_CN[k]}: pooled={pool[k]}, inventory={inventory[k]}")
    lines.append("")
    for K in (2,3):
        s=selected[K]; t=theoretical[K]
        lines.append(f"K={K} 理论最少中继无人机方案: relay_drone={t.totals['relay_drone']}, relay_energy={t.totals['relay_energy']}, copies={t.relay_copies}, CV_T={t.cv_transport:.6f}")
        lines.append(f"K={K} 主提交方案: relay_drone={s.totals['relay_drone']}, relay_energy={s.totals['relay_energy']}, copies={s.relay_copies}, CV_T={s.cv_transport:.6f}, ΔE_R={s.relay_delta_energy:.6f}kWh")
        for gi,g in enumerate(s.group_evals,start=1):
            lines.append(f"  组{gi}: {','.join(g.services)} | " + ", ".join(f"{RESOURCE_CN[r]}={g.resources[r]}" for r in RESOURCE_ORDER))
        lines.append("  总需求: " + ", ".join(f"{RESOURCE_CN[r]}={s.totals[r]}" for r in RESOURCE_ORDER))
        lines.append("  缺口: " + ", ".join(f"{RESOURCE_CN[r]}={s.shortage[r]}" for r in RESOURCE_ORDER if s.shortage[r]>0) if any(s.shortage.values()) else "  缺口: 无")
        if resource_issues[K]:
            lines.append("  资源重编号冲突审计: FAILED")
            lines.extend("    "+x for x in resource_issues[K])
        else:
            lines.append("  资源重编号冲突审计: PASS")
        lines.append("")
    path.write_text("\n".join(lines),encoding="utf-8")


def jsonable_solution(s: PartitionEval) -> Dict[str,Any]:
    return {
        "K":s.K,"canonical":s.canonical,"groups":[list(g) for g in s.groups],
        "group_services":[list(g.services) for g in s.group_evals],
        "group_resources":[g.resources for g in s.group_evals],
        "totals":s.totals,"shortage":s.shortage,"redundancy_vs_pool":s.redundancy_vs_pool,
        "relay_copies":s.relay_copies,"relay_delta_energy":s.relay_delta_energy,
        "cv_transport":s.cv_transport,"cv_joint":s.cv_joint,
        "transport_zero_augmentation":s.transport_zero_augmentation,
        "all_inventory_feasible":s.all_inventory_feasible,
    }


# ============================================================
# 11. 主程序
# ============================================================

def parse_args() -> argparse.Namespace:
    ap=argparse.ArgumentParser(description="问题四：任务分区与独立资源配置完整枚举求解器")
    ap.add_argument("--root",type=Path,default=DEFAULT_ROOT,help="项目根目录")
    ap.add_argument("--q3-zip",type=Path,default=None,help="问题三 output.zip")
    ap.add_argument("--q3-detail",type=Path,default=None,help="Q3详细结果.xlsx（包含Q3_运输架次等工作表）")
    ap.add_argument("--q3-dir",type=Path,default=None,help="问题三输出目录；自动寻找详细xlsx")
    ap.add_argument("--template",type=Path,default=None,help="结果提交模板.xlsx")
    ap.add_argument("--nodes-xlsx",type=Path,default=None,help="调度中心与服务区xlsx，用于地理图")
    ap.add_argument("--dem-mat",type=Path,default=None,help="30m DEM mat，用于地形分区图")
    ap.add_argument("--output-dir",type=Path,default=None,help="问题四输出目录")
    ap.add_argument("--no-pdf",action="store_true",help="不同时保存PDF矢量图")
    return ap.parse_args()


def main() -> None:
    args=parse_args()
    root=args.root.resolve()
    print(root)
    outdir=(args.output_dir or (SCRIPT_PATH.parent/"output")).resolve()
    outdir.mkdir(parents=True,exist_ok=True)
    figdir=outdir/"figures"
    figdir.mkdir(parents=True,exist_ok=True)

    print("[1/11] 定位并读取问题三最终详细结果...")
    q3_detail=None
    extracted=None
    if args.q3_detail and args.q3_detail.exists():
        q3_detail=args.q3_detail.resolve()
    elif args.q3_zip and args.q3_zip.exists():
        extracted=extract_q3_zip(args.q3_zip.resolve(),outdir/"_q3_extracted")
        q3_detail=find_q3_detail_xlsx(extracted)
    elif args.q3_dir and args.q3_dir.exists():
        q3_detail=find_q3_detail_xlsx(args.q3_dir.resolve())
    else:
        guesses=[root/"code"/"3"/"output",root/"output",root]
        for g in guesses:
            if g.exists():
                q3_detail=find_q3_detail_xlsx(g)
                if q3_detail:break
    if q3_detail is None:
        raise FileNotFoundError("未找到包含 Q3_运输架次/Q3_中继架次/Q3_通信保障 的问题三详细结果xlsx。请指定 --q3-detail 或 --q3-zip。")
    data=load_q3_detail(q3_detail)
    audits=q3_integrity_checks(data)
    print(f"  运输{len(data.transports)}架次，中继{len(data.relays)}架次，逐箱{len(data.deliveries)}条。")

    print("[2/11] 构造运输不可拆分单元与实际中继超图...")
    units,service_to_unit,trip_to_unit=build_transport_units(data)
    relay_to_trips,relay_to_units,trip_to_relays=build_actual_relay_relation(data,trip_to_unit)
    strict_comps_list=strict_components(len(units),relay_to_units)
    print(f"  15个服务区压缩为 {len(units)} 个运输不可拆分单元；严格中继关联连通分量={len(strict_comps_list)}。")

    print("[3/11] 定位可选原始节点/DEM与全部上游文件，建立来源清单...")
    named=discover_named_inputs(root)
    node_path=find_node_xlsx(root,args.nodes_xlsx)
    dem_path=find_dem_mat(root,args.dem_mat)
    node_xy=load_node_coordinates(node_path)
    dem,latv,lonv=load_dem_grid(dem_path)
    template=find_template(root,args.template)
    manifest=dict(named)
    manifest.update({"q3_detail":q3_detail,"q3_zip":args.q3_zip.resolve() if args.q3_zip and args.q3_zip.exists() else None,
                     "template":template,"node_xlsx_used":node_path,"dem_mat_used":dem_path})
    write_input_manifest(outdir/"Q4_input_manifest.csv",manifest)
    print(f"  节点坐标: {'已读取 '+str(len(node_xy))+' 个节点' if node_xy else '未找到（不影响核心求解）'}；DEM: {'已读取' if dem is not None else '未找到（不影响核心求解）'}。")
    if not node_xy:
        print("  [提示] 未识别节点坐标：本次将不生成地理类图。可显式指定 --nodes-xlsx <调度中心与服务区.xlsx>。")
    if dem is None:
        print("  [提示] 未识别DEM：地理图若生成将无DEM背景。可显式指定 --dem-mat <30mDEM.mat>。")

    print("[4/11] 计算固定时刻表统一共享情况下的最少资源基准...")
    pool_eval=build_group_eval(tuple(range(len(units))),units,data,trip_to_unit,relay_to_units,node_xy)
    pool_resources=dict(pool_eval.resources)
    inventory=dict(DEFAULT_INVENTORY)
    print("  pooled:",pool_resources)

    print("[5/11] 完整枚举 K=2、K=3 全部分区并精确核算独立资源...")
    results:Dict[int,List[PartitionEval]]={}
    for K in (2,3):
        arr=[]
        for groups in set_partitions_k(len(units),K):
            arr.append(evaluate_partition(groups,units,data,trip_to_unit,relay_to_units,inventory,pool_resources,node_xy))
        pareto_front(arr)
        results[K]=arr
        print(f"  K={K}: {len(arr)} 个分区；库存全可行={sum(r.all_inventory_feasible for r in arr)}；运输侧零增配={sum(r.transport_zero_augmentation for r in arr)}。")

    print("[6/11] 选取理论中继下界方案与主提交方案...")
    selected={2:select_main_solution(results[2]),3:select_main_solution(results[3])}
    theoretical={2:select_theoretical_relay_min(results[2]),3:select_theoretical_relay_min(results[3])}
    for K in (2,3):
        print(f"  K={K}: theoretical relay={theoretical[K].totals['relay_drone']}；main relay={selected[K].totals['relay_drone']}，energy={selected[K].totals['relay_energy']}，copies={selected[K].relay_copies}。")

    print("[7/11] 为主方案重新分配组内实体资源并执行时间冲突审计...")
    resource_issues:Dict[int,List[str]]={}
    assignment_data={}
    for K in (2,3):
        tr_rows,re_rows=allocate_selected_resources(selected[K],data)
        issues=[]
        issues += check_no_resource_conflict(tr_rows,"new_drone","start","return")
        issues += check_no_resource_conflict(tr_rows,"new_battery","start","charge_end")
        issues += check_no_resource_conflict(re_rows,"new_relay_drone","start","drone_release")
        issues += check_no_resource_conflict(re_rows,"new_energy","start","charge_end")
        resource_issues[K]=issues
        assignment_data[K]=(tr_rows,re_rows)
        write_assignments(outdir,K,tr_rows,re_rows)
        if issues:
            raise AssertionError(f"K={K} 实体资源重编号存在冲突：\n"+"\n".join(issues[:10]))

    print("[8/11] 输出枚举结果、超图关系、主方案与详细审计文件...")
    write_units_csv(outdir/"Q4_运输不可拆分单元.csv",units,data,trip_to_unit,relay_to_units)
    write_hyperedges_csv(outdir/"Q4_中继关联超边.csv",relay_to_units,relay_to_trips,data)
    for K in (2,3):
        write_csv(outdir/f"Q4_全部分区_K{K}.csv",partition_csv_header(K),(partition_to_flat_row(r) for r in results[K]))
        front=[r for r in results[K] if r.pareto]
        write_csv(outdir/f"Q4_Pareto候选_K{K}.csv",partition_csv_header(K),(partition_to_flat_row(r) for r in front))
    write_selected_config_csv(outdir/"Q4_分区配置_提交值.csv",selected)
    write_detail_xlsx(outdir/"Q4_详细结果.xlsx",units,relay_to_units,relay_to_trips,pool_resources,inventory,selected,theoretical,audits)

    if template is not None:
        write_template_q4(template,outdir/"结果提交模板_Q4已回填.xlsx",selected)
        print("  已回填 Q4_分区配置。")
    else:
        print("  未找到结果提交模板，仅输出与模板列完全一致的 Q4_分区配置_提交值.csv。")

    print("[9/11] 生成尽可能完整的诊断图、论文备选图与矢量PDF...")
    reg=generate_all_figures(figdir,units,data,trip_to_unit,relay_to_units,results,selected,pool_resources,inventory,node_xy,dem,latv,lonv,not args.no_pdf)
    print(f"  实际生成 {len(reg.items)} 幅不同图表（每幅默认同时保存PNG与PDF）。")

    print("[10/11] 自动生成图表说明TeX、正文数值宏与参考文献条目...")
    write_figure_manifest(outdir/"Q4_图表清单.csv", reg.items)
    write_figure_tex(outdir/"Q4_图表说明.tex",reg.items)
    write_reference_tex(outdir/"Q4_参考文献.tex")
    write_q4_macros(outdir/"Q4_结果宏.tex",units,strict_comps_list,results,selected,theoretical,pool_resources,inventory)

    print("[11/11] 写入最终摘要与审计报告...")
    write_audit_report(outdir/"Q4_审计报告.txt",data,audits,units,strict_comps_list,selected,theoretical,inventory,pool_resources,resource_issues)
    summary={
        "q3_detail":str(q3_detail),
        "transport_units":[list(u) for u in units],
        "strict_components":[[u+1 for u in c] for c in strict_comps_list],
        "strict_K2_feasible":len(strict_comps_list)>=2,
        "strict_K3_feasible":len(strict_comps_list)>=3,
        "partition_counts":{"K2":len(results[2]),"K3":len(results[3])},
        "transport_zero_augmentation_counts":{"K2":sum(r.transport_zero_augmentation for r in results[2]),"K3":sum(r.transport_zero_augmentation for r in results[3])},
        "pool_resources":pool_resources,
        "inventory":inventory,
        "theoretical_relay_min":{"K2":jsonable_solution(theoretical[2]),"K3":jsonable_solution(theoretical[3])},
        "selected":{"K2":jsonable_solution(selected[2]),"K3":jsonable_solution(selected[3])},
        "figure_count":len(reg.items),
        "comm_export_overrun_count":audits["comm_interval_overrun_count"],
    }
    (outdir/"Q4_summary.json").write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding="utf-8")

    print("\n完成。核心结果：")
    for K in (2,3):
        s=selected[K]
        print(f"K={K} 主方案：")
        for gi,g in enumerate(s.group_evals,start=1):
            print(f"  组{gi}: {','.join(g.services)} | A/B/C机={g.resources['A_drone']}/{g.resources['B_drone']}/{g.resources['C_drone']} | "
                  f"A/B/C电池={g.resources['A_battery']}/{g.resources['B_battery']}/{g.resources['C_battery']} | "
                  f"中继={g.resources['relay_drone']} | 中继能源={g.resources['relay_energy']}")
        print("  总缺口:", {RESOURCE_CN[k]:v for k,v in s.shortage.items() if v>0} or "无")
    print(f"\n输出目录：{outdir}")


if __name__ == "__main__":
    main()
