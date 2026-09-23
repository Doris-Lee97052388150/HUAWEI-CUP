from __future__ import annotations
from pathlib import Path
from dataclasses import dataclass
import argparse
import re
from openpyxl import load_workbook

DEFAULT_DIR = Path(r"D:\SF_Dir\第二十三届中国研究生数学建模竞赛 - 中文题目\中文题目\D题\HUAWEI-CUP\数据\无人机应急物资运输基础数据")
GRAVITY = 9.81
EARTH_RADIUS_M = 6_371_000.0
SENSITIVITY = (0.10, 0.15, 0.20, 0.25, 0.30)
ENERGY_ROUND_SCALE = 1_000_000_000  # 字典序比较：10^-9 kWh
TIME_ROUND_SCALE = 1_000           # 字典序比较：10^-3 s


@dataclass(frozen=True)
class Model:
    code: str
    empty_mass: float
    max_payload: float
    max_volume: float
    cruise_speed: float
    range_empty: float
    range_full: float
    usable_energy: float
    base_reserve: float
    preparation: float
    load_each: float
    handover: float
    handover_each: float
    climb_speed: float
    descent_speed: float
    climb_efficiency: float


@dataclass(frozen=True)
class Node:
    code: str
    lon: float
    lat: float
    ground: float


@dataclass(frozen=True)
class Box:
    code: str
    service: str
    kind: str
    weight: float
    volume: float


@dataclass(frozen=True)
class Route:
    service: str
    distance: float
    highest_ground: float
    cruise_altitude: float
    outward_up: float
    outward_down: float
    return_up: float
    return_down: float


@dataclass(frozen=True)
class Pattern:
    model: str
    counts: tuple[int, ...]
    weight: float
    volume: float
    nboxes: int
    energy: float
    time: float
    flight_time: float
    energy_units: int
    time_units: int


class CannotDeliver(ValueError):
    pass


def input_file(directory: Path, stem: str, suffix: str) -> Path:
    """识别原文件及 Windows 下载重名后的 (1)、(1) (2) 等文件名。"""
    rx = re.compile(re.escape(stem) + r"(?P<copies>(?:\s*\(\d+\))*)" + re.escape(suffix), re.I)

    def choose(paths) -> Path | None:
        choices = []
        for path in paths:
            if not path.is_file():
                continue
            match = rx.fullmatch(path.name)
            if match:
                numbers = tuple(int(x) for x in re.findall(r"\((\d+)\)", match.group("copies")))
                choices.append((numbers, path.name, path))
        return max(choices)[2] if choices else None

    selected = choose(directory.iterdir())
    if selected is None and suffix.lower() == ".mat":
        # 用户常将补下载的 DEM 留在“下载”，或将原数据放进子目录。
        selected = choose(directory.rglob("*.mat"))
        for location in (Path(__file__).resolve().parent, Path.home() / "Downloads"):
            if selected is not None:
                break
            if location.is_dir():
                selected = choose(location.iterdir())
    if selected is None:
        if suffix.lower() == ".mat":
            raise FileNotFoundError(
                f"没有找到 DEM 的 .mat 数据文件。已检查 {directory}（含子文件夹）、"
                f"脚本所在目录和“下载”目录。请下载并保留真正的 {stem}(1).mat 文件，"
                "或者通过 --dem-file 指定该文件的完整路径；不能将 .tif 改名为 .mat。"
            )
        raise FileNotFoundError(f"找不到 {stem}{suffix}，请检查输入目录：{directory}")
    print(f"读取 {selected}")
    return selected


def load_models(path: Path) -> dict[str, Model]:
    ws = load_workbook(path, data_only=True, read_only=True).active
    result = {}
    for row in ws.iter_rows(values_only=True):
        if row[0] not in ("A", "B", "C") or len(row) < 17 or row[3] is None:
            continue
        x = list(row)
        r = float(x[9]) / 100.0 if float(x[9]) > 1 else float(x[9])
        result[x[0]] = Model(
            code=x[0], empty_mass=float(x[2]), max_payload=float(x[3]),
            max_volume=float(x[4]), cruise_speed=float(x[5]),
            range_empty=float(x[6]), range_full=float(x[7]),
            usable_energy=float(x[8]), base_reserve=r,
            preparation=float(x[10]), load_each=float(x[11]),
            handover=float(x[12]), handover_each=float(x[13]),
            climb_speed=float(x[14]), descent_speed=float(x[15]),
            climb_efficiency=float(x[16]),
        )
    if set(result) != {"A", "B", "C"}:
        raise ValueError("运输无人机数据中没有完整的 A、B、C 三种机型参数")
    for g in result.values():
        if min(g.max_payload, g.max_volume, g.range_full, g.usable_energy,
               g.climb_speed, g.descent_speed, g.cruise_speed,
               g.climb_efficiency) <= 0 or not 0 <= g.base_reserve < 1:
            raise ValueError(f"{g.code} 的参数存在非正值或不合法余量")
    return result


def load_nodes(path: Path) -> tuple[Node, dict[str, Node]]:
    ws = load_workbook(path, data_only=True, read_only=True).active
    center = None
    sites = {}
    for row in ws.iter_rows(values_only=True):
        code = row[0]
        if code == "O01" or (isinstance(code, str) and re.fullmatch(r"S\d{3}", code)):
            node = Node(code, float(row[2]), float(row[3]), float(row[4]))
            if code == "O01":
                center = node
            else:
                sites[code] = node
    if center is None or not sites:
        raise ValueError("节点表找不到 O01 或服务区")
    return center, sites


def load_boxes(path: Path) -> dict[str, list[Box]]:
    wb = load_workbook(path, data_only=True, read_only=True)
    if "逐箱货箱清单" not in wb.sheetnames:
        raise ValueError("物资表必须包含“逐箱货箱清单”工作表")
    ws = wb["逐箱货箱清单"]
    boxes = defaultdict(list)
    seen = set()
    for row in list(ws.iter_rows(values_only=True))[1:]:
        if row[0] is None:
            continue
        if row[0] in seen:
            raise ValueError(f"重复的货箱编号：{row[0]}")
        seen.add(row[0])
        box = Box(str(row[0]), str(row[1]), str(row[2]), float(row[3]), float(row[4]))
        if box.weight <= 0 or box.volume <= 0:
            raise ValueError(f"货箱 {box.code} 的重量或体积不合法")
        boxes[box.service].append(box)
    if not boxes:
        raise ValueError("逐箱货箱清单为空")
    return dict(boxes)


class DemGrid:
    """用 .mat 内的 DEM 和像元中心经纬度定位穿越的栅格。"""

    def __init__(self, mat_file: Path) -> None:
        raw = loadmat(mat_file, variable_names=["dem", "latitude", "longitude", "nodata", "epsg_code"])
        self.dem = np.asarray(raw["dem"])
        lat = np.asarray(raw["latitude"]).reshape(-1)
        lon = np.asarray(raw["longitude"]).reshape(-1)
        self.nodata = float(np.asarray(raw["nodata"]).item())
        if int(np.asarray(raw["epsg_code"]).item()) != 4326 or self.dem.shape != (len(lat), len(lon)):
            raise ValueError("DEM 经纬度参考系或数组维度与预期不符")
        self.dx = float(np.median(np.diff(lon)))
        self.dy = float(-np.median(np.diff(lat)))
        if self.dx <= 0 or self.dy <= 0 or not np.allclose(np.diff(lon), self.dx, atol=1e-9) or not np.allclose(np.diff(lat), -self.dy, atol=1e-9):
            raise ValueError("DEM 经纬度列表需要是均匀网格，经度升序、纬度降序")
        self.left = float(lon[0] - self.dx / 2)
        self.top = float(lat[0] + self.dy / 2)
        self.nrows, self.ncols = self.dem.shape

    def coordinate(self, node: Node) -> tuple[float, float]:
        x = (node.lon - self.left) / self.dx
        y = (self.top - node.lat) / self.dy
        if not (0 <= x < self.ncols and 0 <= y < self.nrows):
            raise ValueError(f"{node.code} 的经纬度超出 DEM 覆盖范围")
        return x, y

    def line_maximum(self, start: Node, end: Node) -> float:
        """栅格 DDA 遍历；恰好穿过角点时把三个接触栅格均计入。"""
        x0, y0 = self.coordinate(start)
        x1, y1 = self.coordinate(end)
        col, row = floor(x0), floor(y0)
        target_col, target_row = floor(x1), floor(y1)
        delta_x, delta_y = x1 - x0, y1 - y0
        step_x = (delta_x > 0) - (delta_x < 0)
        step_y = (delta_y > 0) - (delta_y < 0)
        tdx = 1 / abs(delta_x) if step_x else inf
        tdy = 1 / abs(delta_y) if step_y else inf
        tmx = ((col + 1 - x0) / delta_x if step_x > 0 else
               (x0 - col) / (-delta_x) if step_x < 0 else inf)
        tmy = ((row + 1 - y0) / delta_y if step_y > 0 else
               (y0 - row) / (-delta_y) if step_y < 0 else inf)
        max_height = -inf

        def update(r: int, c: int) -> None:
            nonlocal max_height
            if not (0 <= r < self.nrows and 0 <= c < self.ncols):
                raise ValueError(f"{start.code}-{end.code} 航段越过 DEM 边界")
            z = float(self.dem[r, c])
            if not np.isfinite(z) or np.isclose(z, self.nodata):
                raise ValueError(f"{start.code}-{end.code} 航段遇到无效 DEM 栅格 {(r, c)}")
            max_height = max(max_height, z)

        update(row, col)
        for _ in range(self.nrows + self.ncols + 5):
            if (col, row) == (target_col, target_row):
                return max_height
            if abs(tmx - tmy) <= 1e-12:
                update(row, col + step_x)
                update(row + step_y, col)
                col += step_x
                row += step_y
                tmx += tdx
                tmy += tdy
            elif tmx < tmy:
                col += step_x
                tmx += tdx
            else:
                row += step_y
                tmy += tdy
            update(row, col)
        raise RuntimeError("DEM 航段遍历未收敛")


def horizontal_distance(a: Node, b: Node) -> float:
    dlat = radians(b.lat - a.lat)
    dlon = radians(b.lon - a.lon)
    s = sin(dlat / 2) ** 2 + cos(radians(a.lat)) * cos(radians(b.lat)) * sin(dlon / 2) ** 2
    return 2 * EARTH_RADIUS_M * atan2(sqrt(s), sqrt(max(0, 1 - s)))


def make_routes(dem: DemGrid, center: Node, sites: dict[str, Node]) -> dict[str, Route]:
    routes = {}
    for code, node in sites.items():
        highest = dem.line_maximum(center, node)
        altitude = highest + 50.0
        out_up = altitude - center.ground
        out_down = altitude - (node.ground + 30.0)
        if min(out_up, out_down) < -1e-7:
            raise ValueError(f"{code} 的巡航高度低于节点作业高度，请核对 DEM 与海拔数据")
        routes[code] = Route(code, horizontal_distance(center, node), highest, altitude,
                             max(out_up, 0), max(out_down, 0), max(out_down, 0), max(out_up, 0))
    return routes


def equivalent_range(model: Model, payload: float) -> float:
    if not 0 <= payload <= model.max_payload + 1e-7:
        raise ValueError(f"机型 {model.code} 的载荷 {payload} kg 超出范围")
    return model.range_empty - (model.range_empty - model.range_full) * (payload / model.max_payload) ** 1.5


def segment_energy(model: Model, distance: float, ascent: float, payload: float) -> float:
    """返回水平巡航与爬升两项能耗之和，单位 kWh。"""
    horizontal = model.usable_energy * distance / equivalent_range(model, payload)
    climbing = (model.empty_mass + payload) * GRAVITY * ascent / (model.climb_efficiency * 3_600_000)
    return horizontal + climbing


def round_trip_energy(model: Model, route: Route, payload: float) -> float:
    return (segment_energy(model, route.distance, route.outward_up, payload)
            + segment_energy(model, route.distance, route.return_up, 0))


def flight_time(model: Model, route: Route) -> float:
    return ((route.outward_up + route.return_up) / model.climb_speed
            + (route.outward_down + route.return_down) / model.descent_speed
            + 2 * route.distance / model.cruise_speed)


def safe_payload(model: Model, route: Route, reserve: float) -> float | None:
    if not 0 <= reserve < 1:
        raise ValueError(f"安全余量应在 [0,1)：{reserve}")
    budget = (1 - reserve) * model.usable_energy
    if round_trip_energy(model, route, 0.0) > budget + 1e-10:
        return None  # 空载也不能安全往返
    if round_trip_energy(model, route, model.max_payload) <= budget:
        return model.max_payload
    low, high = 0.0, model.max_payload
    for _ in range(65):
        mid = (low + high) / 2
        if round_trip_energy(model, route, mid) <= budget + 1e-10:
            low = mid
        else:
            high = mid
    return low  # 向可行侧取值


def distinct_types(boxes: list[Box]) -> tuple[list[tuple[str, float, float]], dict[tuple, deque[str]]]:
    groups: dict[tuple[str, float, float], deque[str]] = {}
    for box in boxes:
        key = (box.kind, box.weight, box.volume)
        groups.setdefault(key, deque()).append(box.code)
    return list(groups), groups


def make_patterns(types: list[tuple[str, float, float]], limits: tuple[int, ...],
                  models: dict[str, Model], route: Route, capacities: dict[str, float | None]) -> list[Pattern]:
    patterns = []
    # 同种货箱仅需记录件数；它们的单箱质量、体积和作业时间完全相同。
    for counts in product(*(range(n + 1) for n in limits)):
        nboxes = sum(counts)
        if not nboxes:
            continue
        weight = sum(n * typ[1] for n, typ in zip(counts, types))
        volume = sum(n * typ[2] for n, typ in zip(counts, types))
        for g in models.values():
            cap = capacities[g.code]
            if cap is None or weight > min(cap, g.max_payload) + 1e-8 or volume > g.max_volume + 1e-9:
                continue
            e = round_trip_energy(g, route, weight)
            t_fly = flight_time(g, route)
            t = g.preparation + nboxes * g.load_each + t_fly + g.handover + nboxes * g.handover_each
            patterns.append(Pattern(g.code, counts, weight, volume, nboxes, e, t, t_fly,
                                    round(e * ENERGY_ROUND_SCALE), round(t * TIME_ROUND_SCALE)))
    return patterns


def optimize_service(boxes: list[Box], models: dict[str, Model], route: Route,
                     capacities: dict[str, float | None]) -> list[tuple[Pattern, list[str]]]:
    types, groups = distinct_types(boxes)
    limits = tuple(len(groups[t]) for t in types)
    patterns = make_patterns(types, limits, models, route, capacities)
    by_first = [[p for p in patterns if p.counts[k] > 0] for k in range(len(types))]
    for k, typ in enumerate(types):
        if not any(p.counts[k] == 1 and p.nboxes == 1 for p in patterns):
            raise CannotDeliver(f"{route.service} 的 {typ[0]} 单箱无法由任何机型安全运输")

    @lru_cache(maxsize=None)
    def dp(remaining: tuple[int, ...]) -> tuple[tuple[int, int, int], Pattern | None]:
        if not any(remaining):
            return (0, 0, 0), None
        first = next(k for k, n in enumerate(remaining) if n)
        best: tuple[int, int, int] | None = None
        chosen = None
        for p in by_first[first]:
            if any(n > r for n, r in zip(p.counts, remaining)):
                continue
            rest = tuple(r - n for r, n in zip(remaining, p.counts))
            value, _ = dp(rest)
            cost = (1 + value[0], p.energy_units + value[1], p.time_units + value[2])
            if best is None or cost < best:
                best, chosen = cost, p
        if best is None:
            raise CannotDeliver(f"{route.service} 找不到完整可行组批")
        return best, chosen

    dp(limits)
    result = []
    remaining = limits
    while any(remaining):
        pattern = dp(remaining)[1]
        if pattern is None:
            raise RuntimeError("组批回溯错误")
        ids = []
        for n, typ in zip(pattern.counts, types):
            ids.extend(groups[typ].popleft() for _ in range(n))
        result.append((pattern, ids))
        remaining = tuple(r - n for r, n in zip(remaining, pattern.counts))
    original_ids = {b.code for b in boxes}
    assigned = [code for _, ids in result for code in ids]
    if len(assigned) != len(original_ids) or set(assigned) != original_ids:
        raise AssertionError(f"{route.service} 存在漏箱或重复配送")
    return result


def sheet(wb: Workbook, name: str, headers: list[str], rows: list[tuple]) -> None:
    ws = wb.create_sheet(name)
    ws.append(headers)
    for row in rows:
        ws.append(row)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor="17365D")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 31
    for col in ws.columns:
        maxlen = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[col[0].column_letter].width = min(max(maxlen * 1.25 + 3, 15), 64)
        for c in col[1:]:
            if isinstance(c.value, float):
                c.number_format = "0.000000" if abs(c.value) < 1 else "0.000"
            if isinstance(c.value, str) and len(c.value) > 55:
                c.alignment = Alignment(wrap_text=True, vertical="center")
                ws.row_dimensions[c.row].height = min(85, max(ws.row_dimensions[c.row].height or 15,
                                                                 15 * ((len(c.value) + 54) // 55)))


def scenarios(models: dict[str, Model]) -> list[tuple[str, dict[str, float]]]:
    base = {code: g.base_reserve for code, g in models.items()}
    items = []
    if not any(all(abs(x - r) < 1e-12 for x in base.values()) for r in SENSITIVITY):
        items.append(("附件基准", base))
    for r in SENSITIVITY:
        label = f"{int(round(r * 100))}%"
        if all(abs(x - r) < 1e-12 for x in base.values()):
            label += "（附件基准）"
        items.append((label, {code: r for code in models}))
    return items


def solve(data_dir: Path, output: Path, dem_file: Path | None = None) -> None:
    models = load_models(input_file(data_dir, "运输无人机数据", ".xlsx"))
    center, sites = load_nodes(input_file(data_dir, "调度中心与服务区", ".xlsx"))
    boxes = load_boxes(input_file(data_dir, "物资需求与配送时限", ".xlsx"))
    dem = DemGrid(dem_file if dem_file is not None else input_file(data_dir, "镇龙乡及周边30米DEM", ".mat"))
    if unknown := set(boxes) - set(sites):
        raise ValueError(f"需求清单中的服务区没有节点坐标：{sorted(unknown)}")
    if missing := set(sites) - set(boxes):
        raise ValueError(f"服务区没有货箱：{sorted(missing)}")
    routes = make_routes(dem, center, sites)
    total_boxes = sum(len(v) for v in boxes.values())
    print(f"读取 {len(sites)} 个服务区、{total_boxes} 个货箱。")

    route_rows, capacity_rows, batch_rows, service_rows, summary_rows = [], [], [], [], []
    for code in sorted(sites):
        r = routes[code]
        route_rows.append((code, center.ground, sites[code].ground, r.distance,
                           r.highest_ground, r.cruise_altitude, r.outward_up,
                           r.outward_down, r.return_up, r.return_down))

    for case, reserves in scenarios(models):
        capacity = {}
        for code in sorted(sites):
            r = routes[code]
            capacity[code] = {}
            for g in models.values():
                rho = reserves[g.code]
                q = safe_payload(g, r, rho)
                capacity[code][g.code] = q
                capacity_rows.append((case, code, g.code, rho, q, "可往返" if q is not None else "空载也不可往返",
                                      (1 - rho) * g.usable_energy,
                                      round_trip_energy(g, r, 0),
                                      round_trip_energy(g, r, q) if q is not None else None))

        case_n = 0
        case_energy = 0.0
        case_time = 0.0
        failures = []
        for code in sorted(sites):
            r = routes[code]
            try:
                batches = optimize_service(boxes[code], models, r, capacity[code])
            except CannotDeliver as exc:
                failures.append(str(exc))
                service_rows.append((case, code, len(boxes[code]), None, None, None, str(exc)))
                continue
            n, e, t = len(batches), sum(p.energy for p, _ in batches), sum(p.time for p, _ in batches)
            case_n += n
            case_energy += e
            case_time += t
            service_rows.append((case, code, len(boxes[code]), n, e, t, "可行"))
            for k, (p, ids) in enumerate(batches, 1):
                g = models[p.model]
                if p.energy > (1 - reserves[p.model]) * g.usable_energy + 1e-8:
                    raise AssertionError(f"{case} {code} 第{k}批：余量校验不通过")
                batch_rows.append((case, code, f"{code}-{k:02d}", p.model,
                                   ", ".join(ids), p.nboxes, p.weight, p.volume,
                                   p.energy, p.flight_time, p.time,
                                   reserves[p.model], (1 - reserves[p.model]) * g.usable_energy - p.energy))
        valid = not failures
        summary_rows.append((case, total_boxes, case_n if valid else None,
                             case_energy if valid else None, case_time if valid else None,
                             "全部交付" if valid else "不可行：" + "；".join(failures)))
        msg = f"{case}: {case_n} 架次，能耗 {case_energy:.4f} kWh，累计作业 {case_time:.1f} s"
        print(msg if valid else f"{case}: 不可行；" + "；".join(failures))

    wb = Workbook()
    wb.remove(wb.active)
    sheet(wb, "余量汇总", ["安全余量情景", "货箱总数", "总架次数", "总能耗(kWh)", "累计作业时间(s)", "状态"], summary_rows)
    sheet(wb, "组批方案", ["情景", "服务区", "架次编号", "机型", "货箱编号（同一服务区）", "箱数", "总质量(kg)",
                      "总体积(m³)", "往返能耗(kWh)", "往返飞行时间(s)", "累计作业时间(s)", "安全余量", "能量富余(kWh)"], batch_rows)
    sheet(wb, "服务区汇总", ["情景", "服务区", "货箱数", "架次数", "能耗(kWh)", "累计作业时间(s)", "状态"], service_rows)
    sheet(wb, "最大安全载荷", ["情景", "服务区", "机型", "安全余量", "最大安全载荷(kg)", "状态",
                          "允许能耗(kWh)", "空载往返能耗(kWh)", "满安全载荷往返能耗(kWh)"], capacity_rows)
    sheet(wb, "航线参数", ["服务区", "O01地面高程(m)", "服务区地面高程(m)", "单程水平距离(m)",
                      "航线最高DEM(m)", "巡航海拔(m)", "去程爬升(m)", "去程下降(m)", "返程爬升(m)", "返程下降(m)"], route_rows)
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    print(f"结果已保存：{output}")


def main() -> None:
    parser = argparse.ArgumentParser(description="第一问：单服务区往返最大安全载荷与货箱组批")
    parser.add_argument("--data-dir", type=Path, default=DEFAULT_DIR, help="含四份输入文件的目录")
    parser.add_argument("--dem-file", type=Path, default=None, help="可选：DEM .mat 文件的完整路径")
    parser.add_argument("--out", type=Path, default=None, help="输出 xlsx 路径，默认在数据目录内")
    args = parser.parse_args()
    if not args.data_dir.is_dir():
        parser.error(f"输入目录不存在：{args.data_dir}。可用 --data-dir 指定实际位置。")
    solve(args.data_dir, args.out or args.data_dir / "问题一_单点往返组批结果.xlsx", args.dem_file)


if __name__ == "__main__":
    main()
